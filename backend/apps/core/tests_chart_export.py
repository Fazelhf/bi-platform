"""The compare window's Excel: the same table, total line and chart as the screen."""
from io import BytesIO

from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.test import APITestCase


class ChartExportTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("ce_user", password="x")
        self.client.force_authenticate(self.user)

    def spec(self, **extra):
        return {
            "title": "فروش ریالی — به تفکیک کارشناس",
            "subtitle": "مقایسه‌ی مهر 1405، آبان 1405",
            "percent": False,
            "chart": {
                "categories": ["علی", "سارا"],
                "series": [{"name": "مهر 1405", "values": [100, 50]},
                           {"name": "آبان 1405", "values": [120, 40]}],
            },
            "table": {
                "head": ["کارشناس", "مهر 1405", "آبان 1405"],
                "rows": [["علی", 100, 120], ["سارا", 50, 40]],
                "foot": ["مجموع", 150, 160],
            },
            "headline": [{"label": "مهر 1405", "value": 150}, {"label": "آبان 1405", "value": 160}],
            "change": 6.7,
            **extra,
        }

    def test_the_file_holds_what_the_window_shows(self):
        res = self.client.post("/api/executive/export/chart/", self.spec(), format="json")
        self.assertEqual(res.status_code, 200)
        self.assertIn(".xlsx", res["Content-Disposition"])
        wb = load_workbook(BytesIO(res.content))
        ws = wb["نمودار"]
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertEqual(ws["A1"].value, "فروش ریالی — به تفکیک کارشناس")
        table = [[c.value for c in row] for row in ws.iter_rows(min_row=4, max_row=7, max_col=3)]
        self.assertEqual(table, [
            ["کارشناس", "مهر 1405", "آبان 1405"],
            ["علی", 100, 120], ["سارا", 50, 40], ["مجموع", 150, 160],
        ])
        self.assertEqual(len(ws._charts), 1)
        data = wb["داده‌ی نمودار"]
        self.assertEqual([c.value for c in data[1]][1:], ["مهر 1405", "آبان 1405"])

    def test_percent_cells_carry_a_percent_format(self):
        res = self.client.post(
            "/api/executive/export/chart/", self.spec(percent=True), format="json",
        )
        ws = load_workbook(BytesIO(res.content))["نمودار"]
        self.assertIn("٪", ws["B5"].number_format)

    def test_nothing_to_export_is_refused(self):
        res = self.client.post("/api/executive/export/chart/", {"title": "x"}, format="json")
        self.assertEqual(res.status_code, 400)
