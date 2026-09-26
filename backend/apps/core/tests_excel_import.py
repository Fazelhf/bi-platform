"""ورود اکسل: sample template → check (nothing saved) → confirm, for each importer."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from apps.core.excel_import import registry
from apps.core.models import DimPeriod, PeriodKind
from apps.crm.models import Customer, Product
from apps.finance.models import BankAccount, CashCategory, CashMovement
from apps.sales2.models import AccountingCost, CustomerAccount, PriceListItem, Receipt

URL = "/api/imports"


def xlsx(header, *rows) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "f.xlsx"
    return buf


class ImportBase(APITestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user("root", password="pw12345!", role="admin", is_superuser=True)
        self.client.force_authenticate(self.admin)
        self.customer = Customer.objects.create(code="c1", name_fa="شرکت الف", national_id="10100000001",
                                                first_contact_at=timezone.now())
        self.tape = Product.objects.create(code="p1", name_fa="چسب نواری")

    def run_import(self, key, file, confirm=False, **params):
        data = {"file": file, **params}
        if confirm:
            data["confirm"] = "1"
        return self.client.post(f"{URL}/{key}/run/", data, format="multipart")


class TemplateTests(ImportBase):
    def test_every_template_downloads_with_a_guide(self):
        for key in registry():
            r = self.client.get(f"{URL}/{key}/template/")
            self.assertEqual(r.status_code, 200, key)
            wb = load_workbook(BytesIO(r.content))
            if key != "sales2-price-list":
                self.assertIn("راهنما", wb.sheetnames, key)

    def test_list_is_filtered_by_section(self):
        keys = {i["key"] for i in self.client.get(f"{URL}/", {"section": "finance"}).data}
        self.assertEqual(keys, {"finance-cash", "finance-budget-actuals"})

    def test_sales_user_cannot_use_sales2_imports(self):
        self.client.force_authenticate(get_user_model().objects.create_user("s", password="x", role="manager"))
        self.assertEqual(self.client.get(f"{URL}/sales2-costs/template/").status_code, 403)

    def test_missing_required_column_is_reported(self):
        r = self.run_import("sales2-costs", xlsx(["نام"], ["x"]), year=1405, month=7)
        self.assertEqual(r.status_code, 400)


class Sales2ImportTests(ImportBase):
    def test_costs_check_then_confirm(self):
        AccountingCost.objects.create(product=self.tape, grammage=0, cost_rial=100, jalali_year=1405, jalali_month=1)
        f = lambda: xlsx(["نام کالا", "گرماژ", "فی حسابداری"], ["چسب نواری", None, 120], ["ناشناخته", None, 5])
        r = self.run_import("sales2-costs", f(), year=1405, month=7)
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["counts"]["changed"], 1)
        self.assertEqual(r.data["counts"]["error"], 1)
        self.assertEqual(AccountingCost.objects.count(), 1)  # preview saved nothing
        r = self.run_import("sales2-costs", f(), confirm=True, year=1405, month=7)
        self.assertEqual(r.data["written"], 1)
        self.assertEqual(AccountingCost.objects.get(jalali_month=7).cost_rial, 120)

    def test_fixed_prices(self):
        f = xlsx(["نام کالا", "قیمت رسمی", "قیمت غیررسمی"], ["چسب نواری", "1,000", 900])
        r = self.run_import("sales2-fixed-prices", f, confirm=True, year=1405, month=7)
        self.assertEqual(r.data["counts"]["new"], 1, r.data)
        self.assertEqual(PriceListItem.objects.filter(product=self.tape).count(), 2)

    def test_credit_by_code_and_duplicate(self):
        f = xlsx(["کد مشتری", "سقف اعتبار", "مهلت پرداخت", "توقف فروش"],
                 ["c1", 5000, 30, "بله"], ["c1", 6000, None, None], ["zz", 1, None, None])
        r = self.run_import("sales2-credit", f, confirm=True)
        self.assertEqual(r.data["counts"], {"new": 1, "changed": 0, "same": 0, "error": 2}, r.data)
        acc = CustomerAccount.objects.get(customer=self.customer)
        self.assertEqual((acc.credit_limit_rial, acc.credit_days, acc.on_hold), (5000, 30, True))

    def test_receipts_pending_finance_and_deduped(self):
        f = lambda: xlsx(["نام مشتری", "تاریخ", "روش", "مبلغ", "شماره پیگیری"],
                         ["شرکت الف", "۱۴۰۵/۰۷/۰۱", "واریز", 500, "T1"],
                         ["شرکت الف", "1405/07/02", "چک", 400, ""])
        r = self.run_import("sales2-receipts", f(), confirm=True)
        self.assertEqual(r.data["counts"]["new"], 1, r.data)
        self.assertIn("سررسید", r.data["rows"][1]["message"])
        rec = Receipt.objects.get()
        self.assertEqual(rec.finance_status, Receipt.FinanceStatus.PENDING)
        r = self.run_import("sales2-receipts", f())
        self.assertEqual(r.data["counts"]["same"], 1)


class FinanceImportTests(ImportBase):
    def setUp(self):
        super().setUp()
        month = DimPeriod.objects.create(jalali_year=1405, jalali_month=6, kind=PeriodKind.MONTH,
                                         start_date=date(2026, 8, 23), end_date=date(2026, 9, 22))
        self.day = DimPeriod.objects.create(jalali_year=1405, jalali_month=6, kind=PeriodKind.DAY, parent=month,
                                            seq=1, start_date=date(2026, 8, 23), end_date=date(2026, 8, 23))
        BankAccount.objects.create(title="جاری ملت")
        self.cat = CashCategory.objects.get(code="sales-other")

    def test_cash_movement_draft_then_update(self):
        f = lambda amount: xlsx(["تاریخ", "جهت", "دسته", "حساب", "مبلغ"],
                                ["1405/06/01", "واریز", "sales-other", "جاری ملت", amount])
        r = self.run_import("finance-cash", f(1000), confirm=True)
        self.assertEqual(r.data["counts"]["new"], 1, r.data)
        mv = CashMovement.objects.get()
        self.assertEqual((mv.period_id, mv.amount_rial, mv.status), (self.day.id, Decimal(1000), "draft"))
        r = self.run_import("finance-cash", f(1500))
        self.assertEqual(r.data["counts"]["changed"], 1)

    def test_outsider_refused(self):
        self.client.force_authenticate(get_user_model().objects.create_user(
            "b2b", password="x", role="manager", department="sales_b2b"))
        self.assertEqual(self.client.get(f"{URL}/finance-cash/template/").status_code, 403)
