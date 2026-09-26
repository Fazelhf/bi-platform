"""
ورود اکسل — one way in for every sheet of rows the platform accepts.

Each kind of import is an `Importer`: its columns, who may use it, how one
row is checked against what is already stored, and how a checked row is
written. The flow is the same for all of them, and nothing is written until
the last step:

    نمونه (template)  → a workbook with the right headers, two sample rows and
                        a «راهنما» sheet describing every column;
    بررسی (preview)   → every row read, typed and checked: new / changed /
                        same / error-with-reason, nothing saved;
    تأیید (apply)     → the same check again, then the valid rows written; rows
                        with an error are skipped and reported.

Columns are found by header text, not position, so a user may reorder them or
add columns of their own; Persian/Arabic letter variants and digits are
folded before comparing.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.core import jalali

_FOLD = str.maketrans({"ي": "ی", "ك": "ک", "‌": " ", "ۀ": "ه",
                       **{d: str(i) for i, d in enumerate("۰۱۲۳۴۵۶۷۸۹")},
                       **{d: str(i) for i, d in enumerate("٠١٢٣٤٥٦٧٨٩")}})


def fold(text) -> str:
    return re.sub(r"\s+", " ", str(text if text is not None else "").translate(_FOLD)).strip()


@dataclass
class Col:
    name: str
    kind: str = "text"  # text | int | money | date | bool | choice
    required: bool = False
    help: str = ""
    choices: dict | None = None  # label → value, for kind="choice"
    aliases: tuple = ()


@dataclass
class Param:
    """Something the whole file shares, asked once in the page (e.g. the month)."""
    name: str
    label: str
    kind: str = "month"  # month | select
    options: object = None  # for select: [{"value", "label"}], or a callable giving them


@dataclass
class Row:
    n: int
    values: dict
    sheet: str = ""
    status: str = "new"  # new | changed | same | error
    message: str = ""
    key: str = ""
    data: dict = field(default_factory=dict)  # what apply() needs, typed

    def as_dict(self) -> dict:
        return {"n": self.n, "sheet": self.sheet, "values": {k: _plain(v) for k, v in self.values.items()},
                "status": self.status, "message": self.message}


def _plain(v):
    if isinstance(v, Decimal):
        return str(v)
    if isinstance(v, (date, datetime)):
        return v.isoformat()[:10]
    return v


class Importer:
    key = ""
    title = ""
    section = ""
    description = ""
    columns: list[Col] = []
    params: list[Param] = []
    sample: list[list] = []
    #: Read every sheet that carries the headers, not only the first — for
    #: company workbooks that split one list over several sheets.
    all_sheets = False

    # -- to override -------------------------------------------------------
    def allowed(self, user) -> bool:
        raise NotImplementedError

    def check(self, row: Row, ctx: dict) -> None:
        """Set row.status / message / data. Typed values are in row.values."""
        raise NotImplementedError

    def write(self, rows: list[Row], ctx: dict, user) -> int:
        raise NotImplementedError

    def context(self, params: dict, user) -> dict:
        return {"params": params}

    # -- shared --------------------------------------------------------------
    def describe(self) -> dict:
        return {
            "key": self.key, "title": self.title, "section": self.section,
            "description": self.description,
            "columns": [{"name": c.name, "required": c.required, "kind": c.kind, "help": c.help,
                         "choices": list(c.choices) if c.choices else None} for c in self.columns],
            "params": [{"name": p.name, "label": p.label, "kind": p.kind,
                        "options": p.options() if callable(p.options) else p.options}
                       for p in self.params],
        }

    def template(self) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "داده"
        ws.sheet_view.rightToLeft = True
        head_fill = PatternFill("solid", fgColor="1C1C1E")
        req_fill = PatternFill("solid", fgColor="7C2D12")
        for i, c in enumerate(self.columns, 1):
            cell = ws.cell(row=1, column=i, value=c.name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = req_fill if c.required else head_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[cell.column_letter].width = max(14, len(c.name) + 6)
        for r, values in enumerate(self.sample, 2):
            for i, v in enumerate(values, 1):
                ws.cell(row=r, column=i, value=v)
        ws.freeze_panes = "A2"

        guide = wb.create_sheet("راهنما")
        guide.sheet_view.rightToLeft = True
        guide.append([self.title])
        guide["A1"].font = Font(bold=True, size=13)
        if self.description:
            guide.append([self.description])
        guide.append([])
        guide.append(["ستون", "الزامی", "نوع", "توضیح"])
        for cell in guide[guide.max_row]:
            cell.font = Font(bold=True)
        kinds = {"text": "متن", "int": "عدد صحیح", "money": "مبلغ (ریال)", "date": "تاریخ شمسی ۱۴۰۵/۰۷/۰۱",
                 "bool": "بله / خیر", "choice": "یکی از گزینه‌ها"}
        for c in self.columns:
            extra = f" — گزینه‌ها: {'، '.join(c.choices)}" if c.choices else ""
            guide.append([c.name, "بله" if c.required else "", kinds[c.kind], c.help + extra])
        guide.append([])
        guide.append(["ستون‌های قرمز الزامی‌اند. ترتیب ستون‌ها مهم نیست. ردیف‌های نمونه را پاک کنید."])
        for col, w in zip("ABCD", (22, 10, 22, 70)):
            guide.column_dimensions[col].width = w
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def read(self, fileobj) -> list[Row]:
        from openpyxl import load_workbook

        try:
            wb = load_workbook(fileobj, data_only=True)
        except Exception:
            raise ValidationError({"file": "فایل اکسل خوانده نشد (فقط .xlsx)."})
        if "داده" in wb.sheetnames:
            sheets = [wb["داده"]]
        else:
            sheets = [s for s in wb.worksheets if s.title != "راهنما"]
            sheets = sheets if self.all_sheets else sheets[:1]
        rows: list[Row] = []
        missing: list[str] = []
        for ws in sheets:
            got = self._read_sheet(ws, len(sheets) > 1)
            if got is None:
                missing = missing or self._missing(ws)
                continue
            rows.extend(got)
        if not rows and missing:
            raise ValidationError({"file": "ستون‌های الزامی پیدا نشد: " + "، ".join(missing)})
        return rows

    def _headers(self, ws):
        required = {c.name for c in self.columns if c.required}
        for r in ws.iter_rows(min_row=1, max_row=6):
            exact, alias = {}, {}
            for cell in r:
                h = fold(cell.value)
                if not h:
                    continue
                for c in self.columns:
                    if h == fold(c.name):
                        exact.setdefault(c.name, cell.column_letter)
                    elif h in {fold(a) for a in c.aliases}:
                        alias.setdefault(c.name, cell.column_letter)
            found = {**alias, **exact}
            if required <= set(found):
                return r[0].row, found
        return None, {}

    def _missing(self, ws) -> list[str]:
        best: set = set()
        for r in ws.iter_rows(min_row=1, max_row=6):
            heads = {fold(c.value) for c in r}
            best = max(best, {c.name for c in self.columns if fold(c.name) in heads}, key=len)
        return [c.name for c in self.columns if c.required and c.name not in best]

    def _read_sheet(self, ws, tag: bool):
        header_row, cols = self._headers(ws)
        if header_row is None:
            return None
        rows = []
        for r in range(header_row + 1, ws.max_row + 1):
            raw = {name: ws[f"{letter}{r}"].value for name, letter in cols.items()}
            if all(v in (None, "") for v in raw.values()):
                continue
            row = Row(n=r, values={}, sheet=ws.title if tag else "")
            errors = []
            for c in self.columns:
                v = raw.get(c.name)
                try:
                    row.values[c.name] = _typed(c, v)
                except ValueError as e:
                    row.values[c.name] = v
                    errors.append(f"{c.name}: {e}")
                if c.required and row.values.get(c.name) in (None, ""):
                    errors.append(f"{c.name} خالی است")
            if errors:
                row.status, row.message = "error", "؛ ".join(errors)
            rows.append(row)
        return rows

    def run(self, fileobj, params: dict, user) -> tuple[list[Row], dict]:
        ctx = self.context(params, user)
        rows = self.read(fileobj)
        seen: dict[str, int] = {}
        for row in rows:
            if row.status == "error":
                continue
            self.check(row, ctx)
            if row.status != "error" and row.key:
                if row.key in seen:
                    row.status, row.message = "error", f"تکرار ردیف {seen[row.key]} در همین فایل"
                else:
                    seen[row.key] = row.n
        return rows, ctx


def _typed(c: Col, v):
    if v in (None, ""):
        return None
    if c.kind == "text":
        return fold(v)
    if c.kind in ("int", "money"):
        s = fold(v).replace(",", "").replace("٬", "")
        try:
            d = Decimal(s)
        except InvalidOperation:
            raise ValueError("عدد نیست")
        if d < 0:
            raise ValueError("منفی است")
        return int(d) if c.kind == "int" else d.quantize(Decimal(1))
    if c.kind == "bool":
        s = fold(v)
        if s in ("بله", "آری", "1", "True", "true", "دارد", "x", "✓"):
            return True
        if s in ("خیر", "نه", "0", "False", "false", "ندارد"):
            return False
        raise ValueError("بله یا خیر")
    if c.kind == "choice":
        s = fold(v)
        for label, value in (c.choices or {}).items():
            if s == fold(label) or s == str(value):
                return value
        raise ValueError(f"یکی از: {'، '.join(c.choices)}")
    if c.kind == "date":
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        m = re.match(r"^(1[34]\d\d)[/\-.](\d\d?)[/\-.](\d\d?)$", fold(v))
        if not m:
            raise ValueError("تاریخ شمسی مثل ۱۴۰۵/۰۷/۰۱")
        try:
            return jalali.to_gregorian(*(int(x) for x in m.groups()))
        except Exception:
            raise ValueError("تاریخ نامعتبر")
    return v


# ---------------------------------------------------------------------------
# Registry and API
# ---------------------------------------------------------------------------
def registry() -> dict[str, Importer]:
    from apps.finance import importers as fin
    from apps.sales2 import importers as s2

    return {i.key: i for i in (*s2.IMPORTERS, *fin.IMPORTERS)}


def _importer(key: str, user) -> Importer:
    imp = registry().get(key)
    if imp is None:
        raise ValidationError({"detail": "نوع ورود اکسل ناشناخته است."})
    if not imp.allowed(user):
        raise PermissionDenied("این ورود اکسل برای شما نیست.")
    return imp


def _params(imp: Importer, data) -> dict:
    out = {}
    for p in imp.params:
        if p.kind == "month":
            y, m = data.get("year"), data.get("month")
            if not (y and m):
                raise ValidationError({"detail": "سال و ماه را انتخاب کنید."})
            out["year"], out["month"] = int(y), int(m)
        else:
            v = data.get(p.name)
            if v in (None, ""):
                raise ValidationError({"detail": f"«{p.label}» را انتخاب کنید."})
            out[p.name] = v
    return out


class ImportListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        section = request.query_params.get("section")
        return Response([i.describe() for i in registry().values()
                         if i.allowed(request.user) and (not section or i.section == section)])


class ImportTemplateView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, key):
        imp = _importer(key, request.user)
        resp = HttpResponse(imp.template(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="template-{key}.xlsx"'
        return resp


class ImportRunView(APIView):
    """POST file (+params). `confirm=1` writes the valid rows; otherwise a preview."""

    permission_classes = [IsAuthenticated]

    def post(self, request, key):
        imp = _importer(key, request.user)
        upload = request.FILES.get("file")
        if not upload:
            raise ValidationError({"file": "فایل اکسل را انتخاب کنید."})
        params = _params(imp, request.data)
        rows, ctx = imp.run(upload, params, request.user)
        counts = {s: sum(1 for r in rows if r.status == s) for s in ("new", "changed", "same", "error")}
        written = None
        if request.data.get("confirm") in ("1", "true", True):
            with transaction.atomic():
                written = imp.write([r for r in rows if r.status in ("new", "changed")], ctx, request.user)
        return Response({"title": imp.title, "counts": counts, "written": written,
                         "columns": [c.name for c in imp.columns],
                         "rows": [r.as_dict() for r in rows]})
