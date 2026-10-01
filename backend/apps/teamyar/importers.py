"""
ورود اکسل فعالیت‌های تیمیار — the Gantt from a sheet, on the shared
check-then-confirm flow of `apps.core.excel_import`.

Rows are matched by title, so the same file can be edited and uploaded
again: an existing bar is updated in place, never duplicated. An empty cell
leaves the stored value alone. Status and deadline changes write the same
diary lines a change made on the page would.
"""
from __future__ import annotations

from datetime import date

from django.utils import timezone

from apps.core import jalali
from apps.core.excel_import import Col, Importer

from . import charter
from .models import LogEntry, Module, Phase, Task
from .views import can_use_teamyar

STATUSES = {label: value for value, label in Task.Status.choices}


def _jdate(d) -> str:
    if not d:
        return "—"
    y, m, day = jalali.from_gregorian(d)
    return f"{y}/{m:02d}/{day:02d}"


class TeamyarTaskImporter(Importer):
    key = "teamyar-tasks"
    title = "فعالیت‌های استقرار تیمیار"
    section = "teamyar"
    description = ("هر ردیف یک فعالیت گانت. فعالیت با «عنوان» شناخته می‌شود: اگر باشد به‌روز می‌شود، "
                   "وگرنه ساخته می‌شود. خانه‌ی خالی مقدار فعلی را عوض نمی‌کند. فاز تازه خودکار ساخته می‌شود. "
                   "اکسل نمونه همان برنامه‌ی منشور پروژه است.")
    columns = [
        Col("عنوان", "raw", required=True, aliases=("فعالیت", "عنوان فعالیت")),
        Col("فاز", "raw", help="اگر نباشد ساخته می‌شود"),
        Col("ماژول", "raw", help="نام ماژول؛ خالی: از روی عنوان خودکار. نام تازه، ماژول تازه می‌سازد"),
        Col("مسئول", "raw", help="نام آزاد — مشاور تیمیار یا همکار خودمان"),
        Col("شروع", "date", help="برای فعالیت تازه اگر خالی باشد همان پایان"),
        Col("پایان", "date", help="ددلاین؛ برای فعالیت تازه الزامی", aliases=("ددلاین", "پایان / ددلاین")),
        Col("پیشرفت", "int", help="۰ تا ۱۰۰", aliases=("پیشرفت ٪", "درصد پیشرفت")),
        Col("وضعیت", "choice", choices=STATUSES),
        Col("نقطه عطف", "bool", help="بله برای مایلستون (یک‌روزه)", aliases=("مایلستون",)),
        Col("توضیح", "raw", aliases=("شرح", "توضیحات")),
    ]

    def allowed(self, user) -> bool:
        return can_use_teamyar(user)

    def template(self) -> bytes:
        from openpyxl.styles import Alignment

        self.sample = [
            [r["title"], r["phase"], r["module"], r["owner"], _jdate(r["start_on"]), _jdate(r["end_on"]),
             100 if r["done"] else 0, "انجام شده" if r["done"] else "شروع نشده",
             "بله" if r["is_milestone"] else "", r["description"]]
            for r in charter.plan()
        ]
        try:
            raw = super().template()
        finally:
            self.sample = []
        # Wider columns and wrapped notes: this sample is the real plan, read
        # and edited as a document rather than skimmed as two example rows.
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(raw))
        ws = wb["داده"]
        for col, w in zip("ABCDEFGHIJ", (52, 30, 22, 34, 12, 12, 9, 12, 10, 70)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.auto_filter.ref = ws.dimensions
        guide = wb["راهنما"]
        guide.cell(row=guide.max_row, column=1,
                   value="ردیف‌ها همان برنامه‌ی منشور پروژه‌اند؛ ویرایش کنید و همین فایل را برگردانید.")
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def context(self, params, user):
        return {
            "params": params,
            "phases": {charter.key(p.title): p for p in Phase.objects.all()},
            "modules": {charter.key(m.title): m for m in Module.objects.all()},
            "tasks": {charter.key(t.title): t for t in Task.objects.select_related("phase")},
        }

    def check(self, row, ctx):
        v = row.values
        row.key = charter.key(v["عنوان"])
        old = ctx["tasks"].get(row.key)
        start, end = v["شروع"], v["پایان"]
        if old is None and end is None:
            row.status, row.message = "error", "فعالیت تازه است؛ «پایان» لازم است"
            return
        if v["نقطه عطف"] and end:
            start = end
        s, e = start or (old.start_on if old else end), end or (old.end_on if old else None)
        if s and e and e < s:
            row.status, row.message = "error", "پایان قبل از شروع است"
            return
        if v["پیشرفت"] is not None and v["پیشرفت"] > 100:
            row.status, row.message = "error", "پیشرفت حداکثر ۱۰۰ است"
            return

        row.data = {"title": v["عنوان"], "phase": v["فاز"], "module": v["ماژول"], "owner": v["مسئول"], "start_on": start,
                    "end_on": end, "progress": v["پیشرفت"], "status": v["وضعیت"],
                    "is_milestone": v["نقطه عطف"], "description": v["توضیح"]}
        notes = []
        if v["فاز"] and charter.key(v["فاز"]) not in ctx["phases"]:
            notes.append(f"فاز تازه «{v['فاز']}»")
        if v["ماژول"] and charter.key(v["ماژول"]) not in ctx["modules"]:
            notes.append(f"ماژول تازه «{v['ماژول']}»")
        if old is None:
            row.status = "new"
            row.message = "؛ ".join([f"فعالیت تازه {_jdate(s)} تا {_jdate(e)}", *notes])
            return
        changes = self._changes(old, row.data)
        if changes:
            row.status, row.message = "changed", "؛ ".join([*changes, *notes])
        else:
            row.status, row.message = "same", "بدون تغییر"

    @staticmethod
    def _changes(old: Task, d: dict) -> list[str]:
        out = []
        if d["phase"] and (old.phase is None or charter.key(old.phase.title) != charter.key(d["phase"])):
            out.append(f"فاز ← {d['phase']}")
        if d["module"] and (old.module is None or charter.key(old.module.title) != charter.key(d["module"])):
            out.append(f"ماژول ← {d['module']}")
        if d["owner"] and old.owner != d["owner"]:
            out.append(f"مسئول ← {d['owner']}")
        if d["start_on"] and old.start_on != d["start_on"]:
            out.append(f"شروع {_jdate(old.start_on)} ← {_jdate(d['start_on'])}")
        if d["end_on"] and old.end_on != d["end_on"]:
            out.append(f"پایان {_jdate(old.end_on)} ← {_jdate(d['end_on'])}")
        if d["progress"] is not None and old.progress != d["progress"]:
            out.append(f"پیشرفت {old.progress}٪ ← {d['progress']}٪")
        if d["status"] and old.status != d["status"]:
            out.append(f"وضعیت ← {Task.Status(d['status']).label}")
        if d["is_milestone"] is not None and old.is_milestone != d["is_milestone"]:
            out.append("نقطه عطف" if d["is_milestone"] else "دیگر نقطه عطف نیست")
        if d["description"] and old.description.strip() != d["description"]:
            out.append("توضیح تازه")
        return out

    def write(self, rows, ctx, user):
        phases = ctx["phases"]
        phase_order = Phase.objects.count()
        order = (Task.objects.order_by("-order").values_list("order", flat=True).first() or 0) + 1
        now = timezone.now()

        def log(task, subject, body=""):
            LogEntry.objects.create(happened_at=now, kind=LogEntry.Kind.SYSTEM, subject=subject,
                                    body=body, task=task, author=user)

        for r in rows:
            d = r.data
            phase = None
            if d["phase"]:
                phase = phases.get(charter.key(d["phase"]))
                if phase is None:
                    phase = Phase.objects.create(title=d["phase"], order=phase_order,
                                                 color=charter.PHASES[phase_order % len(charter.PHASES)][1])
                    phases[charter.key(d["phase"])] = phase
                    phase_order += 1

            module = None
            if d["module"]:
                mods = ctx["modules"]
                module = mods.get(charter.key(d["module"]))
                if module is None:
                    module = Module.objects.create(title=d["module"], order=len(mods))
                    mods[charter.key(d["module"])] = module

            task = ctx["tasks"].get(r.key)
            is_new = task is None
            if is_new:
                task = Task(title=d["title"], order=order, end_on=d["end_on"],
                            start_on=d["start_on"] or d["end_on"])
                order += 1
            old_status, old_status_label, old_end = task.status, task.get_status_display(), task.end_on

            if phase:
                task.phase = phase
            if module:
                task.module = module
            for f in ("owner", "start_on", "end_on", "progress", "status", "is_milestone", "description"):
                if d[f] not in (None, ""):
                    setattr(task, f, d[f])
            if task.status == Task.Status.DONE:
                task.progress = 100
                task.done_on = task.done_on or date.today()
            else:
                task.done_on = None
            task.save()

            if is_new:
                log(task, f"فعالیت «{task.title}» از اکسل تعریف شد")
                continue
            if task.status != old_status:
                log(task, f"وضعیت «{task.title}»: {old_status_label} ← {task.get_status_display()}")
            if task.end_on != old_end:
                log(task, f"ددلاین «{task.title}» تغییر کرد",
                    f"{old_end.isoformat()} ← {task.end_on.isoformat()}")
        return len(rows)


IMPORTERS = [TeamyarTaskImporter()]
