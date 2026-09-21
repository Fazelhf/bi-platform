"""
Excel export for CRM.

Both CRM «خروجی اکسل» buttons used to build a CSV in the browser out of what
the screen happened to be holding. That failed in two ways at once: a CSV is
not a workbook — Excel opens it with no formats, no column widths and, with
Persian text, a fair chance of mojibake — and the drill-down panel only ever
held the 25 rows of its current page, so a manager exporting «۴۱۲ رکورد» got
25 of them with nothing saying so.

So the file is built here instead, from the database, with the same filters
the screen was showing:

  * every matching record, not one page;
  * real .xlsx with number formats, so a Rial column sums in Excel;
  * a «فیلترها» sheet recording exactly which question the numbers answer,
    because an exported file outlives the screen it came from;
  * right-to-left sheets, frozen headers and an autofilter.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.crm import reports as rpt

# --- shared styling (matches apps.core.export, so the two files look alike) -
HEADER_FILL = PatternFill("solid", fgColor="1C1C1E")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14)
SUB_FONT = Font(size=10, color="7A7A7F")
TOTAL_FILL = PatternFill("solid", fgColor="F1F1EE")
TOTAL_FONT = Font(bold=True, size=11)
BAND_FILL = PatternFill("solid", fgColor="FAFAF8")
THIN = Side(style="thin", color="E2E1DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RIAL_FMT = "#,##0"
COUNT_FMT = "#,##0"
PCT_FMT = '0.0"٪"'
DAYS_FMT = "0.0"

FORMATS = {"rial": RIAL_FMT, "count": COUNT_FMT, "pct": PCT_FMT, "days": DAYS_FMT}


def _sheet(wb, title, first=False):
    ws = wb.active if first else wb.create_sheet()
    # Excel forbids : \ / ? * [ ] in a sheet name and truncates past 31 chars.
    ws.title = "".join(c for c in title if c not in ':\\/?*[]')[:31]
    ws.sheet_view.rightToLeft = True
    return ws


def _write_table(ws, columns, rows, start_row=1, total_row=None):
    """
    Write `rows` under a styled header.

    `columns` is a list of {key, label, fmt}; `rows` are dicts. A `total_row`
    dict is written in bold under the data — the same totals the screen shows
    in its footer, so the file and the screen cannot disagree.
    """
    for c, col in enumerate(columns, 1):
        cell = ws.cell(row=start_row, column=c, value=col["label"])
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    r = start_row
    for i, row in enumerate(rows):
        r = start_row + 1 + i
        for c, col in enumerate(columns, 1):
            value = row.get(col["key"])
            if value == "":
                value = None
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            fmt = FORMATS.get(col.get("fmt", "text"))
            if fmt and isinstance(value, (int, float)):
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="left")
            if i % 2:
                cell.fill = BAND_FILL

    if total_row is not None and rows:
        r += 1
        for c, col in enumerate(columns, 1):
            value = total_row.get(col["key"], "جمع" if c == 1 else None)
            cell = ws.cell(row=r, column=c, value=value)
            cell.fill, cell.font, cell.border = TOTAL_FILL, TOTAL_FONT, BORDER
            fmt = FORMATS.get(col.get("fmt", "text"))
            if fmt and isinstance(value, (int, float)):
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="left")

    # Column widths from the longest cell — Persian text is wide, and a column
    # of ### is what an unsized Rial column looks like when it is opened.
    for c, col in enumerate(columns, 1):
        longest = len(str(col["label"]))
        for row in rows[:400]:
            v = row.get(col["key"])
            if v is None:
                continue
            text = f"{v:,.0f}" if isinstance(v, (int, float)) else str(v)
            longest = max(longest, len(text))
        ws.column_dimensions[get_column_letter(c)].width = min(46, max(11, longest + 4))

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        last = get_column_letter(len(columns))
        ws.auto_filter.ref = f"A{start_row}:{last}{start_row + len(rows)}"
    return r


def _title_block(ws, title, subtitle="", width=6):
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    if subtitle:
        ws.cell(row=2, column=1, value=subtitle).font = SUB_FONT
    ws.row_dimensions[1].height = 22
    return 4 if subtitle else 3


def _filters_sheet(wb, context):
    """
    The question the numbers answer, written down.

    A file named «گزارش کلی فروش.xlsx» in someone's downloads folder says
    nothing about which months, whose book or which filters produced it. This
    sheet does, so the file can be forwarded without the context being lost.
    """
    ws = _sheet(wb, "فیلترها")
    ws.cell(row=1, column=1, value="فیلترهای اعمال‌شده").font = TITLE_FONT
    _write_table(
        ws,
        [{"key": "name", "label": "عنوان", "fmt": "text"},
         {"key": "value", "label": "مقدار", "fmt": "text"}],
        [{"name": k, "value": v} for k, v in context if v not in (None, "")],
        start_row=3,
    )
    return ws


# --------------------------------------------------------------------------
# گزارش‌ها
# --------------------------------------------------------------------------
def report_workbook(data: dict, context: list[tuple[str, str]]):
    """One report, as the screen shows it: its columns, its rows, its totals."""
    # The report's own column spec (`k`/`f`) is what the screen's table uses;
    # translate it once here rather than keeping a second copy of the layout.
    axis_label = (data.get("axis_label") or "عنوان").replace("بر محور ", "")
    columns = [{"key": "label", "label": axis_label, "fmt": "text"}]
    columns += [
        {"key": c["k"], "label": c["label"], "fmt": c.get("f", "text")}
        for c in rpt.REPORT_COLUMNS.get(data["key"], [])
    ]

    wb = Workbook()
    ws = _sheet(wb, data.get("title", "گزارش"), first=True)
    window = dict(context).get("بازه زمانی", "")
    start = _title_block(ws, data.get("title", "گزارش"), window, len(columns))

    totals = dict(data.get("totals") or {})
    totals["label"] = "جمع کل"
    _write_table(ws, columns, data.get("rows", []), start_row=start, total_row=totals)

    _filters_sheet(wb, context)
    return wb


# --------------------------------------------------------------------------
# ریز رکوردها (drill-down)
# --------------------------------------------------------------------------
#: Columns per record kind. Wider than what the drawer shows on screen — an
#: export is what somebody works in afterwards, so it carries the contact
#: details and ids the table has no room for.
DRILL_COLUMNS = {
    "deals": [
        {"key": "code", "label": "کد", "fmt": "text"},
        {"key": "title", "label": "عنوان معامله", "fmt": "text"},
        {"key": "customer_name", "label": "مشتری", "fmt": "text"},
        {"key": "owner_name", "label": "کارشناس", "fmt": "text"},
        {"key": "province_name", "label": "استان", "fmt": "text"},
        {"key": "stage_name", "label": "مرحله", "fmt": "text"},
        {"key": "status_display", "label": "وضعیت", "fmt": "text"},
        {"key": "reason_name", "label": "دلیل عدم موفقیت", "fmt": "text"},
        {"key": "amount_rial", "label": "مبلغ (ریال)", "fmt": "rial"},
        {"key": "cost_rial", "label": "هزینه (ریال)", "fmt": "rial"},
        {"key": "profit_rial", "label": "سود (ریال)", "fmt": "rial"},
        {"key": "margin_pct", "label": "حاشیه سود", "fmt": "pct"},
        {"key": "opened_jalali", "label": "تاریخ ایجاد", "fmt": "text"},
        {"key": "closed_jalali", "label": "تاریخ بسته‌شدن", "fmt": "text"},
    ],
    "customers": [
        {"key": "code", "label": "کد", "fmt": "text"},
        {"key": "name_fa", "label": "نام مشتری", "fmt": "text"},
        {"key": "group_name", "label": "گروه", "fmt": "text"},
        {"key": "province_name", "label": "استان", "fmt": "text"},
        {"key": "city", "label": "شهر", "fmt": "text"},
        {"key": "owner_name", "label": "کارشناس", "fmt": "text"},
        {"key": "source_name", "label": "منبع سرنخ", "fmt": "text"},
        {"key": "status_display", "label": "وضعیت", "fmt": "text"},
        {"key": "contact_name", "label": "نام رابط", "fmt": "text"},
        {"key": "phone", "label": "تلفن", "fmt": "text"},
        {"key": "mobile", "label": "موبایل", "fmt": "text"},
        {"key": "first_contact_jalali", "label": "اولین تماس", "fmt": "text"},
        {"key": "first_won_jalali", "label": "اولین خرید", "fmt": "text"},
    ],
    "activities": [
        {"key": "kind_display", "label": "نوع", "fmt": "text"},
        {"key": "customer_name", "label": "مشتری", "fmt": "text"},
        {"key": "owner_name", "label": "کارشناس", "fmt": "text"},
        {"key": "result_display", "label": "نتیجه", "fmt": "text"},
        {"key": "duration_min", "label": "مدت (دقیقه)", "fmt": "count"},
        {"key": "at_jalali", "label": "تاریخ", "fmt": "text"},
        {"key": "subject", "label": "موضوع", "fmt": "text"},
        {"key": "note", "label": "توضیح", "fmt": "text"},
    ],
    "feedback": [
        {"key": "customer_name", "label": "مشتری", "fmt": "text"},
        {"key": "employee_name", "label": "کارشناس", "fmt": "text"},
        {"key": "score", "label": "امتیاز", "fmt": "count"},
        {"key": "at_jalali", "label": "تاریخ", "fmt": "text"},
        {"key": "note", "label": "توضیح", "fmt": "text"},
    ],
    "invoices": [
        {"key": "number", "label": "شماره فاکتور", "fmt": "text"},
        {"key": "kind_display", "label": "نوع", "fmt": "text"},
        {"key": "issued_jalali", "label": "تاریخ صدور", "fmt": "text"},
        {"key": "customer_name", "label": "مشتری", "fmt": "text"},
        {"key": "owner_name", "label": "بازاریاب", "fmt": "text"},
        {"key": "amount_rial", "label": "مبلغ خالص (ریال)", "fmt": "rial"},
        {"key": "vat_rial", "label": "مالیات و عوارض (ریال)", "fmt": "rial"},
        {"key": "unsettled_rial", "label": "تسویه‌نشده (ریال)", "fmt": "rial"},
        {"key": "deal_title", "label": "معامله", "fmt": "text"},
    ],
}

KIND_TITLE = {
    "deals": "معاملات", "customers": "مشتریان", "activities": "فعالیت‌ها",
    "feedback": "بازخوردها", "invoices": "فاکتورها",
}


def _numeric(value):
    """Serialized Decimals arrive as strings; Excel must get numbers."""
    if isinstance(value, str):
        try:
            return float(value)
        except ValueError:
            return value
    return value


def drill_workbook(kind: str, rows: list[dict], title: str,
                   context: list[tuple[str, str]], summary: dict | None = None):
    columns = DRILL_COLUMNS.get(kind, DRILL_COLUMNS["deals"])
    numeric_keys = {c["key"] for c in columns if c.get("fmt") in ("rial", "count", "pct")}
    clean_rows = [
        {k: (_numeric(v) if k in numeric_keys else v) for k, v in row.items()}
        for row in rows
    ]

    wb = Workbook()
    ws = _sheet(wb, KIND_TITLE.get(kind, "رکوردها"), first=True)
    subtitle = f"{len(clean_rows):,} رکورد"
    window = dict(context).get("بازه زمانی", "")
    if window:
        subtitle += f" · {window}"
    start = _title_block(ws, title, subtitle, len(columns))

    # A totals line only where adding up means something.
    total_row = None
    if kind in ("deals", "invoices"):
        total_row = {"label": "جمع کل"}
        first_key = columns[0]["key"]
        total_row[first_key] = "جمع کل"
        for col in columns:
            if col["fmt"] == "rial":
                total_row[col["key"]] = sum(
                    v for r in clean_rows
                    if isinstance(v := r.get(col["key"]), (int, float))
                )
        if summary and summary.get("margin_pct") is not None:
            total_row["margin_pct"] = summary["margin_pct"]

    _write_table(ws, columns, clean_rows, start_row=start, total_row=total_row)

    if summary:
        _summary_sheet(wb, kind, summary)
    _filters_sheet(wb, context)
    return wb


SUMMARY_LABELS = {
    "count": "تعداد", "amount": "مبلغ (ریال)", "cost": "هزینه (ریال)",
    "profit": "سود (ریال)", "margin_pct": "حاشیه سود (٪)", "success": "موفق",
    "success_rate": "نرخ موفقیت (٪)", "customers": "مشتریان", "minutes": "دقیقه",
}


def _summary_sheet(wb, kind, summary):
    """The reconciliation strip from the drawer header, kept with the rows."""
    ws = _sheet(wb, "خلاصه")
    ws.cell(row=1, column=1, value="خلاصه").font = TITLE_FONT
    _write_table(
        ws,
        [{"key": "name", "label": "عنوان", "fmt": "text"},
         {"key": "value", "label": "مقدار", "fmt": "rial"}],
        [{"name": SUMMARY_LABELS.get(k, k), "value": _numeric(v)}
         for k, v in summary.items() if v is not None],
        start_row=3,
    )
    return ws


def to_stream(wb):
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
