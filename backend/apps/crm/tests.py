"""
CRM tests, centred on the thing that actually went wrong in production.

The company's real customer file and a fabricated showroom live in the same
tables, told apart by one column. Two failure modes matter, and only one of
them is visible on screen:

* **A leak.** The showroom shows a real customer, or the real file shows an
  invented one. Loud, and someone reports it.
* **A silent skip.** The importer decides the real data is already loaded when
  it is not, and never runs. Nothing errors, the pages fill with the demo set
  wearing the «واقعی» label, and it reads as working. This is what happened on
  the server: migration 0004 defaults `dataset` to "real", so the generated
  demo set that predated the column arrived labelled real, and a guard that
  asked «are there real rows?» answered yes.

So the guard is tested by its real question — «did *this command* write
anything?» — and the answer keys on the «didar-» codes only the importer
mints.
"""
from collections import defaultdict
from datetime import timedelta
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.db import transaction
from django.db.models import Sum
from django.db.utils import IntegrityError
from django.utils import timezone
from rest_framework.test import APITestCase

from apps.crm import matching, merge as crm_merge, reports as rpt
from apps.crm.management.commands.import_arpa_invoices import (
    Command as InvoiceCommand, dec as inv_dec,
)
from apps.crm.invoice_link import WINDOW_DAYS, link_invoices
from apps.crm.matching import CustomerIndex, Method
from apps.crm.management.commands.import_didar_crm import Command as ImportCommand, fit
from apps.crm.models import (
    Activity, Customer, CustomerExternalRef, CustomerMatchCandidate, Dataset,
    Deal, DealItem, ExternalSource, PipelineStage, Product, SalesInvoice,
    SalesInvoiceItem,
    DismissedParty,
)
from apps.sales.models import DimEmployee, EmployeeChannel, SalesChannel


def _user(username, role, department="", dataset="real"):
    User = get_user_model()
    return User.objects.create_user(
        username=username, password="pw12345!", role=role,
        department=department, crm_dataset=dataset,
    )


class DatasetTestCase(APITestCase):
    """One customer and one deal on each side, deliberately alike."""

    def setUp(self):
        now = timezone.now()
        self.real_customer = Customer.objects.create(
            code="didar-pe-1759", name_fa="بانک ملی خراسان رضوی",
            dataset=Dataset.REAL, first_contact_at=now - timedelta(days=90),
        )
        self.demo_customer = Customer.objects.create(
            code="cust-2082", name_fa="مشتری نمایشی",
            dataset=Dataset.DEMO, first_contact_at=now - timedelta(days=90),
        )
        for customer, ds, code in (
            (self.real_customer, Dataset.REAL, "didar-d-125"),
            (self.demo_customer, Dataset.DEMO, "deal-9001"),
        ):
            Deal.objects.create(
                code=code, title=f"معامله {code}", customer=customer,
                dataset=ds, opened_at=now - timedelta(days=30),
            )

        self.ceo = _user("ceo", "executive")
        self.rep = _user("rep", "manager", "sales_team")


class DatasetIsolationTests(DatasetTestCase):
    """
    The demo showroom was removed. Old databases still hold `dataset` values
    and old accounts still hold a `crm_dataset` preference; neither may bring
    a fabricated row back onto a screen.
    """

    def test_an_old_demo_preference_still_reads_the_real_file(self):
        self.ceo.crm_dataset = "demo"
        self.ceo.save(update_fields=["crm_dataset"])
        self.client.force_authenticate(self.ceo)

        names = [r["name_fa"] for r in self.client.get("/api/crm/customers/").data["results"]]
        self.assertEqual(names, ["بانک ملی خراسان رضوی"])

    def test_new_rows_are_always_real(self):
        self.rep.crm_dataset = "demo"
        self.rep.save(update_fields=["crm_dataset"])
        self.client.force_authenticate(self.rep)

        res = self.client.post("/api/crm/customers/", {"name_fa": "مشتری تازه"})
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(
            Customer.objects.get(name_fa="مشتری تازه").dataset, Dataset.REAL
        )

    def test_the_switch_endpoint_is_gone(self):
        # Checked by URL resolution rather than by requesting it: an unmatched
        # URL renders Django's debug 404 page, whose template-context copy
        # breaks under Python 3.14 on this Django version — an unrelated error
        # that would mask what this test is about.
        from django.urls import Resolver404, resolve
        with self.assertRaises(Resolver404):
            resolve("/api/crm/dataset/")

        self.client.force_authenticate(self.rep)
        self.assertNotIn("dataset", self.client.get("/api/crm/me/").data)

    def test_funnel_stages_do_not_mix_the_two_vocabularies(self):
        """
        Both datasets name their stages «ارتباط مشتری». Unfiltered, the funnel
        drew every stage twice — the duplicate rows reported from the CRM.
        """
        for ds in (Dataset.REAL, Dataset.DEMO):
            PipelineStage.objects.create(
                code=f"st-contact-{ds}", name_fa="ارتباط مشتری",
                order=1, dataset=ds,
            )
        rows = rpt.report_funnel(rpt.Filters.from_query({}, "real"))["rows"]
        self.assertEqual([r["label"] for r in rows], ["ارتباط مشتری"])


class ImportGuardTests(DatasetTestCase):
    """
    `--if-empty` is what lets deploy.sh call the importer on every deploy.
    Wrong in either direction it is dangerous: too eager and it wipes the
    sales team's work, too shy and the real data never arrives.
    """

    def test_guard_sees_imported_rows(self):
        self.assertTrue(ImportCommand()._already_imported())

    def test_guard_ignores_rows_it_did_not_write(self):
        """
        The production bug. A demo customer tagged «real» by the migration
        default is not evidence that the import has run — asking «are there
        real rows?» is what made the importer skip itself in silence.
        """
        Customer.objects.filter(code__startswith="didar-").delete()
        Customer.objects.filter(code="cust-2082").update(dataset=Dataset.REAL)

        self.assertTrue(Customer.objects.filter(dataset=Dataset.REAL).exists())
        self.assertFalse(ImportCommand()._already_imported())

    def test_first_import_clears_the_mislabelled_seed(self):
        Customer.objects.filter(code__startswith="didar-").delete()
        Customer.objects.filter(code="cust-2082").update(dataset=Dataset.REAL)

        cmd = ImportCommand()
        cmd.stdout = type("_S", (), {"write": lambda self, *a, **k: None})()
        cmd.style = type("_T", (), {"WARNING": staticmethod(lambda s: s)})()
        cmd._clear_mislabelled()

        self.assertFalse(Customer.objects.filter(dataset=Dataset.REAL).exists())


class ColumnWidthTests(APITestCase):
    """
    The class of bug that only production could see.

    SQLite ignores `VARCHAR(n)` completely, so an over-long value writes clean
    in development and the import looks correct. PostgreSQL does not, and the
    server stopped on «value too long for type character varying(40)» — after
    the import had already been declared working here.

    `fit` asks the model for the width instead of the call site guessing it,
    which is what these lock in: the two real offenders, and the property that
    a clamp is derived rather than typed.
    """

    def test_clamps_to_the_field_s_own_limit(self):
        long_phone = "02196052633 مدیریت اداره خرید و فروش اقلام عمده شرکت"
        self.assertGreater(len(long_phone), 40)
        self.assertEqual(len(fit(Customer, "phone", long_phone)), 40)

        from apps.crm.models import Activity
        self.assertEqual(len(fit(Activity, "note", "ن" * 900)), 500)

    def test_uses_the_model_rather_than_a_written_number(self):
        """`note` was hand-sliced to 2000 against a 500-wide column; the point
        of deriving the limit is that such a slice cannot drift again."""
        limit = Customer._meta.get_field("name_fa").max_length
        self.assertEqual(len(fit(Customer, "name_fa", "ا" * (limit + 50))), limit)

    def test_none_and_numbers_survive(self):
        self.assertEqual(fit(Customer, "phone", None), "")
        self.assertEqual(fit(Customer, "phone", 2196052633), "2196052633")


class IdentityLayerTests(DatasetTestCase):
    """
    A customer has to be reachable by *each* source system's own id.

    Until accounting arrived there was one source, and `Customer.code` could
    carry both the row's name and its origin («didar-co-46206»). A second
    source breaks that: آرپا's کد طرف حساب has nowhere to live, so the same
    company is imported again under a new code and every per-customer figure
    silently splits in two.
    """

    def test_the_importer_files_the_didar_id_where_a_second_source_can_look(self):
        """
        The id has to end up somewhere queryable — not only inside `code`.
        Derived from `code` so it matches, row for row, what migration 0006
        backfilled onto the customers already in the production database.
        """
        ImportCommand._link(self.real_customer)

        ref = self.real_customer.external_refs.get(source=ExternalSource.DIDAR)
        self.assertEqual(ref.external_id, "pe-1759")
        self.assertEqual(ref.external_name, self.real_customer.name_fa)

    def test_link_is_idempotent_and_survives_a_rename(self):
        """
        The importer runs again whenever a fresher export arrives. A second
        run must update the row it already wrote, not add a second id — and a
        company renamed in دیدار must keep its id rather than fork.
        """
        ImportCommand._link(self.real_customer)
        self.real_customer.name_fa = "نام تازه"
        self.real_customer.save(update_fields=["name_fa", "updated_at"])
        ImportCommand._link(self.real_customer)

        self.assertEqual(self.real_customer.external_refs.count(), 1)
        ref = self.real_customer.external_refs.get(source=ExternalSource.DIDAR)
        self.assertEqual(ref.external_name, "نام تازه")

    def test_one_source_id_cannot_name_two_customers(self):
        """
        The constraint that makes a merge decision final. Without it a later
        import pass could quietly attach آرپا's کد ۴۰۰۷۲۱ to a different
        customer than the reviewer chose, and nothing would complain.
        """
        other = Customer.objects.create(
            code="didar-co-duplicate", name_fa="شرکت دوم",
            first_contact_at=timezone.now(),
        )
        CustomerExternalRef.objects.create(
            source=ExternalSource.ARPA, external_id="400721",
            customer=self.real_customer,
        )
        with self.assertRaises(IntegrityError), transaction.atomic():
            CustomerExternalRef.objects.create(
                source=ExternalSource.ARPA, external_id="400721", customer=other,
            )

    def test_the_same_id_may_repeat_across_sources(self):
        """دیدار and آرپا number their parties independently; «۱۰۰» in one is
        unrelated to «۱۰۰» in the other, so the key is the pair."""
        CustomerExternalRef.objects.create(
            source=ExternalSource.ARPA, external_id="100",
            customer=self.real_customer,
        )
        CustomerExternalRef.objects.create(
            source=ExternalSource.DIDAR, external_id="100",
            customer=self.real_customer,
        )
        self.assertEqual(
            self.real_customer.external_refs.filter(external_id="100").count(), 2
        )


class InvoiceSeparationTests(DatasetTestCase):
    """
    Invoice money and deal money must not be the same number.

    `Deal.amount_rial` comes from دیدار and دیدار's own reports are built on
    it — the importer checks itself against those totals. An invoice is what
    accounting actually billed. Keeping them apart is what lets «چقدر از کاریز
    فاکتور شد» be a question with an answer instead of a tautology.
    """

    def setUp(self):
        super().setUp()
        self.customer = self.real_customer

    def _invoice(self, number, kind, amount):
        return SalesInvoice.objects.create(
            code=f"arpa-inv-{number}", number=number, kind=kind,
            customer=self.customer, issued_at=timezone.now().date(),
            amount_rial=amount,
        )

    def test_invoicing_a_customer_leaves_the_deal_untouched(self):
        deal = Deal.objects.filter(customer=self.customer).first()
        if deal is None:
            self.skipTest("fixture has no deal for this customer")
        before = deal.amount_rial

        self._invoice("1001", SalesInvoice.Kind.SALE, 5_000_000)

        deal.refresh_from_db()
        self.assertEqual(deal.amount_rial, before)

    def test_a_return_sums_against_the_sale_it_reverses(self):
        """
        آرپا writes مرجوعی with a negative مبلغ فروش. Storing the sign as
        given means a plain SUM over a period is already net of returns —
        the alternative is every report remembering to subtract, and one of
        them eventually not doing it.
        """
        self._invoice("1002", SalesInvoice.Kind.SALE, 10_000_000)
        self._invoice("1003", SalesInvoice.Kind.RETURN, -4_000_000)

        total = SalesInvoice.objects.filter(customer=self.customer).aggregate(
            t=Sum("amount_rial")
        )["t"]
        self.assertEqual(total, 6_000_000)

    def test_a_line_survives_a_product_the_catalogue_does_not_have(self):
        """
        The آرپا catalogue and the دیدار one only partly line up. A sale of an
        unmappable product is still revenue, so the line keeps the source's
        own code and name and refuses to be dropped.
        """
        invoice = self._invoice("1004", SalesInvoice.Kind.SALE, 3_000_000)
        item = SalesInvoiceItem.objects.create(
            invoice=invoice, product=None,
            product_code="51300", product_name="49 - 16 - ساده",
            quantity=7200, unit_price_rial=115_500, amount_rial=831_600_000,
        )
        self.assertIsNone(item.product)
        self.assertEqual(invoice.items.count(), 1)


class NormalisationTests(APITestCase):
    """
    The folding is what makes the exact-name tier work at all.

    آرپا writes legacy Persian — Arabic ي (U+064A) and ك (U+0643) — and دیدار
    writes ی and ک. Compared as typed, the same company name from the two
    systems is two different strings, and every tier below «exact name» would
    have to carry the whole customer file.
    """

    def test_the_two_systems_spell_the_same_name_identically_once_folded(self):
        arpa = "شرکت خدمات بهداشتي آواي سلامت پارس"   # ي, ك
        didar = "شرکت خدمات بهداشتی آوای سلامت پارس"   # ی, ک
        self.assertNotEqual(arpa, didar)
        self.assertEqual(matching.name_key(arpa), matching.name_key(didar))

    def test_descriptive_words_do_not_identify_a_company(self):
        self.assertEqual(
            matching.name_key("شرکت بازرگانی ایران ارقام سهامی خاص"),
            matching.name_key("ایران ارقام"),
        )
        self.assertEqual(matching.name_key("آقای انجوی"), matching.name_key("انجوی"))

    def test_zwnj_and_persian_digits_fold(self):
        self.assertEqual(matching.name_key("نمابر\u200cمهر"), "نمابر مهر")
        self.assertEqual(matching.phone_key("۰۲۱۴۴۶۴۱۳۳۰"), "44641330")

    def test_phone_key_sees_through_dialling_prefixes(self):
        for written in ("02144641330", "44641330", "021-4464-1330", "+982144641330"):
            self.assertEqual(matching.phone_key(written), "44641330", written)

    def test_a_short_number_is_not_a_key(self):
        """Four digits shared by hundreds of rows would match everything."""
        self.assertEqual(matching.phone_key("4464"), "")
        self.assertEqual(matching.id_key("123"), "")


class PlaceGuardTests(APITestCase):
    """
    The false positive that set the threshold policy.

    «بانک کشاورزی ایلام» and «بانک کشاورزی گیلان» are 86% alike by character
    overlap — above any fuzzy threshold worth having — and are branches in
    provinces 900km apart. A bank with a branch in every province makes string
    similarity actively dangerous, so a disagreement about *place* vetoes the
    similarity rather than being outweighed by it.
    """

    def test_two_provinces_are_not_one_customer(self):
        self.assertGreater(
            matching.similarity("بانک کشاورزی ایلام", "بانک کشاورزی گیلان"),
            matching.FUZZY_FLOOR,
        )
        self.assertTrue(
            matching.place_conflict("بانک کشاورزی ایلام", "بانک کشاورزی گیلان")
        )

    def test_silence_about_place_is_not_disagreement(self):
        """«بانک سینا» and «بانک سینا تهران» may well be one account."""
        self.assertFalse(matching.place_conflict("بانک سینا", "بانک سینا تهران"))

    def test_the_same_place_written_twice_is_agreement(self):
        self.assertFalse(
            matching.place_conflict("بانک کشاورزی البرز-کرج", "بانک کشاورزی البرز")
        )


class MatchLadderTests(APITestCase):
    """
    Which rung answers decides whether a customer is merged or merely
    suggested. A wrong merge fuses two companies' order history and nothing
    on screen looks wrong afterwards, so only the top rungs may write.
    """

    def setUp(self):
        self.index = CustomerIndex()
        self.index.add(1, name="ایران ارقام", nids=("10100905654",), phones=("02142719000",))
        self.index.add(2, name="بانک کشاورزی ایلام", phones=("08433330000",))
        self.index.add(3, name="پلی کلینیک سوم خرداد", phones=("06153500000",))
        self.index.add(4, name="شبکه بهداشت خرمشهر", phones=("06153500000",))

    def _find(self, **kw):
        kw.setdefault("source", "arpa")
        kw.setdefault("external_id", "9999")
        kw.setdefault("nids", ())
        kw.setdefault("phones", ())
        return self.index.find(**kw)

    def test_a_previous_run_s_link_wins_over_everything(self):
        """Re-running the import must land on the row the reviewer chose, not
        re-derive an answer that may have drifted."""
        self.index.add_ref("arpa", "400721", 2)
        match = self._find(external_id="400721", name="ایران ارقام")
        self.assertEqual(match.customer_id, 2)
        self.assertTrue(match.is_auto)

    def test_the_same_name_spelled_the_other_way_merges_unattended(self):
        match = self._find(name="شرکت ايران ارقام")
        self.assertEqual((match.method, match.customer_id), (Method.NAME, 1))
        self.assertTrue(match.is_auto)

    def test_a_shared_switchboard_is_only_a_suggestion(self):
        """
        The measured failure: a clinic and a health authority in خرمشهر answer
        one number. Phone finds a pair, and the pair must not be written.
        """
        match = self._find(name="درمانگاه تازه", phones=("06153500000",))
        self.assertEqual(match.method, Method.PHONE)
        self.assertTrue(match.found)
        self.assertFalse(match.is_auto)

    def test_a_sister_province_branch_is_not_matched_at_all(self):
        match = self._find(name="بانک کشاورزی گیلان")
        self.assertEqual(match.method, Method.NONE)
        self.assertIsNone(match.customer_id)

    def test_a_name_shared_by_two_customers_is_flagged_not_merged(self):
        """
        A duplicate already inside the CRM. Merging into either half would
        bury it; the reviewer has to settle that first.
        """
        self.index.add(5, name="ایران ارقام")
        match = self._find(name="ایران ارقام")
        self.assertEqual(match.method, Method.AMBIGUOUS)
        self.assertFalse(match.is_auto)

    def test_a_national_id_suggests_but_never_writes(self):
        """
        It names a legal entity, and the customers here are its branches —
        every one of them carrying the head office's number. Trusting it
        merged seventeen پست بانک branches into a single «دولتی پست بانک».
        دیدار carries no national id at all, so the tier could never link the
        two systems anyway: it can only join آرپا parties to each other, which
        is precisely the damage.
        """
        match = self._find(name="نام کاملا متفاوت", nids=("10100905654",))
        self.assertEqual((match.method, match.customer_id), (Method.NATIONAL_ID, 1))
        self.assertFalse(match.is_auto)

    def test_an_unknown_party_is_left_unmatched(self):
        match = self._find(name="شرکت تازه وارد", phones=("02100000000",))
        self.assertEqual(match.method, Method.NONE)


class BranchNationalIdTests(APITestCase):
    """
    The bug a second import run produced, and why the top rung is not enough.

    A شناسه ملی identifies a legal entity. The customers here are largely that
    entity's branches, and every branch of بانک صادرات carries the head
    office's number. Run one created the branches from آرپا with their ids;
    run two matched them to each other by id and merged «بانک صادرات
    کرمانشاه» into «بانک صادرات آذربایجان» — a wrong merge, arriving through
    the rung the ladder trusts most, and invisible on any screen afterwards.

    Two branches may well be one account. That is a decision about how the
    company sells, and an identifier cannot make it.
    """

    def setUp(self):
        self.index = CustomerIndex()
        self.SHARED = "10861904730"
        self.index.add(1, name="بانک صادرات آذربایجان", nids=(self.SHARED,))

    def _find(self, name, nids):
        return self.index.find(
            source="arpa", external_id="1", name=name, nids=nids, phones=(),
        )

    def test_a_second_branch_is_not_merged_into_the_first(self):
        match = self._find("بانک صادرات کرمانشاه", (self.SHARED,))
        self.assertEqual(match.method, Method.BRANCH)
        self.assertFalse(match.is_auto)
        self.assertTrue(match.found, "the pair is still worth showing a reviewer")

    def test_no_id_match_writes_itself_whatever_the_names_say(self):
        """
        The place guard narrowed this and did not close it: a head-office row
        names no place, so it never disagrees, and «پست بانک کرمان» merged
        into «دولتی پست بانک» through a rung that saw no conflict at all.
        """
        for name in ("صادرات آذربایجان", "بانک صادرات", "بانک صادرات کرمانشاه"):
            self.assertFalse(self._find(name, (self.SHARED,)).is_auto, name)

    def test_the_ladder_does_not_reach_the_id_when_a_link_exists(self):
        """A reviewer's decision outranks every heuristic, including this."""
        self.index.add(2, name="بانک صادرات کرمانشاه")
        self.index.add_ref("arpa", "77", 2)
        match = self.index.find(
            source="arpa", external_id="77", name="بانک صادرات کرمانشاه",
            nids=(self.SHARED,), phones=(),
        )
        self.assertEqual((match.method, match.customer_id), (Method.EXISTING, 2))


class InvoiceImportUnitTests(APITestCase):
    """
    The three shapes in the آرپا export that decide whether the totals are
    right, each of which reads as a plausible number when handled wrong.
    """

    def test_the_invoice_key_survives_a_year_rollover(self):
        """
        «شماره برگه» restarts each year and «شماره ثابت سند» is not unique
        either — 868 distinct values for 1,098 rows. Only (kind, number, date)
        is unique on all of them, so the code carries the date.
        """
        from datetime import date
        code_1404 = InvoiceCommand._code("sale", "241", date(2025, 6, 15))
        code_1405 = InvoiceCommand._code("sale", "241", date(2026, 6, 15))
        self.assertNotEqual(code_1404, code_1405)
        self.assertLessEqual(len(code_1404), 50)

    def test_a_totals_row_is_not_an_invoice(self):
        """
        Every file carries one row with no نوع برگه whose amount equals the
        sum of all the others. Counted, it doubles revenue exactly — which
        looks like a very good year rather than like a bug.
        """
        rows = [
            {"نوع برگه": "فاکتور فروش", "مبلغ فروش": 100},
            {"نوع برگه": "فاکتور فروش", "مبلغ فروش": 250},
            {"نوع برگه": None, "مبلغ فروش": 350},
        ]
        real = [r for r in rows if matching.fold(r.get("نوع برگه"))]
        self.assertEqual(len(real), 2)
        self.assertEqual(sum(inv_dec(r["مبلغ فروش"]) for r in real), 350)

    def test_a_return_keeps_the_sign_the_source_gave_it(self):
        self.assertEqual(inv_dec("-1,446,702,200"), Decimal("-1446702200"))
        self.assertEqual(inv_dec("948,000,000"), Decimal("948000000"))
        self.assertEqual(inv_dec(None), Decimal(0))
        self.assertEqual(inv_dec("چیزی نیست"), Decimal(0))


class InvoiceRepMappingTests(APITestCase):
    """
    The same colleague, spelled two ways in two systems. Left unmatched, a
    third of the invoices lose their owner; matched too eagerly, two people
    become one and both their numbers are wrong.
    """

    def setUp(self):
        self.cmd = InvoiceCommand()
        self.cmd.unknown_reps = defaultdict(int)
        self.cmd.employees = {
            matching.name_key(n): n for n in (
                "هانیه منزه", "حامد بهشتی", "بهلول", "مهدیس مومنی",
                "سارا مسگرچیان", "پیام بوساک", "صبا موسوی",
            )
        }

    def test_an_extra_name_part_is_the_same_person(self):
        """«هانیه خواجه منزه» in آرپا is «هانیه منزه» in the CRM."""
        self.assertEqual(self.cmd._rep("هانیه خواجه منزه"), "هانیه منزه")
        self.assertEqual(self.cmd._rep("حامد بهشتی زواره"), "حامد بهشتی")

    def test_a_surname_only_colleague_still_matches(self):
        self.assertEqual(self.cmd._rep("عاطفه بهلول"), "بهلول")

    def test_hamza_is_not_a_different_person(self):
        self.assertEqual(self.cmd._rep("مهدیس موءمنی"), "مهدیس مومنی")

    def test_a_different_surname_is_never_inferred(self):
        """
        A rule loose enough to pair «سارا مسگرقمی» with «سارا مسگرچیان» would
        also fuse two colleagues who merely share a given name, and a fused
        colleague halves a real person's numbers wherever they are reported.
        So the rule refuses, and the invoice is left unowned — visibly.
        """
        self.assertIsNone(self.cmd._rep("سارا مسگری"))
        self.assertIsNone(self.cmd._rep("کاوه بهشتی"))
        self.assertEqual(set(self.cmd.unknown_reps), {"سارا مسگری", "کاوه بهشتی"})

    def test_a_confirmed_pair_is_written_down_not_derived(self):
        """
        «محسن بوساک» is «پیام بوساک» and «سارا مسگرقمی» is «سارا مسگرچیان» —
        confirmed by the sales manager, and underivable, because it is the
        name that differs rather than its spelling. Listing them by hand is
        what keeps the general rule tight.
        """
        self.assertEqual(self.cmd._rep("محسن بوساک"), "پیام بوساک")
        self.assertEqual(self.cmd._rep("سارا مسگرقمی"), "سارا مسگرچیان")
        self.assertFalse(self.cmd.unknown_reps)

    def test_the_placeholder_is_not_a_colleague(self):
        """425 of 1,098 invoices carry «بازاریاب بدون پورسانت»."""
        self.assertIsNone(self.cmd._rep("بازاریاب بدون پورسانت"))
        self.assertNotIn("بازاریاب بدون پورسانت", self.cmd.unknown_reps)


class MergeReviewTests(APITestCase):
    """
    What a reviewer's decision does.

    Both answers must end with the آرپا party carrying an external ref. An
    unresolved party is invisible to the invoice import — 260 invoices worth
    543bn Rial are waiting on this queue — so «reject» meaning «drop it»
    would quietly make that money unimportable for good.
    """

    def setUp(self):
        now = timezone.now()
        self.customer = Customer.objects.create(
            code="didar-co-500", name_fa="بانک سپه سبلان",
            dataset=Dataset.REAL, first_contact_at=now,
        )
        self.payload = {
            "کد": "400721", "نام": "بانک سپه شعبه سبلان شمالی",
            "شماره تلفن": "02188811961", "شهر": "تهران",
            "کد اقتصادی": "411431776417", "شناسه ملی": "10861904730",
            "نام گروه": "نمابر مهر بانکها", "نوع": "حقوقی",
            "شرایط تسویه پیش فرض": "30روزه",
        }
        self.candidate = CustomerMatchCandidate.objects.create(
            source=ExternalSource.ARPA, external_id="400721",
            external_name="بانک سپه شعبه سبلان شمالی",
            customer=self.customer, method="phone", score=Decimal("0.5"),
            payload=self.payload,
        )
        self.user = _user("reviewer", "manager", "sales_team")

    def test_accepting_links_the_party_and_writes_accounting_fields(self):
        customer = crm_merge.accept(self.candidate, self.user)

        self.assertEqual(customer.pk, self.customer.pk)
        self.assertTrue(customer.external_refs.filter(
            source=ExternalSource.ARPA, external_id="400721"
        ).exists())
        customer.refresh_from_db()
        self.assertEqual(customer.economic_code, "411431776417")
        self.assertEqual(customer.payment_terms, "30روزه")

    def test_accepting_does_not_rename_the_account(self):
        """
        The sales team knows customers by the name on their own screen. آرپا's
        legal name is kept on the ref, where the review screen shows it; a
        silent rename mid-quarter is not an update.
        """
        crm_merge.accept(self.candidate, self.user)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.name_fa, "بانک سپه سبلان")
        self.assertEqual(
            self.customer.external_refs.get(source=ExternalSource.ARPA).external_name,
            "بانک سپه شعبه سبلان شمالی",
        )

    def test_rejecting_creates_the_account_rather_than_dropping_it(self):
        created = crm_merge.reject(self.candidate, self.user)

        self.assertNotEqual(created.pk, self.customer.pk)
        self.assertEqual(created.name_fa, "بانک سپه شعبه سبلان شمالی")
        self.assertTrue(created.external_refs.filter(external_id="400721").exists())
        self.assertFalse(
            self.customer.external_refs.filter(source=ExternalSource.ARPA).exists()
        )

    def test_a_decision_cannot_be_taken_twice(self):
        crm_merge.accept(self.candidate, self.user)
        with self.assertRaises(crm_merge.MergeError):
            crm_merge.reject(self.candidate, self.user)

    def test_one_decision_settles_the_rival_suggestions(self):
        """
        An «ambig» party is suggested against every customer sharing its name.
        Ruling once answers the question; leaving the rivals pending invites a
        second reviewer to contradict the first.
        """
        other = Customer.objects.create(
            code="didar-co-501", name_fa="بانک سپه سبلان",
            dataset=Dataset.REAL, first_contact_at=timezone.now(),
        )
        rival = CustomerMatchCandidate.objects.create(
            source=ExternalSource.ARPA, external_id="400721",
            external_name="بانک سپه شعبه سبلان شمالی",
            customer=other, method="ambig", score=Decimal(1),
            payload=self.payload,
        )
        crm_merge.accept(self.candidate, self.user)

        rival.refresh_from_db()
        self.assertEqual(rival.state, CustomerMatchCandidate.State.REJECTED)

    def test_a_reviewer_may_choose_a_different_target(self):
        other = Customer.objects.create(
            code="didar-co-502", name_fa="بانک سپه مرکزی",
            dataset=Dataset.REAL, first_contact_at=timezone.now(),
        )
        customer = crm_merge.accept(self.candidate, self.user, other)

        self.assertEqual(customer.pk, other.pk)
        self.assertTrue(other.external_refs.filter(external_id="400721").exists())

    def test_a_code_already_claimed_is_refused_rather_than_moved(self):
        """Silently re-pointing it would undo an earlier reviewer's decision."""
        # A different customer, because the constraint already forbids two
        # suggestions pairing one party with one customer.
        other = Customer.objects.create(
            code="didar-co-503", name_fa="بانک سپه دیگر",
            dataset=Dataset.REAL, first_contact_at=timezone.now(),
        )
        second = CustomerMatchCandidate.objects.create(
            source=ExternalSource.ARPA, external_id="400721",
            external_name="بانک سپه شعبه سبلان شمالی",
            customer=other, method="fuzzy", score=Decimal("0.9"),
            payload=self.payload,
        )
        crm_merge.reject(self.candidate, self.user)
        second.refresh_from_db()
        second.state = CustomerMatchCandidate.State.PENDING
        second.save(update_fields=["state"])
        with self.assertRaises(crm_merge.MergeError):
            crm_merge.accept(second, self.user)


class MergeReviewApiTests(DatasetTestCase):
    """The queue over HTTP, including who is allowed to empty it."""

    def setUp(self):
        super().setUp()
        self.candidate = CustomerMatchCandidate.objects.create(
            source=ExternalSource.ARPA, external_id="400900",
            external_name="شرکت آزمایشی", customer=self.real_customer,
            method="phone", score=Decimal("0.5"),
            payload={"کد": "400900", "نام": "شرکت آزمایشی", "نوع": "حقوقی"},
        )

    def test_the_queue_lists_both_sides_in_one_request(self):
        self.client.force_authenticate(self.ceo)
        res = self.client.get("/api/crm/match-candidates/")
        self.assertEqual(res.status_code, 200)
        row = res.data["results"][0]
        self.assertEqual(row["arpa"]["name_fa"], "شرکت آزمایشی")
        self.assertEqual(row["crm"]["name_fa"], self.real_customer.name_fa)
        # The weight behind the decision, not just the two names.
        self.assertIn("deals", row["crm"])

    def test_an_old_demo_preference_does_not_hide_the_queue(self):
        """The showroom is gone; an account that last looked at it must still
        see the real queue rather than an empty one."""
        self.client.force_authenticate(self.ceo)
        self.ceo.crm_dataset = "demo"
        self.ceo.save(update_fields=["crm_dataset"])
        res = self.client.get("/api/crm/match-candidates/")
        self.assertEqual(res.data["count"], 1)

    def test_accepting_over_http_links_the_party(self):
        self.client.force_authenticate(self.ceo)
        res = self.client.post(f"/api/crm/match-candidates/{self.candidate.pk}/accept/")
        self.assertEqual(res.status_code, 200, res.data)
        self.assertTrue(
            self.real_customer.external_refs.filter(external_id="400900").exists()
        )

    def test_a_second_decision_is_refused_not_silently_applied(self):
        self.client.force_authenticate(self.ceo)
        self.client.post(f"/api/crm/match-candidates/{self.candidate.pk}/accept/")
        res = self.client.post(f"/api/crm/match-candidates/{self.candidate.pk}/reject/")
        self.assertEqual(res.status_code, 409)

    def test_summary_counts_what_is_still_open(self):
        self.client.force_authenticate(self.ceo)
        res = self.client.get("/api/crm/match-candidates/summary/")
        self.assertEqual(res.data["pending"], 1)


class AbsorbTests(APITestCase):
    """
    Fusing two CRM rows that are one company.

    The duplicate is kept and flagged, never deleted. `Deal.customer` cascades,
    so deleting the loser would take its deals, their lines and their stage
    history with it — a tidier list bought with history, and unnoticed until
    someone asks why a customer's numbers dropped.
    """

    def setUp(self):
        now = timezone.now()
        self.primary = Customer.objects.create(
            code="didar-co-700", name_fa="ایران ارقام", dataset=Dataset.REAL,
            first_contact_at=now, phone="",
        )
        self.dupe = Customer.objects.create(
            code="arpa-700", name_fa="شرکت ایران ارقام", dataset=Dataset.REAL,
            first_contact_at=now - timedelta(days=400),
            phone="02142719000", national_id="10100905654",
        )
        self.deal = Deal.objects.create(
            code="didar-d-700", title="معامله", customer=self.dupe,
            dataset=Dataset.REAL, opened_at=now,
        )

    def test_the_survivor_takes_the_records(self):
        crm_merge.absorb(self.primary, self.dupe)

        self.deal.refresh_from_db()
        self.assertEqual(self.deal.customer_id, self.primary.pk)
        self.assertEqual(self.primary.deals.count(), 1)

    def test_the_duplicate_is_flagged_not_deleted(self):
        crm_merge.absorb(self.primary, self.dupe)

        self.dupe.refresh_from_db()
        self.assertEqual(self.dupe.merged_into_id, self.primary.pk)
        self.assertFalse(self.dupe.is_active)
        self.assertTrue(Customer.objects.filter(pk=self.dupe.pk).exists())

    def test_what_only_the_duplicate_knew_is_carried_over(self):
        """
        Two rows usually exist because each system learned something the other
        did not. Keeping only the survivor's blanks throws that away.
        """
        crm_merge.absorb(self.primary, self.dupe)

        self.primary.refresh_from_db()
        self.assertEqual(self.primary.phone, "02142719000")
        self.assertEqual(self.primary.national_id, "10100905654")

    def test_the_relationship_starts_at_the_earlier_of_the_two(self):
        crm_merge.absorb(self.primary, self.dupe)

        self.primary.refresh_from_db()
        self.assertEqual(self.primary.first_contact_at, self.dupe.first_contact_at)

    def test_a_merged_row_cannot_be_merged_again(self):
        crm_merge.absorb(self.primary, self.dupe)
        third = Customer.objects.create(
            code="arpa-701", name_fa="ایران ارقام ۳", dataset=Dataset.REAL,
            first_contact_at=timezone.now(),
        )
        with self.assertRaises(crm_merge.MergeError):
            crm_merge.absorb(third, self.dupe)

    def test_a_customer_cannot_absorb_itself(self):
        with self.assertRaises(crm_merge.MergeError):
            crm_merge.absorb(self.primary, self.primary)

    def test_a_merged_row_leaves_the_customer_list(self):
        crm_merge.absorb(self.primary, self.dupe)
        user = _user("lister", "manager", "sales_team")
        self.client.force_authenticate(user)

        names = [
            r["name_fa"]
            for r in self.client.get("/api/crm/customers/").data["results"]
        ]
        self.assertIn("ایران ارقام", names)
        self.assertNotIn("شرکت ایران ارقام", names)


class BulkReviewTests(APITestCase):
    """Sending customers from the list into the merge queue."""

    def setUp(self):
        now = timezone.now()
        self.a = Customer.objects.create(
            code="didar-co-800", name_fa="پارس رول", dataset=Dataset.REAL,
            first_contact_at=now, phone="02155551234",
        )
        self.b = Customer.objects.create(
            code="arpa-800", name_fa="شرکت پارس رول", dataset=Dataset.REAL,
            first_contact_at=now, phone="02155551234",
        )
        self.c = Customer.objects.create(
            code="arpa-801", name_fa="هیچ‌کس", dataset=Dataset.REAL,
            first_contact_at=now,
        )
        self.user = _user("bulk", "manager", "sales_team")
        self.client.force_authenticate(self.user)

    def test_two_selected_rows_are_queued_as_a_pair(self):
        res = self.client.post(
            "/api/crm/customers/bulk-review/", {"ids": [self.a.pk, self.b.pk]},
            format="json",
        )
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data["queued"], 1)
        candidate = CustomerMatchCandidate.objects.get(source=ExternalSource.CRM)
        self.assertEqual(
            {candidate.customer_id, candidate.duplicate_id}, {self.a.pk, self.b.pk}
        )

    def test_the_same_pair_is_not_queued_twice(self):
        """
        Selecting the pair the other way round is the same question. Two rows
        would mean two reviewers ruling separately, and nothing stopping them
        from disagreeing.
        """
        for ids in ([self.a.pk, self.b.pk], [self.b.pk, self.a.pk]):
            self.client.post(
                "/api/crm/customers/bulk-review/", {"ids": ids}, format="json"
            )
        self.assertEqual(
            CustomerMatchCandidate.objects.filter(source=ExternalSource.CRM).count(), 1
        )

    def test_one_selected_row_makes_the_matcher_hunt_for_its_twin(self):
        res = self.client.post(
            "/api/crm/customers/bulk-review/", {"ids": [self.b.pk]}, format="json"
        )
        self.assertEqual(res.data["queued"], 1)
        candidate = CustomerMatchCandidate.objects.get(source=ExternalSource.CRM)
        self.assertEqual(
            {candidate.customer_id, candidate.duplicate_id}, {self.a.pk, self.b.pk}
        )

    def test_a_row_with_no_twin_is_reported_not_silently_dropped(self):
        res = self.client.post(
            "/api/crm/customers/bulk-review/",
            {"ids": [self.c.pk, self.a.pk, self.b.pk]}, format="json",
        )
        reasons = [s["reason"] for s in res.data["skipped"]]
        self.assertTrue(any("پیدا نشد" in r for r in reasons), res.data)

    def test_queueing_never_merges_on_its_own(self):
        """
        Even an exact name match only queues. Across two systems a repeated
        name means one company; inside one file it more often means someone
        typed it twice, and the person who pressed the button can say which.
        """
        self.client.post(
            "/api/crm/customers/bulk-review/", {"ids": [self.a.pk, self.b.pk]},
            format="json",
        )
        self.a.refresh_from_db()
        self.b.refresh_from_db()
        self.assertIsNone(self.a.merged_into_id)
        self.assertIsNone(self.b.merged_into_id)

    def test_accepting_a_crm_pair_fuses_the_two_rows(self):
        self.client.post(
            "/api/crm/customers/bulk-review/", {"ids": [self.a.pk, self.b.pk]},
            format="json",
        )
        candidate = CustomerMatchCandidate.objects.get(source=ExternalSource.CRM)
        res = self.client.post(f"/api/crm/match-candidates/{candidate.pk}/accept/")
        self.assertEqual(res.status_code, 200, res.data)

        self.b.refresh_from_db()
        self.assertEqual(self.b.merged_into_id, self.a.pk)

    def test_rejecting_a_crm_pair_creates_nothing(self):
        """Both sides already exist, so «different customers» is just the
        answer — there is no party to file and no account to invent."""
        self.client.post(
            "/api/crm/customers/bulk-review/", {"ids": [self.a.pk, self.b.pk]},
            format="json",
        )
        before = Customer.objects.count()
        candidate = CustomerMatchCandidate.objects.get(source=ExternalSource.CRM)
        self.client.post(f"/api/crm/match-candidates/{candidate.pk}/reject/")

        self.assertEqual(Customer.objects.count(), before)
        self.b.refresh_from_db()
        self.assertIsNone(self.b.merged_into_id)


class BulkDeleteTests(APITestCase):
    """
    Deleting from the list, and refusing to when it would cost history.

    `Deal.customer` cascades and `SalesInvoice.customer` is PROTECT, so a
    naive bulk delete either destroys deals silently or dies halfway with a
    database error naming no row. Both are worse than saying which customers
    could not go, and why.
    """

    def setUp(self):
        now = timezone.now()
        self.empty = Customer.objects.create(
            code="arpa-900", name_fa="بی‌سابقه", dataset=Dataset.REAL,
            first_contact_at=now,
        )
        self.with_deal = Customer.objects.create(
            code="arpa-901", name_fa="با معامله", dataset=Dataset.REAL,
            first_contact_at=now,
        )
        Deal.objects.create(
            code="d-901", title="م", customer=self.with_deal,
            dataset=Dataset.REAL, opened_at=now,
        )
        self.with_invoice = Customer.objects.create(
            code="arpa-902", name_fa="با فاکتور", dataset=Dataset.REAL,
            first_contact_at=now,
        )
        SalesInvoice.objects.create(
            code="arpa-inv-902", number="902", customer=self.with_invoice,
            issued_at=now.date(), amount_rial=1000, dataset=Dataset.REAL,
        )
        self.client.force_authenticate(_user("deleter", "manager", "sales_team"))

    def _delete(self, *rows):
        return self.client.post(
            "/api/crm/customers/bulk-delete/",
            {"ids": [r.pk for r in rows]}, format="json",
        )

    def test_an_empty_customer_is_deleted(self):
        res = self._delete(self.empty)
        self.assertEqual(res.data["deleted"], 1)
        self.assertFalse(Customer.objects.filter(pk=self.empty.pk).exists())

    def test_history_blocks_the_delete_and_says_what_held_it(self):
        res = self._delete(self.with_deal, self.with_invoice)

        self.assertEqual(res.data["deleted"], 0)
        self.assertTrue(Customer.objects.filter(pk=self.with_deal.pk).exists())
        self.assertTrue(Customer.objects.filter(pk=self.with_invoice.pk).exists())
        reasons = {b["name_fa"]: b["reason"] for b in res.data["blocked"]}
        self.assertIn("معامله", reasons["با معامله"])
        self.assertIn("فاکتور", reasons["با فاکتور"])

    def test_a_mixed_selection_deletes_what_it_can(self):
        """A refusal on one row must not strand the rest — a reviewer clearing
        1,600 dormant accounts cannot be made to retry them one at a time."""
        res = self._delete(self.empty, self.with_deal)

        self.assertEqual(res.data["deleted"], 1)
        self.assertEqual(len(res.data["blocked"]), 1)
        self.assertFalse(Customer.objects.filter(pk=self.empty.pk).exists())
        self.assertTrue(Customer.objects.filter(pk=self.with_deal.pk).exists())

    def test_an_empty_selection_is_refused(self):
        res = self.client.post(
            "/api/crm/customers/bulk-delete/", {"ids": []}, format="json"
        )
        self.assertEqual(res.status_code, 400)


# --------------------------------------------------------------------------
# Row-level scope
# --------------------------------------------------------------------------
class ScopeTests(APITestCase):
    """
    Each کارشناس has their own login and their own book.

    This is the whole reason a rep can be given an account: the guarantee is
    not "the screen does not show a button for other people's customers", it
    is "the API does not answer for them". So every test here goes at the API,
    by-id routes included — a list that filters and a `retrieve` that does not
    is a leak you find by guessing an integer.
    """

    def setUp(self):
        now = timezone.now()
        self.mine = DimEmployee.objects.create(code="e-mine", full_name_fa="کارشناس من")
        self.theirs = DimEmployee.objects.create(code="e-theirs", full_name_fa="کارشناس دیگر")

        self.rep = _user("rep1", "operator", "sales_team")
        self.mine.user = self.rep
        self.mine.save(update_fields=["user"])

        self.other_rep = _user("rep2", "operator", "sales_team")
        self.theirs.user = self.other_rep
        self.theirs.save(update_fields=["user"])

        self.boss = _user("boss", "manager", "sales_team")
        # In a CRM department but never linked to a salesperson row.
        self.unlinked = _user("ghost", "operator", "sales_b2b")

        self.stage = PipelineStage.objects.create(
            code="s-open", name_fa="ارتباط مشتری", kind="open", order=1,
            probability_pct=10, dataset=Dataset.REAL,
        )
        self.product = Product.objects.create(
            code="p-1", name_fa="کالا", dataset=Dataset.REAL,
            list_price_rial=1000, unit_cost_rial=400,
        )

        self.my_customer = Customer.objects.create(
            code="c-mine", name_fa="مشتری من", owner=self.mine,
            dataset=Dataset.REAL, first_contact_at=now - timedelta(days=10),
        )
        self.their_customer = Customer.objects.create(
            code="c-theirs", name_fa="مشتری دیگری", owner=self.theirs,
            dataset=Dataset.REAL, first_contact_at=now - timedelta(days=10),
        )
        self.my_deal = Deal.objects.create(
            code="d-mine", title="معامله من", customer=self.my_customer,
            owner=self.mine, stage=self.stage, dataset=Dataset.REAL,
            opened_at=now - timedelta(days=5),
        )
        self.their_deal = Deal.objects.create(
            code="d-theirs", title="معامله دیگری", customer=self.their_customer,
            owner=self.theirs, stage=self.stage, dataset=Dataset.REAL,
            opened_at=now - timedelta(days=5),
        )
        self.their_item = DealItem.objects.create(
            deal=self.their_deal, product=self.product, quantity=1,
            unit_price_rial=1000, unit_cost_rial=400, dataset=Dataset.REAL,
        )
        for owner, customer in ((self.mine, self.my_customer),
                                (self.theirs, self.their_customer)):
            Activity.objects.create(
                kind="call_out", customer=customer, owner=owner,
                at=now - timedelta(days=1), dataset=Dataset.REAL,
            )

    # ---- reading ---------------------------------------------------------
    def test_rep_lists_only_their_own_customers(self):
        self.client.force_authenticate(self.rep)
        rows = self.client.get("/api/crm/customers/").data["results"]
        self.assertEqual([r["name_fa"] for r in rows], ["مشتری من"])

    def test_rep_lists_only_their_own_deals_and_activities(self):
        self.client.force_authenticate(self.rep)
        deals = self.client.get("/api/crm/deals/", {"status": "open"}).data["results"]
        self.assertEqual([d["title"] for d in deals], ["معامله من"])

        acts = self.client.get("/api/crm/activities/").data["results"]
        self.assertEqual([a["customer_name"] for a in acts], ["مشتری من"])

    def test_manager_lists_the_whole_team(self):
        self.client.force_authenticate(self.boss)
        names = {r["name_fa"] for r in self.client.get("/api/crm/customers/").data["results"]}
        self.assertEqual(names, {"مشتری من", "مشتری دیگری"})

    def test_rep_cannot_fetch_another_reps_record_by_id(self):
        """The half a filtered list does not cover."""
        self.client.force_authenticate(self.rep)
        self.assertEqual(
            self.client.get("/api/crm/deals/%d/" % self.their_deal.id).status_code, 404
        )
        self.assertEqual(
            self.client.get("/api/crm/customers/%d/" % self.their_customer.id).status_code, 404
        )
        self.assertEqual(
            self.client.get("/api/crm/deal-items/%d/" % self.their_item.id).status_code, 404
        )

    def test_rep_cannot_read_another_reps_deal_lines(self):
        self.client.force_authenticate(self.rep)
        rows = self.client.get(
            "/api/crm/deal-items/", {"deal": self.their_deal.id}
        ).data["results"]
        self.assertEqual(rows, [])

    def test_an_owner_param_cannot_widen_a_reps_view(self):
        """
        The scope comes from the account; the query string does not get a
        vote. Asking for someone else's book returns your own rather than
        theirs — the forced owner simply overwrites whatever was asked for,
        which is the property worth having.
        """
        self.client.force_authenticate(self.rep)
        rows = self.client.get(
            "/api/crm/customers/", {"owner": self.theirs.id}
        ).data["results"]
        self.assertEqual([r["name_fa"] for r in rows], ["مشتری من"])

    def test_an_account_with_no_employee_row_sees_nothing(self):
        """A half-finished setup must fail closed, not open."""
        self.client.force_authenticate(self.unlinked)
        self.assertEqual(self.client.get("/api/crm/customers/").data["results"], [])
        self.assertEqual(self.client.get("/api/crm/deals/").data["results"], [])
        me = self.client.get("/api/crm/me/").data
        self.assertFalse(me["sees_all"])
        self.assertTrue(me["unlinked"])

    def test_analytics_are_scoped_too(self):
        """A report is just another way to read the rows."""
        self.client.force_authenticate(self.rep)
        board = self.client.get("/api/crm/pipeline/").data["columns"]
        self.assertEqual([d["title"] for col in board for d in col["deals"]], ["معامله من"])

        self.client.force_authenticate(self.unlinked)
        board = self.client.get("/api/crm/pipeline/").data["columns"]
        self.assertEqual([d for col in board for d in col["deals"]], [])

    # ---- writing ---------------------------------------------------------
    def test_a_rep_owns_what_they_create_whatever_the_payload_says(self):
        self.client.force_authenticate(self.rep)
        res = self.client.post(
            "/api/crm/customers/",
            {"name_fa": "سرنخ تازه", "owner": self.theirs.id},
            format="json",
        )
        self.assertEqual(res.status_code, 201)
        self.assertEqual(Customer.objects.get(name_fa="سرنخ تازه").owner_id, self.mine.id)

    def test_a_rep_cannot_hand_a_record_to_someone_else(self):
        self.client.force_authenticate(self.rep)
        res = self.client.patch(
            "/api/crm/customers/%d/" % self.my_customer.id,
            {"owner": self.theirs.id}, format="json",
        )
        self.assertEqual(res.status_code, 200)
        self.my_customer.refresh_from_db()
        self.assertEqual(self.my_customer.owner_id, self.mine.id)

    def test_a_manager_may_enter_on_someone_elses_behalf(self):
        self.client.force_authenticate(self.boss)
        res = self.client.post(
            "/api/crm/customers/",
            {"name_fa": "ثبت مدیر", "owner": self.theirs.id},
            format="json",
        )
        self.assertEqual(res.status_code, 201)
        self.assertEqual(Customer.objects.get(name_fa="ثبت مدیر").owner_id, self.theirs.id)

    def test_merge_queue_is_for_managers_only(self):
        self.client.force_authenticate(self.rep)
        self.assertEqual(self.client.get("/api/crm/match-candidates/").status_code, 403)
        self.client.force_authenticate(self.boss)
        self.assertEqual(self.client.get("/api/crm/match-candidates/").status_code, 200)


class DepartmentAccessTests(APITestCase):
    """Which departments may open CRM at all."""

    def test_every_sales_department_is_in(self):
        for dept in ("sales_team", "sales_b2b", "sales_org"):
            self.client.force_authenticate(_user("in-" + dept, "manager", dept))
            self.assertEqual(self.client.get("/api/crm/me/").status_code, 200, dept)

    def test_the_other_departments_are_not(self):
        for dept in ("production", "finance", "commercial"):
            self.client.force_authenticate(_user("out-" + dept, "manager", dept))
            self.assertEqual(self.client.get("/api/crm/me/").status_code, 403, dept)


# --------------------------------------------------------------------------
# Separate books
# --------------------------------------------------------------------------
class ChannelTests(APITestCase):
    """
    Three sales departments, three customer files.

    فروش همکار, فروش بانکی and فروش B2B sell to different books through the
    same screens. The channel column already existed on Customer and Deal and
    on the sales facts these reports sit beside, so the question is only
    whether CRM honours it — everywhere, including the paths that build their
    queryset by hand rather than through `Filters`.
    """

    def setUp(self):
        now = timezone.now()
        self.team_mgr = _user("m-team", "manager", "sales_team")
        self.bank_mgr = _user("m-bank", "manager", "sales_org")
        self.b2b_mgr = _user("m-b2b", "manager", "sales_b2b")
        self.ceo = _user("ceo-all", "executive")

        self.stage = PipelineStage.objects.create(
            code="s-open", name_fa="ارتباط مشتری", kind="open", order=1,
            probability_pct=10, dataset=Dataset.REAL,
        )

        self.books = {}
        for key, channel, label in (
            ("team", SalesChannel.TEAM, "مشتری همکار"),
            ("bank", SalesChannel.ORGANIZATIONAL, "مشتری بانکی"),
            ("b2b", SalesChannel.B2B, "مشتری بی‌تو‌بی"),
        ):
            customer = Customer.objects.create(
                code="c-" + key, name_fa=label, channel=channel,
                dataset=Dataset.REAL, first_contact_at=now - timedelta(days=10),
            )
            Deal.objects.create(
                code="d-" + key, title="معامله " + label, customer=customer,
                channel=channel, stage=self.stage, dataset=Dataset.REAL,
                opened_at=now - timedelta(days=5),
            )
            Activity.objects.create(
                kind="call_out", customer=customer, at=now - timedelta(days=1),
                dataset=Dataset.REAL,
            )
            self.books[key] = customer

    def names(self, url, key="name_fa", **params):
        return [r[key] for r in self.client.get(url, params).data["results"]]

    # ---- reading ---------------------------------------------------------
    def test_each_department_sees_only_its_own_book(self):
        for user, expected in (
            (self.team_mgr, "مشتری همکار"),
            (self.bank_mgr, "مشتری بانکی"),
            (self.b2b_mgr, "مشتری بی‌تو‌بی"),
        ):
            self.client.force_authenticate(user)
            self.assertEqual(self.names("/api/crm/customers/"), [expected])
            self.assertEqual(
                self.names("/api/crm/deals/", "title", status="open"),
                ["معامله " + expected],
            )
            self.assertEqual(
                self.names("/api/crm/activities/", "customer_name"), [expected]
            )

    def test_the_ceo_reads_every_book(self):
        self.client.force_authenticate(self.ceo)
        self.assertEqual(
            set(self.names("/api/crm/customers/")),
            {"مشتری همکار", "مشتری بانکی", "مشتری بی‌تو‌بی"},
        )

    def test_a_channel_param_cannot_reach_another_book(self):
        """Same rule as `owner`: the query string may narrow, never widen."""
        self.client.force_authenticate(self.bank_mgr)
        self.assertEqual(
            self.names("/api/crm/customers/", channel="team"), ["مشتری بانکی"]
        )

    def test_the_ceo_may_narrow_to_one_book(self):
        self.client.force_authenticate(self.ceo)
        self.assertEqual(
            self.names("/api/crm/customers/", channel="b2b"), ["مشتری بی‌تو‌بی"]
        )

    def test_the_pipeline_board_is_per_book(self):
        """A hand-built queryset is exactly where a channel filter gets lost."""
        self.client.force_authenticate(self.b2b_mgr)
        board = self.client.get("/api/crm/pipeline/").data["columns"]
        titles = [d["title"] for col in board for d in col["deals"]]
        self.assertEqual(titles, ["معامله مشتری بی‌تو‌بی"])

    def test_a_deal_cannot_be_fetched_across_books_by_id(self):
        self.client.force_authenticate(self.team_mgr)
        theirs = Deal.objects.get(code="d-bank")
        self.assertEqual(
            self.client.get("/api/crm/deals/%d/" % theirs.id).status_code, 404
        )

    # ---- writing ---------------------------------------------------------
    def test_a_new_customer_joins_the_filing_departments_book(self):
        self.client.force_authenticate(self.bank_mgr)
        res = self.client.post(
            "/api/crm/customers/", {"name_fa": "مشتری تازه بانکی"}, format="json"
        )
        self.assertEqual(res.status_code, 201)
        self.assertEqual(
            Customer.objects.get(name_fa="مشتری تازه بانکی").channel,
            SalesChannel.ORGANIZATIONAL,
        )

    def test_a_deal_takes_its_channel_from_its_customer(self):
        """
        Even when the CEO — who is in every book and therefore pinned to none
        — is the one entering it. The model default is «همکار», so without
        this the deal would file itself into the wrong department.
        """
        self.client.force_authenticate(self.ceo)
        res = self.client.post("/api/crm/deals/", {
            "customer": self.books["b2b"].id,
            "title": "معامله جدید",
            "stage": self.stage.id,
        }, format="json")
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(Deal.objects.get(title="معامله جدید").channel, SalesChannel.B2B)


class RosterScopeTests(APITestCase):
    """The کارشناس list a department is offered belongs to that department."""

    def setUp(self):
        now = timezone.now()
        self.mine = DimEmployee.objects.create(code="r-bank", full_name_fa="کارشناس بانکی")
        self.theirs = DimEmployee.objects.create(code="r-team", full_name_fa="کارشناس همکار")
        EmployeeChannel.objects.create(
            employee=self.mine, channel=SalesChannel.ORGANIZATIONAL, is_active=True
        )
        EmployeeChannel.objects.create(
            employee=self.theirs, channel=SalesChannel.TEAM, is_active=True
        )
        # Nobody's roster, but they already own a customer in the bank book —
        # the nine unrostered salespeople the real data actually has.
        self.legacy = DimEmployee.objects.create(code="r-old", full_name_fa="کارشناس قدیمی")
        Customer.objects.create(
            code="c-legacy", name_fa="مشتری قدیمی", owner=self.legacy,
            channel=SalesChannel.ORGANIZATIONAL, dataset=Dataset.REAL,
            first_contact_at=now - timedelta(days=400),
        )

    def options_names(self):
        return {e["name"] for e in self.client.get("/api/crm/options/").data["employees"]}

    def test_a_department_is_offered_its_own_roster(self):
        self.client.force_authenticate(_user("bank-mgr", "manager", "sales_org"))
        names = self.options_names()
        self.assertIn("کارشناس بانکی", names)
        self.assertNotIn("کارشناس همکار", names)

    def test_someone_who_already_owns_a_row_is_still_offered(self):
        """
        Otherwise opening an existing customer shows an owner picker that does
        not contain that customer's owner.
        """
        self.client.force_authenticate(_user("bank-mgr2", "manager", "sales_org"))
        self.assertIn("کارشناس قدیمی", self.options_names())

    def test_the_ceo_is_offered_everyone(self):
        self.client.force_authenticate(_user("ceo-roster", "executive"))
        names = self.options_names()
        self.assertIn("کارشناس بانکی", names)
        self.assertIn("کارشناس همکار", names)


class ReimportRespectsTheAppTests(APITestCase):
    """
    What a second accounting import may and may not undo.

    The failure these pin down was measured, not imagined: deploy re-ran the
    import on every release, 18 of 20 customers deleted in the app came back
    after one run, and ~1,800 accounts were rewritten from the workbook over
    whatever the team had corrected by hand.
    """

    HEADER = [
        "کد", "نام", "نام گروه", "شماره تلفن", "موبایل", "کد ملی",
        "کد اقتصادی", "نوع", "غیر فعال", "خوش حساب", "شرایط تسویه پیش فرض",
    ]

    def setUp(self):
        import tempfile
        self.dir = tempfile.mkdtemp()
        self.user = _user("reimporter", "manager", "sales_team")
        self.client.force_authenticate(self.user)

    def tearDown(self):
        import shutil
        shutil.rmtree(self.dir, ignore_errors=True)

    def _workbook(self, parties, buyers=()):
        import openpyxl
        from pathlib import Path

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(self.HEADER)
        for p in parties:
            ws.append([p.get(h, "") for h in self.HEADER])
        wb.save(Path(self.dir, "اشخاص کلی.xlsx"))

        sales = openpyxl.Workbook()
        ws = sales.active
        ws.append(["کد طرف حساب", "نوع برگه"])
        for code in buyers:
            ws.append([code, "فاکتور فروش"])
        sales.save(Path(self.dir, "فروش کل.xlsx"))

    def _import(self):
        from io import StringIO
        from django.core.management import call_command
        call_command("import_arpa_parties", dir=self.dir, stdout=StringIO())

    PARTY = {
        "کد": "555001", "نام": "کاغذ پردازان البرز", "نام گروه": "سایر طرف حسابها",
        "شماره تلفن": "02633334444", "کد ملی": "10101010101",
        "کد اقتصادی": "411111111111", "نوع": "حقوقی", "غیر فعال": "False",
        "خوش حساب": "False", "شرایط تسویه پیش فرض": "نقدی",
    }

    def _customer(self):
        return Customer.objects.get(code="arpa-555001")

    def test_a_deleted_customer_stays_deleted(self):
        self._workbook([self.PARTY])
        self._import()
        customer = self._customer()

        res = self.client.delete(f"/api/crm/customers/{customer.pk}/")
        self.assertIn(res.status_code, (200, 204), getattr(res, "data", None))
        self._import()

        self.assertFalse(Customer.objects.filter(code="arpa-555001").exists())
        self.assertTrue(DismissedParty.objects.filter(
            source=ExternalSource.ARPA, external_id="555001"
        ).exists())

    def test_a_bulk_deleted_customer_stays_deleted(self):
        self._workbook([self.PARTY])
        self._import()
        customer = self._customer()

        self.client.post(
            "/api/crm/customers/bulk-delete/", {"ids": [customer.pk]}, format="json"
        )
        self._import()

        self.assertFalse(Customer.objects.filter(code="arpa-555001").exists())

    def test_a_hand_corrected_field_survives_the_next_import(self):
        self._workbook([self.PARTY], buyers=["555001"])
        self._import()
        customer = self._customer()

        res = self.client.patch(
            f"/api/crm/customers/{customer.pk}/",
            {"phone": "02699998888", "national_id": "20202020202"}, format="json",
        )
        self.assertEqual(res.status_code, 200, res.data)
        self._import()

        customer.refresh_from_db()
        self.assertEqual(customer.phone, "02699998888")
        self.assertEqual(customer.national_id, "20202020202")

    def test_an_accounting_only_field_still_follows_accounting(self):
        """Payment terms cannot be edited in the CRM, so nobody's work is
        overwritten by keeping them current — and a stale term is wrong."""
        self._workbook([self.PARTY], buyers=["555001"])
        self._import()
        self._workbook([{**self.PARTY, "شرایط تسویه پیش فرض": "30روزه"}], buyers=["555001"])
        self._import()

        self.assertEqual(self._customer().payment_terms, "30روزه")

    def test_an_account_the_team_reopened_is_not_closed_again(self):
        """
        A party with no invoices arrives closed. If a rep reopens it — they
        are about to sell to it — the next import used to close it again,
        because «no invoices in the file» was re-derived on every run.
        """
        self._workbook([self.PARTY])
        self._import()
        customer = self._customer()
        self.assertFalse(customer.is_active)

        customer.is_active = True
        customer.save(update_fields=["is_active"])
        self._import()

        customer.refresh_from_db()
        self.assertTrue(customer.is_active)

    def test_an_account_that_starts_buying_is_opened(self):
        self._workbook([self.PARTY])
        self._import()
        self.assertFalse(self._customer().is_active)

        self._workbook([self.PARTY], buyers=["555001"])
        self._import()
        self.assertTrue(self._customer().is_active)

    def test_a_reviewer_created_account_arrives_open(self):
        """
        «Reject» on the review screen creates the account without the invoice
        workbooks at hand. Unknown sales used to read as «no sales» and the
        account was created closed — hidden from the list the reviewer had
        just added it to.
        """
        someone = Customer.objects.create(
            code="didar-co-555", name_fa="کاغذ پردازان", dataset=Dataset.REAL,
            first_contact_at=timezone.now(),
        )
        candidate = CustomerMatchCandidate.objects.create(
            source=ExternalSource.ARPA, external_id="555001",
            external_name=self.PARTY["نام"], customer=someone,
            method="fuzzy", score=Decimal("0.9"), payload=self.PARTY,
        )
        created = crm_merge.reject(candidate, self.user)
        self.assertTrue(created.is_active)


class DeployDoesNotLoadDataTests(APITestCase):
    """The deploy script must not run the accounting import on its own."""

    def test_deploy_script_does_not_call_the_arpa_importers(self):
        from pathlib import Path
        from django.conf import settings

        script = Path(settings.BASE_DIR).parent / "deploy.sh"
        commands = [
            line for line in script.read_text(encoding="utf-8").splitlines()
            if not line.lstrip().startswith("#")
        ]
        self.assertFalse(
            any("import_arpa_" in line for line in commands),
            "deploy.sh runs an آرپا import again — every release would undo "
            "deletes and hand edits made in the app since the last one.",
        )


class InvoiceLinkTests(APITestCase):
    """
    Attaching an invoice to the deal it billed — and refusing to guess.

    The 30-day window is measured, not chosen: gaps to the nearest won deal
    cluster within a month and then spread evenly to 90 days and beyond.
    Linking past that would attach invoices to whichever deal was least far
    away, and every per-deal figure built on it would be precise and wrong.
    """

    def setUp(self):
        from datetime import date
        self.day = date(2025, 10, 1)
        self.customer = Customer.objects.create(
            code="didar-co-l1", name_fa="مشتری لینک", dataset=Dataset.REAL,
            first_contact_at=timezone.now() - timedelta(days=500),
        )

    def _deal(self, code, days_from_invoice, amount=1_000_000, status="won"):
        closed = timezone.make_aware(
            timezone.datetime.combine(self.day + timedelta(days=days_from_invoice),
                                      timezone.datetime.min.time())
        )
        return Deal.objects.create(
            code=code, title=code, customer=self.customer, dataset=Dataset.REAL,
            status=status, opened_at=closed - timedelta(days=30), closed_at=closed,
            amount_rial=amount,
        )

    def _invoice(self, number, amount=1_000_000, customer=None):
        return SalesInvoice.objects.create(
            code=f"arpa-inv-l-{number}", number=number,
            customer=customer or self.customer, issued_at=self.day,
            amount_rial=amount, dataset=Dataset.REAL,
        )

    def test_the_nearest_won_deal_inside_the_window_is_linked(self):
        near = self._deal("d-near", 10)
        self._deal("d-far", 25)
        inv = self._invoice("1")

        stats = link_invoices()

        inv.refresh_from_db()
        self.assertEqual(inv.deal_id, near.pk)
        self.assertEqual(stats.linked, 1)

    def test_a_deal_outside_the_window_is_not_guessed(self):
        self._deal("d-old", -(WINDOW_DAYS + 5))
        inv = self._invoice("2")

        stats = link_invoices()

        inv.refresh_from_db()
        self.assertIsNone(inv.deal_id)
        self.assertEqual(stats.out_of_window, 1)

    def test_an_open_or_lost_deal_is_not_what_an_invoice_bills(self):
        self._deal("d-open", 2, status="open")
        self._deal("d-lost", 2, status="lost")
        inv = self._invoice("3")

        stats = link_invoices()

        inv.refresh_from_db()
        self.assertIsNone(inv.deal_id)
        self.assertEqual(stats.no_won_deal, 1)

    def test_on_a_tie_the_closer_amount_wins(self):
        """Two deals won the same day for one customer are two orders; the
        amount is what tells them apart."""
        self._deal("d-small", 5, amount=1_000_000)
        big = self._deal("d-big", 5, amount=90_000_000)
        inv = self._invoice("4", amount=88_000_000)

        link_invoices()

        inv.refresh_from_db()
        self.assertEqual(inv.deal_id, big.pk)

    def test_an_existing_link_is_never_replaced(self):
        first = self._deal("d-first", 20)
        inv = self._invoice("5")
        SalesInvoice.objects.filter(pk=inv.pk).update(deal=first)
        self._deal("d-closer", 1)

        stats = link_invoices()

        inv.refresh_from_db()
        self.assertEqual(inv.deal_id, first.pk)
        self.assertEqual(stats.already, 1)


class MergedSalesTests(APITestCase):
    """
    One sales figure: آرپا's invoices, attributed through دیدار.

    Replaces the side-by-side version. Showing won deals and invoices both
    as «فروش» meant every reader picked a number, and in 1404 the two were
    2x to 70x apart. The tests below pin the three things that make the one
    figure trustworthy: it reconciles to آرپا, it is attributed to a person
    wherever دیدار knows one, and it lands in the right department.
    """

    def setUp(self):
        from datetime import date
        self.ceo = _user("ceo-m", "executive")
        now = timezone.now()
        self.rep = DimEmployee.objects.create(code="e-m1", full_name_fa="کارشناس الف")
        self.deal_owner = DimEmployee.objects.create(code="e-m2", full_name_fa="کارشناس معامله")
        self.bank_rep = DimEmployee.objects.create(code="e-m3", full_name_fa="کارشناس بانکی")

        self.customer = Customer.objects.create(
            code="didar-co-m1", name_fa="مشتری فروش", dataset=Dataset.REAL,
            owner=self.rep, first_contact_at=now - timedelta(days=400),
        )
        self.bank = Customer.objects.create(
            code="arpa-m3", name_fa="بانک", dataset=Dataset.REAL,
            first_contact_at=now - timedelta(days=400),
            channel=SalesChannel.ORGANIZATIONAL,
        )
        self.start, self.end = date(2025, 9, 23), date(2025, 10, 22)  # مهر 1404
        self.day = date(2025, 10, 1)
        self.won = Deal.objects.create(
            code="d-m1", title="م", customer=self.customer, owner=self.deal_owner,
            dataset=Dataset.REAL, status="won", amount_rial=10_000_000,
            opened_at=now - timedelta(days=60),
            closed_at=timezone.make_aware(timezone.datetime(2025, 10, 5)),
        )
        self.unbilled = Deal.objects.create(
            code="d-m2", title="بی‌فاکتور", customer=self.bank, owner=self.bank_rep,
            dataset=Dataset.REAL, status="won", amount_rial=4_000_000,
            opened_at=now - timedelta(days=60),
            closed_at=timezone.make_aware(timezone.datetime(2025, 10, 6)),
        )

        def invoice(code, amount, **kw):
            kw.setdefault("customer", self.customer)
            kw.setdefault("channel", SalesChannel.TEAM)
            kw.setdefault("issued_at", self.day)
            return SalesInvoice.objects.create(
                code=f"arpa-inv-{code}", number=code, amount_rial=amount,
                unsettled_rial=amount, dataset=Dataset.REAL, **kw,
            )

        self.i_own = invoice("i1", 30_000_000, owner=self.rep)
        self.i_deal = invoice("i2", 5_000_000, deal=self.won)          # rep from deal
        self.i_customer = invoice("i3", 2_000_000)                     # rep from customer
        self.i_sister = invoice("i4", 900_000_000, is_intercompany=True)
        self.i_bank = invoice("i5", 7_000_000, customer=self.bank,
                              channel=SalesChannel.ORGANIZATIONAL, owner=self.bank_rep)
        # Arrived before its آرپا party was matched — still a sale.
        self.i_waiting = invoice("i6", 3_000_000, customer=None,
                                 party_code="777001", party_name="هنوز تطبیق‌نخورده")
        # The window's first day, which a datetime comparison drops.
        self.i_edge = invoice("i7", 1_000_000, owner=self.rep, issued_at=self.start)

    def _window(self):
        return {"date_from": self.start.isoformat(), "date_to": self.end.isoformat()}

    def _f(self, **extra):
        return rpt.Filters.from_query({**self._window(), **extra})

    SALES = 30_000_000 + 5_000_000 + 2_000_000 + 7_000_000 + 3_000_000 + 1_000_000

    # -- reconciliation -------------------------------------------------
    def test_sales_is_every_invoice_except_intercompany(self):
        total = self._f().invoices().aggregate(s=Sum("amount_rial"))["s"]
        self.assertEqual(total, self.SALES)

    def test_an_unmatched_invoice_is_still_a_sale(self):
        """
        An invoice whose party waits in the review queue used to be skipped,
        and that hid 24% of 1405 from every total.
        """
        self.assertTrue(self._f().invoices().filter(number="i6").exists())

    def test_won_deals_are_not_added_to_sales(self):
        f = self._f()
        rows = [r for r in rpt.report_sales(f, "time")["rows"] if r["amount"]]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["amount"], self.SALES)
        self.assertNotIn("invoiced", rows[0])

    def test_every_breakdown_sums_to_the_same_total(self):
        for axis in ("user", "team", "province", "group", "source", "customer"):
            data = rpt.report_sales(self._f(), axis)
            self.assertEqual(data["totals"]["amount"], self.SALES, axis)

    def test_an_invoice_on_the_first_day_of_the_window_is_inside_it(self):
        self.assertTrue(self._f().invoices().filter(number="i7").exists())

    # -- attribution ----------------------------------------------------
    def test_the_rep_falls_back_to_the_deal_then_the_customer_owner(self):
        reps = dict(self._f().invoices().values_list("number", "rep_id"))
        self.assertEqual(reps["i1"], self.rep.pk)          # invoice's own بازاریاب
        self.assertEqual(reps["i2"], self.deal_owner.pk)   # the linked deal
        self.assertEqual(reps["i3"], self.rep.pk)          # the customer's owner
        self.assertIsNone(reps["i6"])                      # nothing to go on

    def test_by_salesperson_names_what_is_left_unattributed(self):
        rows = {r["label"]: r["amount"] for r in rpt.report_sales(self._f(), "user")["rows"]}
        self.assertEqual(rows["کارشناس الف"], 30_000_000 + 2_000_000 + 1_000_000)
        self.assertEqual(rows["کارشناس معامله"], 5_000_000)
        self.assertEqual(rows["بدون کارشناس"], 3_000_000)
        self.assertNotIn("—", rows)

    def test_the_customer_axis_shows_the_review_queue_as_what_it_is(self):
        labels = {r["label"] for r in rpt.report_sales(self._f(), "customer")["rows"]}
        self.assertIn("در انتظار تطبیق مشتری", labels)

    # -- the dashboard --------------------------------------------------
    def test_the_dashboard_has_one_sales_card_that_reconciles(self):
        self.client.force_authenticate(self.ceo)
        res = self.client.get("/api/crm/dashboard/", self._window())
        self.assertEqual(res.status_code, 200)
        cards = {c["key"]: c for c in res.data["cards"]}
        self.assertNotIn("won", cards)
        self.assertNotIn("invoiced", cards)
        sales = cards["sales"]
        self.assertEqual(sales["value"], self.SALES)
        self.assertEqual(sales["sub"]["intercompany"], 900_000_000)
        self.assertEqual(sales["sub"]["unmatched"], 3_000_000)
        self.assertEqual(sales["drill"]["kind"], "invoices")

    def test_won_but_unbilled_is_pipeline_and_drills_to_those_deals(self):
        self.client.force_authenticate(self.ceo)
        cards = {c["key"]: c for c in self.client.get(
            "/api/crm/dashboard/", self._window()).data["cards"]}
        unbilled = cards["unbilled"]
        self.assertEqual(unbilled["value"], 4_000_000)

        res = self.client.get("/api/crm/deals/", unbilled["drill"]["params"])
        self.assertEqual({d["code"] for d in res.data["results"]}, {"d-m2"})

    def test_profit_is_labelled_as_deal_profit(self):
        self.assertEqual(rpt.report_profit(self._f(), "user")["title"], "سود معاملات")

    # -- scope ----------------------------------------------------------
    def test_a_department_sees_its_own_sales_by_the_invoice_department(self):
        self.client.force_authenticate(_user("bank-mgr-m", "manager", "sales_org"))
        res = self.client.get("/api/crm/invoices/", self._window())
        self.assertEqual({r["number"] for r in res.data["results"]}, {"i5"})

    def test_a_salesperson_sees_the_sales_attributed_to_them(self):
        """Including those credited through their customer — without the
        fallback a rep would not see their own accounts being billed."""
        user = _user("rep-m", "operator", "sales_team")
        self.rep.user = user
        self.rep.save(update_fields=["user"])
        self.client.force_authenticate(user)

        res = self.client.get("/api/crm/invoices/", self._window())
        self.assertEqual({r["number"] for r in res.data["results"]}, {"i1", "i3", "i7"})

    def test_the_drawer_names_an_unmatched_invoice_by_its_party(self):
        self.client.force_authenticate(self.ceo)
        rows = {r["number"]: r for r in self.client.get(
            "/api/crm/invoices/", self._window()).data["results"]}
        self.assertEqual(rows["i6"]["customer_name"], "هنوز تطبیق‌نخورده")
        self.assertTrue(rows["i6"]["awaiting_match"])
        self.assertEqual(rows["i3"]["owner_name"], "کارشناس الف")


class InvoiceAttachesOnMatchTests(APITestCase):
    """An invoice imported before its party is matched gets its customer
    the moment the party is resolved — by any path."""

    def setUp(self):
        from datetime import date
        self.waiting = SalesInvoice.objects.create(
            code="arpa-inv-w1", number="w1", customer=None, party_code="888001",
            party_name="طرف منتظر", issued_at=date(2025, 10, 1),
            amount_rial=5_000_000, dataset=Dataset.REAL,
        )
        self.customer = Customer.objects.create(
            code="didar-co-w1", name_fa="مشتری دیدار", dataset=Dataset.REAL,
            first_contact_at=timezone.now(),
        )
        self.payload = {"کد": "888001", "نام": "طرف منتظر", "نوع": "حقوقی"}

    def _candidate(self):
        return CustomerMatchCandidate.objects.create(
            source=ExternalSource.ARPA, external_id="888001",
            external_name="طرف منتظر", customer=self.customer,
            method="phone", score=Decimal("0.5"), payload=self.payload,
        )

    def test_accepting_a_match_attaches_the_waiting_invoice(self):
        crm_merge.accept(self._candidate())
        self.waiting.refresh_from_db()
        self.assertEqual(self.waiting.customer_id, self.customer.pk)

    def test_rejecting_a_match_attaches_it_to_the_new_account(self):
        created = crm_merge.reject(self._candidate())
        self.waiting.refresh_from_db()
        self.assertEqual(self.waiting.customer_id, created.pk)

    def test_an_intercompany_customer_marks_its_waiting_invoices(self):
        self.customer.is_intercompany = True
        self.customer.save(update_fields=["is_intercompany"])
        crm_merge.accept(self._candidate())
        self.waiting.refresh_from_db()
        self.assertTrue(self.waiting.is_intercompany)


class ChannelResolverTests(APITestCase):
    """
    Department from evidence, strongest first. Every imported customer used
    to be «همکار», so the bank department's manager saw none of its sales.
    """

    def setUp(self):
        from apps.crm.channels import ChannelResolver
        self.team_rep = DimEmployee.objects.create(code="c-t", full_name_fa="همکار")
        self.b2b_rep = DimEmployee.objects.create(code="c-b", full_name_fa="بی‌تو‌بی")
        self.two = DimEmployee.objects.create(code="c-2", full_name_fa="دو کانال")
        EmployeeChannel.objects.create(employee=self.team_rep, channel=SalesChannel.TEAM, is_active=True)
        EmployeeChannel.objects.create(employee=self.b2b_rep, channel=SalesChannel.B2B, is_active=True)
        EmployeeChannel.objects.create(employee=self.two, channel=SalesChannel.ORGANIZATIONAL, is_active=True)
        EmployeeChannel.objects.create(employee=self.two, channel=SalesChannel.PSP, is_active=True)
        self.r = ChannelResolver()

    def test_arpa_organizational_label_is_the_b2b_book(self):
        """«گروه فروش سازمانی» was mapped to `organizational`; the one rep
        آرپا tags with it sits on the B2B roster."""
        self.assertEqual(self.r.invoice("گروه فروش سازمانی", None, ""), SalesChannel.B2B)
        self.assertEqual(self.r.invoice("گروه فروش بانکی", None, ""), SalesChannel.ORGANIZATIONAL)

    def test_the_invoice_label_beats_the_roster(self):
        self.assertEqual(
            self.r.invoice("گروه فروش همکار", self.b2b_rep.pk, ""), SalesChannel.TEAM
        )

    def test_a_single_roster_channel_decides(self):
        self.assertEqual(self.r.invoice("", self.b2b_rep.pk, "سایر طرف حسابها"), SalesChannel.B2B)

    def test_a_rep_in_two_channels_decides_nothing(self):
        self.assertEqual(
            self.r.invoice("", self.two.pk, "نمابر مهر بانکها"), SalesChannel.ORGANIZATIONAL
        )
        self.assertEqual(self.r.invoice("", self.two.pk, "سایر طرف حسابها"), SalesChannel.TEAM)

    def test_the_group_decides_when_nothing_else_does(self):
        self.assertEqual(self.r.customer(None, "نمابر مهر سازمانها"), SalesChannel.B2B)
        self.assertEqual(self.r.customer(None, "مشتریان مشترک"), SalesChannel.TEAM)

    def test_rederiving_moves_imported_customers_and_their_deals(self):
        from apps.crm.channels import rederive_customer_channels
        now = timezone.now()
        imported = Customer.objects.create(
            code="arpa-ch1", name_fa="وارد شده", dataset=Dataset.REAL,
            owner=self.b2b_rep, first_contact_at=now,
        )
        CustomerExternalRef.objects.create(
            customer=imported, source=ExternalSource.ARPA, external_id="ch1",
        )
        typed = Customer.objects.create(
            code="c-typed", name_fa="ساخته در برنامه", dataset=Dataset.REAL,
            owner=self.b2b_rep, first_contact_at=now, channel=SalesChannel.TEAM,
        )
        deal = Deal.objects.create(
            code="d-ch1", title="م", customer=imported, dataset=Dataset.REAL,
            opened_at=now,
        )

        rederive_customer_channels()

        imported.refresh_from_db(); typed.refresh_from_db(); deal.refresh_from_db()
        self.assertEqual(imported.channel, SalesChannel.B2B)
        self.assertEqual(deal.channel, SalesChannel.B2B)
        # A customer a department created in the app is left where they put it.
        self.assertEqual(typed.channel, SalesChannel.TEAM)


class WorkScreensTests(APITestCase):
    """
    The screens added on top of the lists — کارتابل امروز, global search, the
    Excel exports and the bulk actions — each answer from the same scope as
    the lists. None of them may become the side door a list closed: a rep's
    search, worklist or export holds their own book and nothing else.

    Reuses ScopeTests' fixture: two reps with a customer and a deal each, and
    a manager over both.
    """

    setUp = ScopeTests.setUp

    def _task(self, owner, customer, days_ago, title):
        from apps.crm.models import Task
        return Task.objects.create(
            title=title, customer=customer, owner=owner, dataset=Dataset.REAL,
            due_at=timezone.now() - timedelta(days=days_ago),
        )

    # ---- کارتابل امروز ----------------------------------------------------
    def test_today_lists_only_the_reps_own_work(self):
        self._task(self.mine, self.my_customer, 2, "کار من")
        self._task(self.theirs, self.their_customer, 2, "کار دیگری")
        self.client.force_authenticate(self.rep)
        data = self.client.get("/api/crm/today/").data
        self.assertEqual([t["title"] for t in data["overdue"]], ["کار من"])

    def test_today_ignores_an_owner_param_from_a_rep(self):
        self._task(self.theirs, self.their_customer, 2, "کار دیگری")
        self.client.force_authenticate(self.rep)
        data = self.client.get("/api/crm/today/", {"owner": self.theirs.id}).data
        self.assertEqual(data["overdue"], [])

    def test_old_overdue_tasks_are_counted_not_listed(self):
        """Imported history must not bury today's work (see BACKLOG_DAYS)."""
        self._task(self.mine, self.my_customer, 2, "تازه")
        self._task(self.mine, self.my_customer, 200, "قدیمی")
        self.client.force_authenticate(self.rep)
        data = self.client.get("/api/crm/today/").data
        self.assertEqual([t["title"] for t in data["overdue"]], ["تازه"])
        self.assertEqual(data["counters"]["backlog"], 1)

    def test_unlinked_account_gets_an_empty_worklist(self):
        self._task(self.theirs, self.their_customer, 2, "کار دیگری")
        self.client.force_authenticate(self.unlinked)
        data = self.client.get("/api/crm/today/").data
        self.assertEqual(data["overdue"], [])
        self.assertEqual(data["counters"]["open_count"], 0)

    # ---- search -----------------------------------------------------------
    def test_search_finds_only_what_the_lists_would_show(self):
        self.client.force_authenticate(self.rep)
        data = self.client.get("/api/crm/search/", {"q": "معامله"}).data
        self.assertEqual([d["title"] for d in data["deals"]], ["معامله من"])
        data = self.client.get("/api/crm/search/", {"q": "مشتری"}).data
        self.assertEqual([c["name_fa"] for c in data["customers"]], ["مشتری من"])

    # ---- Excel ------------------------------------------------------------
    def _sheet_rows(self, res):
        from io import BytesIO
        from openpyxl import load_workbook
        self.assertEqual(res.status_code, 200, getattr(res, "data", None))
        ws = load_workbook(BytesIO(res.content)).worksheets[0]
        # Title block, blank line, header — data starts on row 5. The bold
        # «جمع کل» line under the data is a total, not a record.
        return [
            r for r in ws.iter_rows(min_row=5, values_only=True)
            if any(r) and r[0] != "جمع کل"
        ]

    def test_drill_export_holds_every_row_not_one_page(self):
        now = timezone.now()
        for i in range(40):
            Deal.objects.create(
                code=f"d-bulk-{i}", title=f"انبوه {i}", customer=self.my_customer,
                owner=self.mine, stage=self.stage, dataset=Dataset.REAL,
                opened_at=now - timedelta(days=2),
            )
        self.client.force_authenticate(self.rep)
        res = self.client.get("/api/crm/export/drill/", {"kind": "deals", "status": "open", "page_size": 25})
        titles = [r[1] for r in self._sheet_rows(res)]
        self.assertEqual(len(titles), 41)
        self.assertNotIn("معامله دیگری", titles)

    def test_drill_export_can_be_narrowed_to_ticked_rows(self):
        self.client.force_authenticate(self.boss)
        res = self.client.get("/api/crm/export/drill/", {
            "kind": "deals", "status": "open", "ids": str(self.their_deal.id),
        })
        self.assertEqual([r[1] for r in self._sheet_rows(res)], ["معامله دیگری"])

    def test_report_export_is_a_workbook(self):
        self.client.force_authenticate(self.boss)
        res = self.client.get("/api/crm/reports/incoming/export/", {"axis": "user"})
        self.assertEqual(res.status_code, 200)
        self.assertIn("spreadsheetml", res["Content-Type"])

    # ---- bulk actions -------------------------------------------------------
    def test_a_rep_cannot_reassign_in_bulk(self):
        self.client.force_authenticate(self.rep)
        res = self.client.post("/api/crm/deals/bulk-assign/", {
            "ids": [self.my_deal.id], "owner": self.theirs.id,
        }, format="json")
        self.assertEqual(res.status_code, 403)

    def test_bulk_move_cannot_reach_someone_elses_deal(self):
        other_stage = PipelineStage.objects.create(
            code="s-next", name_fa="مرحله بعد", kind="open", order=2,
            probability_pct=50, dataset=Dataset.REAL,
        )
        self.client.force_authenticate(self.rep)
        res = self.client.post("/api/crm/deals/bulk-move/", {
            "ids": [self.my_deal.id, self.their_deal.id], "stage": other_stage.id,
        }, format="json")
        self.assertEqual(res.data["moved"], 1)
        self.their_deal.refresh_from_db()
        self.assertEqual(self.their_deal.stage_id, self.stage.id)
        # And the one it did move left a stage event, like a drag would.
        self.assertTrue(self.my_deal.stage_events.filter(to_stage=other_stage).exists())

    # ---- lists --------------------------------------------------------------
    def test_lists_honour_the_page_size_they_ask_for(self):
        now = timezone.now()
        for i in range(12):
            Deal.objects.create(
                code=f"d-page-{i}", title=f"صفحه {i}", customer=self.my_customer,
                owner=self.mine, stage=self.stage, dataset=Dataset.REAL,
                opened_at=now - timedelta(days=2),
            )
        self.client.force_authenticate(self.rep)
        data = self.client.get("/api/crm/deals/", {"status": "open", "page_size": 5}).data
        self.assertEqual(len(data["results"]), 5)
        self.assertEqual(data["count"], 13)

    def test_unknown_ordering_is_ignored_not_obeyed(self):
        self.client.force_authenticate(self.boss)
        res = self.client.get("/api/crm/deals/", {"status": "open", "ordering": "owner__user__password"})
        self.assertEqual(res.status_code, 200)
