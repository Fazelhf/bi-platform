"""
Acting on a reviewer's decision about a suspected duplicate.

The matcher deliberately refuses to merge anything it is not certain of, which
leaves a queue: 160 pairs a person has to rule on. This is what happens when
they do.

Both answers end the same way — the party ends up with an external ref — and
that is the point. An unresolved party is invisible to the invoice import, so
its invoices are skipped; 260 invoices worth 543bn Rial are waiting on this
queue. «Reject» is not «discard»: it means *this is a different customer*, and
a different customer still needs an account.

Decisions never delete. Accepting writes accounting's fields onto the customer
the reviewer chose; rejecting creates a new customer. The one delete path,
`delete_customers`, lives here too so it can leave a tombstone behind. If a decision turns out
to be wrong, the ref moves — no history has been destroyed in the meantime.
"""
from __future__ import annotations

from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from apps.crm.models import (
    Customer, CustomerExternalRef, CustomerMatchCandidate, ExternalSource,
)
from apps.crm.party_sync import PartyWriter


class MergeError(Exception):
    """A decision that cannot be carried out as asked."""


def _settle(candidate, state, user, customer):
    candidate.state = state
    candidate.customer = customer
    candidate.decided_at = timezone.now()
    candidate.decided_by = user if user and user.is_authenticated else None
    candidate.save(update_fields=[
        "state", "customer", "decided_at", "decided_by", "updated_at",
    ])
    # The same party may have been suggested against several customers — an
    # «ambig» name matches every customer carrying it. One decision settles
    # the party, so the rival suggestions are not left for someone to rule on
    # a second time and contradict the first.
    CustomerMatchCandidate.objects.filter(
        source=candidate.source, external_id=candidate.external_id,
        state=CustomerMatchCandidate.State.PENDING,
    ).exclude(pk=candidate.pk).update(
        state=CustomerMatchCandidate.State.REJECTED,
        decided_at=candidate.decided_at,
        decided_by=candidate.decided_by,
    )


#: Everything that hangs off a customer and has to travel when two are fused.
#: Missing one would leave records pointing at a row nobody can see any more —
#: invisible, but still counted by any report that joins from the other end.
#: The sales2 rows are finance's: documents and receipts issued to the
#: absorbed customer would otherwise keep a balance on a hidden account.
OWNED = (
    ("deals", "customer"),
    ("activities", "customer"),
    ("tasks", "customer"),
    ("feedback", "customer"),
    ("invoices", "customer"),
    ("external_refs", "customer"),
    ("sales2_documents", "customer"),
    ("sales2_receipts", "customer"),
)

#: Filled on the survivor from the absorbed row when the survivor has nothing.
#: A merge that keeps only one side's phone number loses the other, and the
#: whole reason two rows existed is usually that each knew something.
CARRIED = (
    "phone", "mobile", "email", "address", "city", "national_id",
    "economic_code", "registration_no", "postal_code", "payment_terms",
    "contact_name",
)


def _snapshot(customer: Customer) -> dict:
    import json

    from django.core.serializers.json import DjangoJSONEncoder
    from django.forms.models import model_to_dict

    data = model_to_dict(customer, exclude=["tags"])
    data["tags"] = list(customer.tags.values_list("pk", flat=True))
    return json.loads(json.dumps(data, cls=DjangoJSONEncoder))


def _move_accounts(primary: Customer, duplicate: Customer, moved: dict) -> None:
    """
    Finance's per-customer accounts, which allow one per customer (per book).

    Re-pointed when the survivor has none. When both have one, the balances
    are two ledgers finance has to reconcile — choosing one here would
    silently drop the other, so the merge stops and says which.
    """
    from apps.sales2.models import CustomerAccount, LedgerAccount

    mine = set(LedgerAccount.objects.filter(customer=primary).values_list("book", flat=True))
    theirs = list(LedgerAccount.objects.filter(customer=duplicate))
    clash = [a.get_book_display() for a in theirs if a.book in mine]
    if clash:
        raise MergeError(
            "هر دو مشتری در «" + "، ".join(clash) + "» حساب دفتری دارند؛ "
            "پیش از ادغام، مالی باید حساب‌ها را یکی کند."
        )
    if theirs:
        LedgerAccount.objects.filter(pk__in=[a.pk for a in theirs]).update(customer=primary)
        moved["ledger_accounts"] = [a.pk for a in theirs]

    account = CustomerAccount.objects.filter(customer=duplicate).first()
    if account and CustomerAccount.objects.filter(customer=primary).exists():
        raise MergeError(
            "هر دو مشتری حساب اعتباری (سقف اعتبار) دارند؛ "
            "پیش از ادغام، مالی باید یکی را حذف کند."
        )
    if account:
        CustomerAccount.objects.filter(pk=account.pk).update(customer=primary)
        moved["sales2_account"] = [account.pk]


def _repoint_candidates(primary: Customer, duplicate: Customer) -> None:
    """
    Open review rows about the absorbed customer now concern the survivor.
    One that would pair the survivor with itself, or repeat a question
    already filed, is closed instead of left for someone to rule on twice.
    """
    for c in CustomerMatchCandidate.objects.filter(
        state=CustomerMatchCandidate.State.PENDING,
    ).filter(Q(customer=duplicate) | Q(duplicate=duplicate)):
        customer_id = primary.pk if c.customer_id == duplicate.pk else c.customer_id
        duplicate_id = primary.pk if c.duplicate_id == duplicate.pk else c.duplicate_id
        clash = customer_id == duplicate_id or CustomerMatchCandidate.objects.filter(
            source=c.source, external_id=c.external_id, customer_id=customer_id,
        ).exclude(pk=c.pk).exists()
        if clash:
            c.state = CustomerMatchCandidate.State.REJECTED
            c.decided_at = timezone.now()
            c.save(update_fields=["state", "decided_at", "updated_at"])
        else:
            CustomerMatchCandidate.objects.filter(pk=c.pk).update(
                customer_id=customer_id, duplicate_id=duplicate_id,
            )


@transaction.atomic
def absorb(primary: Customer, duplicate: Customer, user=None) -> Customer:
    """
    Fold one customer into another.

    The duplicate is kept, not deleted. Its deals and invoices move to the
    survivor, so nothing is orphaned, and the row itself stays behind carrying
    `merged_into` — which is what makes the decision reversible. Deleting it
    would take its `Deal` rows with it by cascade, and a merge that destroys
    history to tidy a list is a worse outcome than the duplicate was.

    Every re-pointed id and both rows' fields as they stood go to
    `CustomerMerge`, and the survivor's cached dates are recomputed from the
    records it now owns — a merge used to leave «آخرین فعالیت» at the
    survivor's own old date while the absorbed side's recent calls sat in its
    timeline.
    """
    from apps.crm.models import CustomerMerge

    if primary.pk == duplicate.pk:
        raise MergeError("یک مشتری را نمی‌شود در خودش ادغام کرد.")
    # Re-read under lock: two managers merging the same pair at once must not
    # both pass the checks below.
    primary = Customer.objects.select_for_update().get(pk=primary.pk)
    duplicate = Customer.objects.select_for_update().get(pk=duplicate.pk)
    if duplicate.merged_into_id:
        raise MergeError(f"«{duplicate.name_fa}» قبلاً ادغام شده است.")
    if primary.merged_into_id:
        raise MergeError(
            f"«{primary.name_fa}» خودش در مشتری دیگری ادغام شده است."
        )

    before = {"survivor": _snapshot(primary), "absorbed": _snapshot(duplicate)}
    moved: dict[str, list] = {}
    _move_accounts(primary, duplicate, moved)
    for accessor, field in OWNED:
        qs = getattr(duplicate, accessor).all()
        ids = list(qs.values_list("pk", flat=True))
        if ids:
            qs.model.objects.filter(pk__in=ids).update(**{field: primary})
            moved[accessor] = ids
    tags = list(duplicate.tags.values_list("pk", flat=True))
    if tags:
        primary.tags.add(*tags)
        moved["tags"] = tags
    # Rows merged into the absorbed one earlier now point at the survivor, so
    # `merged_into` never has to be followed more than one step.
    chained = list(duplicate.merged_from.values_list("pk", flat=True))
    if chained:
        Customer.objects.filter(pk__in=chained).update(merged_into=primary)
        moved["merged_from"] = chained
    _repoint_candidates(primary, duplicate)

    if duplicate.note and duplicate.note not in primary.note:
        primary.note = "\n\n".join(filter(None, [primary.note, duplicate.note]))
    primary.is_good_payer = primary.is_good_payer or duplicate.is_good_payer
    primary.is_intercompany = primary.is_intercompany or duplicate.is_intercompany
    if duplicate.vat_cert_expires_at and (
        not primary.vat_cert_expires_at
        or duplicate.vat_cert_expires_at > primary.vat_cert_expires_at
    ):
        primary.vat_cert_expires_at = duplicate.vat_cert_expires_at
    if primary.status == Customer.Status.LEAD and duplicate.status != Customer.Status.LEAD:
        primary.status = duplicate.status
    for field in CARRIED:
        if not getattr(primary, field) and getattr(duplicate, field):
            setattr(primary, field, getattr(duplicate, field))
    for field in ("group_id", "province_id", "owner_id", "lead_source_id"):
        if not getattr(primary, field) and getattr(duplicate, field):
            setattr(primary, field, getattr(duplicate, field))
    # The relationship started whenever the earlier of the two started.
    if duplicate.first_contact_at and (
        not primary.first_contact_at
        or duplicate.first_contact_at < primary.first_contact_at
    ):
        primary.first_contact_at = duplicate.first_contact_at
    primary.save()

    duplicate.merged_into = primary
    duplicate.is_active = False
    duplicate.save(update_fields=["merged_into", "is_active", "updated_at"])
    Customer.refresh_rollups([primary.pk])
    CustomerMerge.objects.create(
        survivor=primary, absorbed=duplicate,
        by=user if user and user.is_authenticated else None,
        moved=moved, before=before,
    )
    primary.refresh_from_db()
    return primary


@transaction.atomic
def accept(candidate: CustomerMatchCandidate, user=None, customer=None) -> Customer:
    """
    «They are the same customer.»

    `customer` overrides the suggestion. The screen offers it for the
    «ambig» tier, where the party's name matches several CRM rows and the
    matcher has no basis to prefer one — that choice is the reviewer's whole
    contribution and refusing it would send them to the admin to do it by hand.
    """
    if candidate.state != CustomerMatchCandidate.State.PENDING:
        raise MergeError("این مورد قبلاً تعیین تکلیف شده است.")

    # A pair raised from the customer list has no external party to file —
    # both sides are CRM rows, so «same customer» means fusing them.
    if candidate.source == ExternalSource.CRM:
        if not candidate.duplicate_id:
            raise MergeError("این مورد طرف دومی ندارد.")
        primary = customer or candidate.customer
        other = (
            candidate.duplicate
            if primary.pk != candidate.duplicate_id
            else candidate.customer
        )
        survivor = absorb(primary, other, user)
        _settle(candidate, CustomerMatchCandidate.State.ACCEPTED, user, survivor)
        return survivor

    target = customer or candidate.customer
    if target.merged_into_id:
        raise MergeError(
            f"«{target.name_fa}» خودش در مشتری دیگری ادغام شده است."
        )

    taken = CustomerExternalRef.objects.filter(
        source=candidate.source, external_id=candidate.external_id,
    ).exclude(customer=target).first()
    if taken:
        raise MergeError(
            f"این کد از قبل به «{taken.customer.name_fa}» وصل است."
        )

    writer = PartyWriter()
    writer.apply(target, candidate.payload)
    writer.link(target, candidate.payload)
    _settle(candidate, CustomerMatchCandidate.State.ACCEPTED, user, target)
    return target


@transaction.atomic
def reject(candidate: CustomerMatchCandidate, user=None) -> Customer:
    """
    «They are different customers.»

    Which makes the party a customer in its own right, so it is created here
    rather than dropped. Skipping that would leave its invoices unimportable
    for good, and the reviewer would have said «not a duplicate» and watched
    the account disappear.

    The new account arrives open. Whether it has ever traded is not knowable
    from this screen — the invoice files are not loaded here — and a reviewer
    who has just said «this is a customer of its own» should not find it
    hidden from the list. A later import opens accounts that start buying but
    never closes one, so closing it stays a decision for the team.
    """
    if candidate.state != CustomerMatchCandidate.State.PENDING:
        raise MergeError("این مورد قبلاً تعیین تکلیف شده است.")

    # Both sides already exist as accounts, so «different customers» is simply
    # the answer — there is nothing to create and nothing to undo.
    if candidate.source == ExternalSource.CRM:
        _settle(candidate, CustomerMatchCandidate.State.REJECTED, user,
                candidate.customer)
        return candidate.customer

    existing = CustomerExternalRef.objects.filter(
        source=candidate.source, external_id=candidate.external_id,
    ).first()
    if existing:
        customer = existing.customer
    else:
        writer = PartyWriter()
        customer = writer.create(candidate.payload)
        writer.link(customer, candidate.payload)

    _settle(candidate, CustomerMatchCandidate.State.REJECTED, user, candidate.customer)
    return customer


def queue_pair(primary: Customer, duplicate: Customer, method="manual",
               confidence=None, reason=""):
    """
    File two CRM rows as a suspected duplicate.

    Returns None when the pair is already known — in the queue, or already
    ruled on. Re-queueing a rejected pair would let a list of «these two look
    alike» quietly reopen a decision someone already made.
    """
    if primary.pk == duplicate.pk:
        return None
    # Order the pair so «A then B» and «B then A» are one row rather than two
    # questions about the same thing. The survivor is chosen at accept time,
    # so which side is stored first carries no meaning and must not.
    if primary.pk > duplicate.pk:
        primary, duplicate = duplicate, primary
    external_id = f"cust-{duplicate.pk}"
    if CustomerMatchCandidate.objects.filter(
        source=ExternalSource.CRM, external_id=external_id, customer=primary,
    ).exists():
        return None
    return CustomerMatchCandidate.objects.create(
        source=ExternalSource.CRM,
        external_id=external_id,
        external_name=duplicate.name_fa[:200],
        external_phone=(duplicate.phone or duplicate.mobile)[:40],
        external_city=duplicate.city[:100],
        customer=primary,
        duplicate=duplicate,
        method=method,
        score=1 if method == "manual" else 0,
        confidence=confidence or CustomerMatchCandidate.Confidence.SUSPECT,
        reason=reason[:200],
    )


def queue_scan(customers, dataset: str):
    """
    Hunt for the selected customers' twins and queue whatever turns up.

    For when someone suspects a row is a duplicate without knowing of what.
    It is `find_pairs` — the whole-file sweep — narrowed to pairs touching
    the selection, so the two paths cannot disagree about what a duplicate is.

    This used to run its own ladder with the selected rows left *out* of the
    index, so two copies that were both selected could never find each other,
    and each row stopped at its first match. Selecting a page of a file full
    of duplicates queued almost nothing, which is why the button looked broken.

    Nothing here merges: every pair waits for a person.
    """
    picked = {c.pk for c in customers}
    pairs, _ = find_pairs(dataset)
    mine = [p for p in pairs if p["a"]["pk"] in picked or p["b"]["pk"] in picked]
    rows = Customer.objects.in_bulk({p["a"]["pk"] for p in mine} | {p["b"]["pk"] for p in mine})

    queued, touched = [], set()
    for p in mine:
        touched |= {p["a"]["pk"], p["b"]["pk"]}
        row = queue_pair(rows[p["a"]["pk"]], rows[p["b"]["pk"]],
                         _method_for(p["reason"]), p["confidence"], p["reason"])
        if row:
            queued.append(row)
    skipped = [
        {"name_fa": c.name_fa,
         "reason": "این جفت از قبل در صف است یا تعیین تکلیف شده."
         if c.pk in touched else "مشابهی پیدا نشد."}
        for c in customers
        if c.pk not in {x.customer_id for x in queued} | {x.duplicate_id for x in queued}
    ]
    return queued, skipped


def survivor_of(a: Customer, b: Customer) -> tuple[Customer, Customer]:
    """
    Which of two rows to keep when nobody picked: (keep, absorb).

    The row with trading history first — invoices, then deals — because that
    is the one reports and the invoice import already point at. Then the
    accounting row, whose code the next آرپا import looks for. Then the older.
    """
    def rank(c):
        return (-c.invoices.count(), -c.deals.count(),
                0 if c.code.startswith("arpa-") else 1, c.pk)
    keep, other = sorted((a, b), key=rank)
    return keep, other


def accept_definite(dataset: str, user=None, dry_run=False) -> dict:
    """
    Merge every open CRM pair that is definite — the bulk half of review.

    «Definite» is `matching.confidence` re-run on the two rows as they stand
    *now*, not the label stored when the pair was queued: a phone corrected
    since then must not be merged on. Each pair goes through `accept`, so each
    is its own transaction, leaves the absorbed row behind with `merged_into`
    and a `CustomerMerge` record, and one failure does not undo the others.
    Suspect pairs are never touched.
    """
    from apps.crm.matching import confidence

    def definite(c):
        return bool(c.duplicate_id) and confidence(
            _as_dict(c.customer), _as_dict(c.duplicate)
        )[0] == "definite"

    pending = CustomerMatchCandidate.objects.filter(
        source=ExternalSource.CRM,
        state=CustomerMatchCandidate.State.PENDING,
        customer__dataset=dataset,
    ).select_related("customer", "duplicate").order_by("pk")
    sure = [c.pk for c in pending if definite(c)]
    if dry_run:
        return {"definite": len(sure), "merged": 0, "failed": []}

    merged, failed = 0, []
    for pk in sure:
        # Re-read: an earlier merge in the same group re-points or closes
        # the pairs that mentioned the row it absorbed.
        c = CustomerMatchCandidate.objects.select_related("customer", "duplicate").get(pk=pk)
        if c.state != CustomerMatchCandidate.State.PENDING or not definite(c):
            continue
        keep, _ = survivor_of(c.customer, c.duplicate)
        try:
            accept(c, user, keep)
            merged += 1
        except MergeError as exc:
            failed.append({"name_fa": c.customer.name_fa, "reason": str(exc)})
    return {"definite": len(sure), "merged": merged, "failed": failed}


_FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def _mark(value) -> str:
    """A cell as a mark: «۱», 1, 1.0 and " 1 " are all the same mark."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).translate(_FA_DIGITS).strip()


def read_merge_marks(fileobj) -> tuple[list[dict], list[str]]:
    """
    The rows a reader marked in a customers export: [{mark, id, keep, row}].

    The export carries an empty «ادغام» column. Rows given the same mark are
    one customer; a mark ending in `*` («1*») says which row survives. The
    header is looked for, not assumed at a fixed row, because the export has
    a title block above it and a reader may delete or add rows there.
    """
    from openpyxl import load_workbook

    from apps.crm.export import MERGE_ID_LABEL, MERGE_MARK_LABEL

    try:
        wb = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception:
        raise MergeError("فایل اکسل خوانده نشد. همان فایل خروجی مشتریان را بفرستید.")

    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        for n, header in enumerate(rows, 1):
            labels = [str(v).strip() if v is not None else "" for v in header]
            if MERGE_MARK_LABEL in labels and MERGE_ID_LABEL in labels:
                break
            if n > 20:
                break
        else:
            continue
        if MERGE_MARK_LABEL not in labels or MERGE_ID_LABEL not in labels:
            continue
        mark_col = labels.index(MERGE_MARK_LABEL)
        id_col = labels.index(MERGE_ID_LABEL)
        marked, problems = [], []
        for r, values in enumerate(rows, n + 1):
            mark = _mark(values[mark_col] if mark_col < len(values) else None)
            if not mark:
                continue
            keep = mark.endswith("*")
            mark = mark.rstrip("*").strip()
            raw = _mark(values[id_col] if id_col < len(values) else None)
            if not raw.isdigit():
                problems.append(f"ردیف {r}: علامت «{mark}» دارد ولی شناسه‌ی مشتری ندارد.")
                continue
            marked.append({"mark": mark, "id": int(raw), "keep": keep, "row": r})
        return marked, problems

    raise MergeError(
        "ستون‌های «ادغام» و «شناسه» در فایل نیست. خروجی تازه‌ی مشتریان را بگیرید، "
        "ردیف‌های تکراری را در ستون «ادغام» با یک عدد مشترک علامت بزنید و همان را بفرستید."
    )


def merge_marked(marked: list[dict], dataset: str, user=None, dry_run=True) -> dict:
    """
    Merge each group of rows that share a mark — the reader already decided.

    Nothing here guesses whether two rows are the same company; the mark is
    the decision. What it does pick, unless a row is starred, is the survivor:
    `survivor_of`'s order, so the row the invoices hang off stays. Each group
    is its own transaction, so one bad group does not undo the rest, and every
    merge goes through `absorb` — reversible, with a `CustomerMerge` record.
    """
    groups: dict[str, list[dict]] = {}
    for m in marked:
        groups.setdefault(m["mark"], []).append(m)

    found = Customer.objects.filter(
        dataset=dataset, pk__in=[m["id"] for m in marked]
    ).in_bulk()

    def side(c):
        return {"id": c.pk, "code": c.code, "name_fa": c.name_fa}

    out, merged = [], 0
    for mark, items in groups.items():
        entry = {"mark": mark, "keep": None, "absorb": [], "error": ""}
        out.append(entry)
        missing = [m["row"] for m in items if m["id"] not in found]
        rows = list({m["id"]: found[m["id"]] for m in items if m["id"] in found}.values())
        # A row merged earlier counts as the customer it went into.
        rows = list({(c.merged_into or c).pk: (c.merged_into or c) for c in rows}.values())
        if missing:
            entry["error"] = "مشتری ردیف " + "، ".join(map(str, missing)) + " پیدا نشد."
        elif len(rows) < 2:
            entry["error"] = "فقط یک مشتری این علامت را دارد؛ دست‌کم دو ردیف لازم است."
        starred = {m["id"] for m in items if m["keep"]}
        if not entry["error"] and len(starred) > 1:
            entry["error"] = "بیش از یک ردیف با * علامت خورده؛ فقط ماندنی را ستاره بزنید."
        if entry["error"]:
            entry["names"] = [c.name_fa for c in rows]
            continue

        keep = next((c for c in rows if c.pk in starred), None)
        if keep is None:
            keep = rows[0]
            for c in rows[1:]:
                keep, _ = survivor_of(keep, c)
        others = [c for c in rows if c.pk != keep.pk]
        entry["keep"] = side(keep)
        entry["absorb"] = [side(c) for c in others]
        if dry_run:
            continue
        try:
            with transaction.atomic():
                for c in others:
                    absorb(keep, c, user)
            merged += len(others)
        except MergeError as exc:
            entry["error"] = str(exc)

    ready = [g for g in out if not g["error"]]
    return {
        "groups": out,
        "ready": len(ready),
        "rows": sum(len(g["absorb"]) for g in ready),
        "merged": merged,
    }


class ImportTwins:
    """
    For importers: is the customer just created someone the CRM already has?

    The ladder that answered this was removed with the old party importer,
    so دیدار, تیم‌یار and آرپا have each been adding a second «آقای سلطانی»
    beside the first. A suspected twin is queued for review — not merged, and
    not refused: the import still lands, a person decides. Off when
    «تنظیمات تشخیص تکراری» says so. The index is built on first use, so an
    upload that creates nothing costs nothing.
    """

    def __init__(self, dataset: str):
        self.dataset = dataset
        self.index = None
        self.queued = 0

    def _build(self):
        from apps.crm.matching import configured_index

        index = configured_index()
        for pk, name, nid, eco, phone, mobile in Customer.objects.filter(
            dataset=self.dataset, merged_into__isnull=True,
        ).values_list("pk", "name_fa", "national_id", "economic_code", "phone", "mobile"):
            index.add(pk, name=name, nids=(nid, eco), phones=(phone, mobile))
        return index

    def check(self, customer: Customer) -> None:
        from apps.crm.matching import confidence
        from apps.crm.models import CrmConfig

        if not CrmConfig.get().queue_on_import:
            return
        if self.index is None:
            self.index = self._build()
        nids = (customer.national_id, customer.economic_code)
        phones = (customer.phone, customer.mobile)
        match = self.index.find(
            source=ExternalSource.CRM, external_id=f"cust-{customer.pk}",
            name=customer.name_fa, nids=nids, phones=phones,
        )
        if match.found and match.customer_id != customer.pk:
            twin = Customer.objects.filter(pk=match.customer_id).first()
            if twin:
                sure, why = confidence(_as_dict(customer), _as_dict(twin))
                if queue_pair(twin, customer, match.method, sure, why):
                    self.queued += 1
        self.index.add(customer.pk, name=customer.name_fa, nids=nids, phones=phones)


def _as_dict(c) -> dict:
    return {f: getattr(c, f) for f in ("name_fa", "national_id", "economic_code", "mobile", "phone")}


#: A key shared by more rows than this is a head office's id or a shared
#: switchboard, not a duplicate — pairing every branch with every other would
#: bury the real duplicates under hundreds of non-questions.
CROWD = 8


def find_pairs(dataset: str):
    """
    Every suspected duplicate pair in the customer file, with its confidence.

    The whole-file version of `queue_scan`: rows sharing a normalised name,
    id or phone, plus names within the fuzzy threshold among those sharing
    their opening letters. Read-only — the caller decides whether to queue.
    Returns (pairs, crowded), crowded being keys skipped as too common.
    """
    from collections import defaultdict
    from difflib import SequenceMatcher
    from itertools import combinations

    from apps.crm.matching import confidence, place_conflict, typo_apart
    from apps.crm.models import CrmConfig

    cfg = CrmConfig.get()
    rows = list(Customer.objects.filter(dataset=dataset, merged_into__isnull=True).values(
        "pk", "code", "name_key", "nid_key", "mobile_key", "phone_key",
        "name_fa", "national_id", "economic_code", "mobile", "phone", "city",
    ))
    keys = ("name_key", "nid_key", "mobile_key") + (("phone_key",) if cfg.match_phone else ())
    buckets = defaultdict(list)
    for r in rows:
        for k in keys:
            if r[k]:
                buckets[(k, r[k])].append(r)
    pairs, crowded = {}, 0
    for grp in buckets.values():
        if len(grp) > CROWD:
            crowded += 1
            continue
        for a, b in combinations(grp, 2):
            pairs.setdefault(tuple(sorted((a["pk"], b["pk"]))), (a, b))

    threshold = float(cfg.fuzzy_threshold)
    heads = defaultdict(list)
    for r in rows:
        if len(r["name_key"]) >= 4:
            heads[r["name_key"][:4]].append(r)
    for grp in heads.values():
        for a, b in combinations(grp, 2):
            if a["name_key"] == b["name_key"] or place_conflict(a["name_fa"], b["name_fa"]):
                continue
            # Overlap alone pairs every bank in a province with every other;
            # only a typo-shaped difference is worth a reviewer's time.
            if not typo_apart(a["name_key"], b["name_key"]):
                continue
            if SequenceMatcher(None, a["name_key"], b["name_key"]).ratio() >= threshold:
                pairs.setdefault(tuple(sorted((a["pk"], b["pk"]))), (a, b))

    out = []
    for a, b in pairs.values():
        sure, why = confidence(a, b)
        out.append({"a": a, "b": b, "confidence": sure, "reason": why})
    return out, crowded


def _method_for(reason: str) -> str:
    M = CustomerMatchCandidate.Method
    if "نام یکسان" in reason:
        return M.NAME
    if "شناسه" in reason:
        return M.NATIONAL_ID
    if "تلفن" in reason:
        return M.PHONE
    return M.FUZZY


def queue_found(pairs) -> int:
    """File `find_pairs` output in the review queue. Never merges."""
    rows = Customer.objects.in_bulk(
        {p["a"]["pk"] for p in pairs} | {p["b"]["pk"] for p in pairs}
    )
    return sum(
        1 for p in pairs
        if queue_pair(rows[p["a"]["pk"]], rows[p["b"]["pk"]],
                      _method_for(p["reason"]), p["confidence"], p["reason"])
    )


def alternatives(candidate: CustomerMatchCandidate, limit: int = 6):
    """
    Other customers the party might be, for the «ambig» tier.

    That tier means the name matched more than one CRM row — a duplicate the
    دیدار import left behind. The reviewer's job there is to pick which half
    is real, so the halves have to be on screen.
    """
    from apps.crm.matching import name_key

    key = name_key(candidate.external_name)
    if not key:
        return Customer.objects.none()
    ids = [
        c.pk for c in Customer.objects.filter(
            dataset=candidate.customer.dataset, merged_into__isnull=True,
        ).exclude(pk=candidate.customer_id).only("pk", "name_fa")
        if name_key(c.name_fa) == key
    ][:limit]
    return Customer.objects.filter(pk__in=ids)


@transaction.atomic
def delete_customers(queryset, user=None) -> int:
    """
    Delete customers, leaving a tombstone for every source id they carried.

    The only sanctioned way to remove a customer. A bare `.delete()` cascades
    the external refs away, and with them the fact that the party was ever
    imported — so the next accounting load recreates it. Both the list's bulk
    action and the single-row DELETE come through here for that reason.
    """
    from apps.crm.models import DismissedParty

    ids = list(queryset.values_list("pk", flat=True))
    if not ids:
        return 0
    refs = CustomerExternalRef.objects.filter(customer_id__in=ids).values_list(
        "source", "external_id", "external_name"
    )
    DismissedParty.objects.bulk_create(
        [
            DismissedParty(
                source=source, external_id=external_id,
                external_name=name,
                dismissed_by=user if user and user.is_authenticated else None,
            )
            for source, external_id, name in refs
        ],
        ignore_conflicts=True,
    )
    return Customer.objects.filter(pk__in=ids).delete()[0]
