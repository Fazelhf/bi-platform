"""
The weekly price workbook, in and out — in the company's own layout.

The workbook has one sheet per paper weight and رسمی/غیر رسمی: a title in A1,
the date and «فی 02» in row 2, then «ردیف · سایز · متراژ · اجرت برش · اجرت
چاپ · 200 · 50 · 10» with the price columns as formulas over «فی 02». Import
reads it by header text (one sheet has an extra unlabeled column); export
writes the same shape back, formulas included, so the file sent to customers
still recalculates when someone edits «فی 02» in Excel.
"""
from __future__ import annotations

import re
from decimal import Decimal
from io import BytesIO

from django.db import transaction

from apps.core import jalali
from apps.sales2.models import PriceSheet, PriceSheetRow

_NORM = str.maketrans({"ي": "ی", "ك": "ک", "‌": " "})


def _norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "").translate(_NORM)).strip()


def read_workbook(fileobj) -> list[dict]:
    """Every sheet with 48/55 in its name → base, sizes and fees."""
    from openpyxl import load_workbook

    wb = load_workbook(fileobj)
    out = []
    for ws in wb.worksheets:
        title = _norm(ws.title)
        m = re.search(r"(48|55)", title)
        if not m:
            continue
        base = None
        for c in ws[2]:
            if isinstance(c.value, (int, float)) and c.value > 1000:
                base = Decimal(str(c.value))
                break
        hdr = next((r for r in ws.iter_rows(min_row=1, max_row=6)
                    if any(_norm(c.value) == "سایز" for c in r)), None)
        if base is None or hdr is None:
            continue
        cols: dict[str, str] = {}
        for c in hdr:
            if c.value is not None and _norm(c.value) and _norm(c.value) not in cols:
                cols[_norm(c.value)] = c.column_letter
        rows = []
        for r in range(hdr[0].row + 1, ws.max_row + 1):
            w, ln = ws[f"{cols['سایز']}{r}"].value, ws[f"{cols['متراژ']}{r}"].value
            if not isinstance(w, (int, float)) or not isinstance(ln, (int, float)):
                continue
            formula = str(ws[f"{cols['200']}{r}"].value or "") if "200" in cols else ""
            wm = re.search(r"\*\s*(\d+(?:\.\d+)?)%\)", formula)
            note = " ".join(_norm(c.value) for c in ws[r]
                            if isinstance(c.value, str) and not c.value.startswith("="))
            rows.append({
                "width_mm": int(w), "length_m": int(ln),
                "cut_fee_rial": Decimal(str(ws[f"{cols['اجرت برش']}{r}"].value or 0)),
                "print_fee_rial": Decimal(str(ws[f"{cols.get('اجرت چاپ', cols['اجرت برش'])}{r}"].value or 0))
                if "اجرت چاپ" in cols else Decimal(0),
                "waste_pct": Decimal(wm.group(1)) if wm else None,
                "note": note[:60],
            })
        out.append({"grammage": int(m.group(1)), "is_official": "غیر" not in title,
                    "base_fi_rial": base, "rows": rows})
    return out


@transaction.atomic
def save_month(sheets: list[dict], jy: int, jm: int) -> list[PriceSheet]:
    """The month's sheets, replacing that month's own (earlier months untouched)."""
    saved = []
    for sh in sheets:
        PriceSheet.objects.filter(grammage=sh["grammage"], is_official=sh["is_official"],
                                  jalali_year=jy, jalali_month=jm).delete()
        obj = PriceSheet.objects.create(
            name=f"{'رسمی' if sh['is_official'] else 'غیر رسمی'} {sh['grammage']}",
            grammage=sh["grammage"], is_official=sh["is_official"],
            base_fi_rial=sh["base_fi_rial"], jalali_year=jy, jalali_month=jm,
        )
        default_waste = obj.waste_pct
        for i, r in enumerate(sh["rows"]):
            waste = r.get("waste_pct")
            PriceSheetRow.objects.create(
                sheet=obj, sort_order=i, width_mm=r["width_mm"], length_m=r["length_m"],
                cut_fee_rial=r["cut_fee_rial"], print_fee_rial=r["print_fee_rial"],
                waste_pct=None if waste is None or waste == default_waste else waste,
                note=r.get("note", ""),
            )
        saved.append(obj)
    return saved


def copy_to_month(sheets: list[PriceSheet], jy: int, jm: int) -> list[PriceSheet]:
    """Carry the sheets in force into a month of its own, ready to edit."""
    data = [{
        "grammage": s.grammage, "is_official": s.is_official, "base_fi_rial": s.base_fi_rial,
        "rows": [{"width_mm": r.width_mm, "length_m": r.length_m, "cut_fee_rial": r.cut_fee_rial,
                  "print_fee_rial": r.print_fee_rial, "waste_pct": r.waste_pct, "note": r.note}
                 for r in s.rows.all()],
    } for s in sheets]
    created = save_month(data, jy, jm)
    for new, old in zip(created, sheets):
        new.waste_pct, new.qty_tiers = old.waste_pct, old.qty_tiers
        new.save(update_fields=["waste_pct", "qty_tiers"])
    return created


def export_workbook(sheets: list[PriceSheet], jy: int, jm: int) -> bytes:
    """The company's layout: A1 title, row 2 date and «فی 02», formula prices."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    wb.remove(wb.active)
    order = sorted(sheets, key=lambda s: (-s.is_official, -s.grammage))
    for sh in order:
        ws = wb.create_sheet(sh.name[:30])
        ws.sheet_view.rightToLeft = True
        tiers = sorted(sh.qty_tiers or [], key=lambda t: -t["min_qty"])
        ws["A1"] = sh.name
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = f"{jy}/{jm:02d}/01"
        ws["C2"] = "تاریخ:"
        ws["D2"] = int(sh.base_fi_rial)
        ws["E2"] = f"فی 02 {'رسمی' if sh.is_official else 'غیر رسمی'}:"
        headers = ["ردیف", "سایز", "متراژ", "اجرت برش", "اجرت چاپ"] + \
                  [str(t["min_qty"]) if t["min_qty"] else "کمتر" for t in tiers] + ["توضیح"]
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=i, value=h)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
        for n, r in enumerate(sh.rows.all(), start=1):
            row = 3 + n
            waste = float(r.waste_pct if r.waste_pct is not None else sh.waste_pct)
            ws.cell(row=row, column=1, value=n)
            ws.cell(row=row, column=2, value=r.width_mm)
            ws.cell(row=row, column=3, value=r.length_m)
            ws.cell(row=row, column=4, value=int(r.cut_fee_rial))
            ws.cell(row=row, column=5, value=int(r.print_fee_rial))
            base = f"(((B{row}*C{row}*$D$2)*{waste:g}%)/1000)+D{row}"
            for k, t in enumerate(tiers):
                col = 6 + k
                formula = f"={base}" if not t["pct"] else f"=({base})*{100 + t['pct']:g}%"
                c = ws.cell(row=row, column=col, value=formula)
                c.number_format = "#,##0"
            ws.cell(row=row, column=6 + len(tiers), value=r.note or None)
        for col, width in zip("ABCDEFGHI", (6, 8, 8, 12, 12, 14, 14, 14, 14)):
            ws.column_dimensions[col].width = width
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def month_of(on) -> tuple[int, int]:
    y, m, _ = jalali.from_gregorian(on)
    return y, m
