"""
مطالبات طبق دفتر آرپا — receivables from accounting's own ledger.

Finance's «گزارش وصولی» is typed by hand each week from آرپا's «دفتر حساب»
of every customer: what is still open, aged by the month of the فاکتور it
belongs to. This module does that from the ledger itself:

1. **Read** the export as آرپا writes it (by header text, not column letter):
   شماره سند، شماره ثابت سند، تاریخ سند، مبلغ بدهکار، مبلغ بستانکار، مانده،
   وضعیت، شرح ردیف سند، تیک. The first row, «موجودی قبل از دوره», is the
   opening balance.

2. **Classify** each row from its description: «بابت فاکتور فروش 168» is
   that فاکتور; «دریافت 1 فقره چک بابت تسویه فاکتور فروش168» is a دریافت by
   cheque that settles فاکتور 168.

3. **Settle**: a دریافت pays the فاکتور it names first. What names nothing,
   names a فاکتور already paid, or is left over, pays the oldest open item
   (FIFO). A دریافت naming a فاکتور that is not in the file pays the
   opening balance — the فاکتور is from before the period — or, with no
   opening left, waits as پیش دریافت.

4. **Reconcile**: the balance is recomputed row by row and checked against
   آرپا's own «مانده» on every row; any disagreement is reported, never
   smoothed over. A دریافت split over several فاکتورها under one سند is
   checked against the «مبلغ …» its description states.

5. **Age**: what is left of each item is put in the column of its month, as
   the workbook does — but per فاکتور, so a customer whose old فاکتور is half
   paid and whose new one is untouched shows in two columns, not one.

Two books: دفتر اصلی and دفتر دوم (رول مارکت) are separate آرپا files, so
every account is one customer in one book (`LedgerAccount`), and every
figure is per book; «هر دو» adds them per customer.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.utils import timezone

from apps.core import jalali
from apps.crm.matching import fold
from apps.sales2.models import ZERO, LedgerAccount, LedgerBook, LedgerEntry

Kind = LedgerEntry.Kind
Method = LedgerEntry.Method

BOOK_LABEL = dict(LedgerBook.choices)
BOOK_SHORT = {LedgerBook.MAIN: "اصلی", LedgerBook.SECOND: "دوم"}


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------
class LedgerFileError(ValueError):
    pass


#: header text (folded) → field. First match wins, so «شماره ثابت سند» is
#: tried before «شماره سند».
HEADERS = (
    ("fixed_no", ("شماره ثابت",)),
    ("doc_no", ("شماره سند",)),
    ("date", ("تاریخ",)),
    ("debit", ("بدهکار",)),
    ("credit", ("بستانکار",)),
    ("balance", ("مانده",)),
    ("status", ("وضعیت", "تشخیص")),
    ("description", ("شرح",)),
    ("ticked", ("تیک",)),
)

INVOICE_RE = re.compile(r"فاکتور\s*(?:فروش)?\s*(?:شماره)?\s*(\d+)")
RETURN_RE = re.compile(r"برگشت\s*از\s*فروش|مرجوع")
CHEQUE_COUNT_RE = re.compile(r"(\d+)\s*فقره\s*چک")
STATED_RE = re.compile(r"مبلغ\s*([\d.,٬]+)")


@dataclass
class Row:
    row_no: int
    doc_no: str
    fixed_no: str
    date: date | None
    debit: Decimal
    credit: Decimal
    balance: Decimal | None
    description: str
    ticked: bool
    kind: str = Kind.OTHER
    method: str = ""
    invoice_no: str = ""
    cheque_count: int = 0


@dataclass
class Parsed:
    rows: list[Row]
    opening: Decimal
    sheet: str

    @property
    def first_date(self):
        return min((r.date for r in self.rows if r.date), default=None)

    @property
    def last_date(self):
        return max((r.date for r in self.rows if r.date), default=None)


def _money(v) -> Decimal:
    if v is None or v == "":
        return ZERO
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v)).quantize(Decimal(1))
    text = fold(v).replace(",", "").replace("٬", "").replace(" ", "")
    neg = text.startswith("(") and text.endswith(")") or text.startswith("-")
    text = text.strip("()-")
    try:
        d = Decimal(text or 0).quantize(Decimal(1))
    except InvalidOperation:
        return ZERO
    return -d if neg else d


def parse_jalali(v) -> date | None:
    """«1405/02/06» (or a date cell) → a Gregorian date."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date):
        # Excel holds a real date only when someone retyped the cell; a year
        # under 1500 is a Jalali date stored as one.
        if v.year < 1500:
            return jalali.to_gregorian(v.year, v.month, v.day)
        return v
    m = re.search(r"(\d{4})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{1,2})", fold(v))
    if not m:
        return None
    y, mo, d = (int(x) for x in m.groups())
    if not (1 <= mo <= 12 and 1 <= d <= 31):
        return None
    return jalali.to_gregorian(y, mo, d)


def _signed(balance, status) -> Decimal | None:
    if balance is None or balance == "":
        return None
    amount = abs(_money(balance))
    s = fold(status)
    if s.startswith("بس"):
        return -amount
    return amount


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return fold(v)


def classify(description: str, debit: Decimal, credit: Decimal) -> tuple[str, str, str, int]:
    """(kind, method, invoice_no, cheque_count) read off a row's description."""
    text = fold(description)
    m = INVOICE_RE.search(text)
    invoice_no = m.group(1) if m else ""
    if "قبل از دوره" in text:
        return Kind.OPENING, "", "", 0
    cheques = CHEQUE_COUNT_RE.search(text)
    count = int(cheques.group(1)) if cheques else (1 if "چک" in text else 0)
    if RETURN_RE.search(text):
        return Kind.RETURN, "", invoice_no, 0
    if credit > 0 and ("دریافت" in text or "وصول" in text or "واریز" in text):
        if "چک" in text:
            method = Method.CHEQUE
        elif "نقد" in text:
            method = Method.CASH
        else:
            method = Method.TRANSFER
        return Kind.RECEIPT, method, invoice_no, count if method == Method.CHEQUE else 0
    if debit > 0 and invoice_no and "دریافت" not in text:
        return Kind.INVOICE, "", invoice_no, 0
    return Kind.OTHER, "", invoice_no, 0


def read_file(fileobj) -> Parsed:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(fileobj, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 — any unreadable file is the same answer
        raise LedgerFileError("فایل اکسل خوانده نشد.") from exc
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for h, header in enumerate(rows[:15]):
            names = [fold(c) for c in header]
            cols: dict[str, int] = {}
            for i, name in enumerate(names):
                for key, needles in HEADERS:
                    if key not in cols and any(n in name for n in needles):
                        cols[key] = i
                        break
            if {"debit", "credit", "date"} <= cols.keys():
                return _rows(rows[h + 1:], cols, ws.title)
    raise LedgerFileError(
        "ستون‌های «تاریخ سند»، «مبلغ بدهکار» و «مبلغ بستانکار» پیدا نشد؛ "
        "خروجی «دفتر حساب» یک طرف حساب را از آرپا بدهید."
    )


def _rows(raw, cols, sheet) -> Parsed:
    def get(r, key):
        i = cols.get(key)
        return r[i] if i is not None and i < len(r) else None

    out: list[Row] = []
    opening = ZERO
    for r in raw:
        desc = _text(get(r, "description"))
        debit, credit = _money(get(r, "debit")), _money(get(r, "credit"))
        when = parse_jalali(get(r, "date"))
        balance = _signed(get(r, "balance"), get(r, "status"))
        if "قبل از دوره" in desc:
            # Its own مانده is the opening; its بدهکار/بستانکار are zero.
            opening = balance if balance is not None else debit - credit
            continue
        if when is None:
            # The totals row (=SUM…) and blank rows have no date.
            continue
        if debit == 0 and credit == 0:
            continue
        kind, method, inv, count = classify(desc, debit, credit)
        out.append(Row(
            row_no=len(out) + 1, doc_no=_text(get(r, "doc_no")), fixed_no=_text(get(r, "fixed_no")),
            date=when, debit=debit, credit=credit, balance=balance, description=desc[:500],
            ticked=fold(get(r, "ticked")).lower() in ("true", "1", "بله", "✓"),
            kind=kind, method=method, invoice_no=inv, cheque_count=count,
        ))
    if not out and not opening:
        raise LedgerFileError("ردیفی با تاریخ و مبلغ در فایل نیست.")
    return Parsed(rows=out, opening=opening, sheet=sheet)


# ---------------------------------------------------------------------------
# Saving
# ---------------------------------------------------------------------------
@transaction.atomic
def save(parsed: Parsed, customer, book: str, user=None, filename: str = "",
         party_code: str = "") -> LedgerAccount:
    """
    Replace the customer's ledger in `book` with the file's. The export is
    the whole account for its period, so a fresh one supersedes the last;
    nothing is merged row by row.
    """
    acc, _ = LedgerAccount.objects.get_or_create(book=book, customer=customer)
    acc.entries.all().delete()
    LedgerEntry.objects.bulk_create([
        LedgerEntry(
            account=acc, row_no=r.row_no, doc_no=r.doc_no[:20], fixed_no=r.fixed_no[:20], date=r.date,
            debit_rial=r.debit, credit_rial=r.credit, file_balance_rial=r.balance,
            description=r.description, ticked=r.ticked, kind=r.kind, method=r.method,
            invoice_no=r.invoice_no[:20], cheque_count=r.cheque_count,
        )
        for r in parsed.rows
    ])
    acc.opening_rial = parsed.opening
    acc.first_date, acc.last_date = parsed.first_date, parsed.last_date
    acc.source_file = filename[:200]
    acc.imported_at, acc.imported_by = timezone.now(), user if getattr(user, "is_authenticated", False) else None
    if party_code:
        acc.party_code = party_code[:32]
    elif not acc.party_code and book == LedgerBook.MAIN:
        acc.party_code = arpa_code(customer)
    acc.save()
    return acc


def arpa_code(customer) -> str:
    from apps.crm.models import CustomerExternalRef, ExternalSource

    ref = CustomerExternalRef.objects.filter(customer=customer, source=ExternalSource.ARPA) \
        .order_by("pk").values_list("external_id", flat=True).first()
    return ref or ""


def suggest_customers(filename: str, limit: int = 5) -> list[dict]:
    """Customers whose name holds the file's name («اجاقلو.xlsx» → اجاقلو)."""
    from apps.crm.models import Customer

    stem = re.sub(r"\.(xlsx|xlsm|xls)$", "", fold(filename), flags=re.I)
    stem = re.sub(r"\(\d+\)|\d{6,}|دفتر|حساب|معین|تفصیلی|آقای|خانم|شرکت|[_\-]", " ", stem)
    words = [w for w in stem.split() if len(w) > 1]
    if not words:
        return []
    qs = Customer.objects.all()
    for w in words:
        qs = qs.filter(name_fa__icontains=w)
    found = list(qs.order_by("name_fa")[:limit])
    if not found:
        longest = max(words, key=len)
        found = list(Customer.objects.filter(name_fa__icontains=longest).order_by("name_fa")[:limit])
    return [{"id": c.id, "name": c.name_fa} for c in found]


# ---------------------------------------------------------------------------
# Settling: which دریافت paid which فاکتور
# ---------------------------------------------------------------------------
@dataclass
class Payment:
    date: date
    doc_no: str
    amount: Decimal
    method: str
    how: str          # named · fifo · opening
    row_no: int


@dataclass
class Item:
    """Something owed: a فاکتور, the opening balance, or another بدهکار row."""
    key: str
    kind: str
    number: str
    date: date
    amount: Decimal
    row_no: int
    doc_no: str = ""
    description: str = ""
    payments: list[Payment] = field(default_factory=list)
    returned: Decimal = ZERO

    @property
    def paid(self) -> Decimal:
        return sum((p.amount for p in self.payments), ZERO)

    @property
    def remaining(self) -> Decimal:
        return self.amount - self.paid

    @property
    def settled_on(self) -> date | None:
        if self.remaining > 0 or not self.payments:
            return None
        return max(p.date for p in self.payments)


@dataclass
class Credit:
    date: date
    doc_no: str
    fixed_no: str
    amount: Decimal
    kind: str
    method: str
    invoice_no: str
    row_no: int
    description: str
    cheque_count: int = 0
    left: Decimal = ZERO
    #: names a فاکتور not booked (yet): waits as پیش دریافت, never FIFO'd
    earmarked: bool = False
    allocations: list[tuple[str, Decimal, str]] = field(default_factory=list)  # (item key, amount, how)


@dataclass
class Issue:
    level: str        # error · warn · info
    message: str
    row_no: int | None = None


@dataclass
class Analysis:
    account: LedgerAccount
    as_of: date | None
    items: list[Item]
    credits: list[Credit]
    entries: list[LedgerEntry]
    running: list[tuple[LedgerEntry, Decimal]]   # each row and the computed مانده after it
    prepayment: Decimal
    balance: Decimal
    file_balance: Decimal | None
    issues: list[Issue]

    @property
    def open_items(self) -> list[Item]:
        return [i for i in self.items if i.remaining > 0]

    @property
    def oldest_open(self) -> date | None:
        return min((i.date for i in self.open_items), default=None)

    @property
    def reconciled(self) -> bool:
        return not any(i.level == "error" for i in self.issues)


def _stated_amount(text: str) -> Decimal | None:
    m = STATED_RE.search(fold(text))
    if not m:
        return None
    digits = re.sub(r"\D", "", m.group(1))
    return Decimal(digits) if digits else None


def analyse(account: LedgerAccount, as_of: date | None = None, entries=None) -> Analysis:
    if entries is None:
        entries = list(account.entries.all())
    entries = [e for e in entries if as_of is None or e.date <= as_of]
    entries.sort(key=lambda e: e.row_no)
    issues: list[Issue] = []

    # ---- reconcile against آرپا's own مانده, row by row -----------------
    running: list[tuple[LedgerEntry, Decimal]] = []
    bal = account.opening_rial
    first_bad = None
    for e in entries:
        bal += e.debit_rial - e.credit_rial
        running.append((e, bal))
        if first_bad is None and e.file_balance_rial is not None and e.file_balance_rial != bal:
            first_bad = e
    if first_bad is not None:
        issues.append(Issue("error", f"مانده‌ی محاسبه‌شده در ردیف {first_bad.row_no} (سند {first_bad.doc_no}) "
                                     f"با «مانده»ی آرپا نمی‌خواند؛ دفتر کامل نیست یا ردیفی جا افتاده.",
                            first_bad.row_no))
    file_balance = next((e.file_balance_rial for e in reversed(entries) if e.file_balance_rial is not None),
                        account.opening_rial if not entries else None)

    # ---- what is owed --------------------------------------------------
    items: dict[str, Item] = {}
    order: list[Item] = []
    first_day = account.first_date or (entries[0].date if entries else as_of or date.today())
    if account.opening_rial > 0:
        it = Item("opening", Kind.OPENING, "", first_day, account.opening_rial, 0,
                  description="موجودی قبل از دوره")
        items[it.key] = it
        order.append(it)
    credits: list[Credit] = []
    if account.opening_rial < 0:
        credits.append(Credit(first_day, "", "", -account.opening_rial, Kind.OPENING, "", "", 0,
                              "موجودی بستانکار قبل از دوره", left=-account.opening_rial))
    for e in entries:
        if e.debit_rial > 0:
            if e.kind == Kind.INVOICE and e.invoice_no:
                key = f"inv:{e.invoice_no}"
                if key in items:
                    # One فاکتور booked over two rows (or a correction): one item.
                    items[key].amount += e.debit_rial
                    continue
                it = Item(key, Kind.INVOICE, e.invoice_no, e.date, e.debit_rial, e.row_no,
                          doc_no=e.doc_no, description=e.description)
            else:
                it = Item(f"row:{e.row_no}", Kind.OTHER, e.invoice_no, e.date, e.debit_rial, e.row_no,
                          doc_no=e.doc_no, description=e.description)
            items[it.key] = it
            order.append(it)
        if e.credit_rial > 0:
            credits.append(Credit(e.date, e.doc_no, e.fixed_no, e.credit_rial, e.kind, e.method,
                                  e.invoice_no, e.row_no, e.description, e.cheque_count, left=e.credit_rial))

    def pay(c: Credit, it: Item, amount: Decimal, how: str):
        if amount <= 0:
            return
        it.payments.append(Payment(c.date, c.doc_no, amount, c.method or c.kind, how, c.row_no))
        if c.kind == Kind.RETURN:
            it.returned += amount
        c.left -= amount
        c.allocations.append((it.key, amount, how))

    # 1. each credit pays the فاکتور it names
    for c in credits:
        if not c.invoice_no:
            continue
        it = items.get(f"inv:{c.invoice_no}")
        if it is None:
            opening = items.get("opening")
            if opening is not None and opening.remaining > 0:
                pay(c, opening, min(c.left, opening.remaining), "opening")
            else:
                c.earmarked = True
                issues.append(Issue("info", f"سند {c.doc_no}: فاکتور {c.invoice_no} تا این تاریخ در دفتر نیست؛ "
                                            f"مبلغ پیش‌دریافت حساب شد.", c.row_no))
            continue
        take = min(c.left, it.remaining)
        if take < c.left:
            issues.append(Issue("warn", f"سند {c.doc_no}: بیش از مانده‌ی فاکتور {c.invoice_no} دریافت شده؛ "
                                        f"{int(c.left - take):,} ریال به قدیمی‌ترین بدهی رفت.", c.row_no))
        pay(c, it, take, "named")

    # 2. the rest pays the oldest open item first
    pending = sorted((c for c in credits if c.left > 0 and not c.earmarked), key=lambda c: (c.date, c.row_no))
    for c in pending:
        for it in sorted(order, key=lambda i: (i.date, i.row_no)):
            if c.left <= 0:
                break
            if it.remaining > 0:
                pay(c, it, min(c.left, it.remaining), "fifo")

    prepayment = sum((c.left for c in credits), ZERO)
    balance = account.opening_rial + sum((e.debit_rial - e.credit_rial for e in entries), ZERO)
    owed = sum((i.remaining for i in order), ZERO)
    if owed - prepayment != balance:  # pragma: no cover — an invariant, not a branch
        issues.append(Issue("error", "جمع اقلام باز با مانده‌ی حساب نمی‌خواند."))
    if file_balance is not None and file_balance != balance:
        issues.append(Issue("error", f"مانده‌ی پایانی محاسبه‌شده {int(balance):,} و طبق آرپا {int(file_balance):,} است."))

    # ---- a دریافت over several فاکتورها under one سند: check the stated sum
    by_doc: dict[tuple, list[Credit]] = {}
    for c in credits:
        if c.kind == Kind.RECEIPT and c.doc_no:
            by_doc.setdefault((c.doc_no, c.date), []).append(c)
    for (doc, _), group in by_doc.items():
        stated = _stated_amount(group[0].description)
        total = sum((c.amount for c in group), ZERO)
        # «2 فقره چک … به مبلغ X» states each cheque's amount, not the sum.
        count = max(1, group[0].cheque_count)
        if stated is not None and total not in (stated, stated * count) and len(group) > 1:
            issues.append(Issue("warn", f"سند {doc}: شرح «مبلغ {int(stated):,}» می‌گوید ولی جمع ردیف‌ها "
                                        f"{int(total):,} است.", group[0].row_no))

    return Analysis(account=account, as_of=as_of, items=order, credits=credits, entries=entries,
                    running=running, prepayment=prepayment, balance=balance,
                    file_balance=file_balance, issues=issues)


# ---------------------------------------------------------------------------
# Invoices as آرپا's sales list has them — due date and salesperson
# ---------------------------------------------------------------------------
def sales_invoices(customer_ids) -> dict[tuple[int, str], dict]:
    """(customer, شماره برگه) → due date, total and salesperson from CRM's آرپا copy."""
    from apps.crm.models import SalesInvoice

    out = {}
    for inv in SalesInvoice.objects.filter(customer_id__in=customer_ids, kind=SalesInvoice.Kind.SALE) \
            .select_related("owner").order_by("issued_at"):
        out[(inv.customer_id, str(inv.number).strip())] = {
            "due_date": inv.due_date, "total": inv.total_rial, "issued_at": inv.issued_at,
            "owner": inv.owner.full_name_fa if inv.owner_id else "",
        }
    return out


# ---------------------------------------------------------------------------
# The collections report («گزارش وصولی»)
# ---------------------------------------------------------------------------
def month_columns(as_of: date, months: int) -> list[dict]:
    """The aging columns: «ماه‌های قبل», then `months` months ending with as_of's."""
    y, m, _ = jalali.from_gregorian(as_of)
    keys = []
    for _ in range(months):
        keys.append((y, m))
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    keys.reverse()
    cols = [{"key": "before", "label": "ماه های قبل", "year": None, "month": None}]
    for yy, mm in keys:
        label = jalali.MONTHS_FA[mm] if yy == keys[-1][0] else f"{jalali.MONTHS_FA[mm]} {yy}"
        cols.append({"key": f"{yy}-{mm:02d}", "label": label, "year": yy, "month": mm})
    return cols


def bucket_of(d: date, cols: list[dict]) -> str:
    y, m, _ = jalali.from_gregorian(d)
    key = f"{y}-{m:02d}"
    return key if any(c["key"] == key for c in cols) else "before"


def jdate(d: date | None) -> str:
    if not d:
        return ""
    y, m, dd = jalali.from_gregorian(d)
    return f"{y}/{m:02d}/{dd:02d}"


def accounts(book: str):
    qs = LedgerAccount.objects.select_related("customer__owner").prefetch_related("entries")
    if book in (LedgerBook.MAIN, LedgerBook.SECOND):
        qs = qs.filter(book=book)
    return qs


def report(book: str, as_of: date, months: int = 11) -> dict:
    """
    One row per account (or per customer, for «all»): the workbook's columns
    — کد تفصیل، نام، تاریخ (oldest open), the month buckets, پیش دریافت,
    مانده بدهکاران — plus what the workbook could not say: آرپا's own
    balance, whether the two agree, and the hand-typed figure beside it.
    """
    cols = month_columns(as_of, months)
    rows: dict = {}
    for acc in accounts(book):
        a = analyse(acc, as_of)
        key = acc.customer_id if book == "all" else acc.pk
        row = rows.get(key)
        if row is None:
            owner = acc.customer.owner.full_name_fa if acc.customer.owner_id else ""
            row = rows[key] = {
                "id": acc.pk, "customer": acc.customer_id, "name": acc.customer.name_fa,
                "book": acc.book if book != "all" else "all",
                "code": acc.party_code, "codes": {}, "accounts": [],
                "buckets": {c["key"]: ZERO for c in cols},
                "prepayment": ZERO, "balance": ZERO, "file_balance": ZERO,
                "by_book": {}, "oldest_open": None,
                "manual_balance": None, "manual_as_of": None,
                "note": acc.note, "marketer": acc.marketer or owner,
                "last_date": None, "entries": 0, "open_items": 0,
                "issues": [], "reconciled": True, "has_ledger": False,
            }
        row["accounts"].append(acc.pk)
        row["codes"][acc.book] = acc.party_code
        if not row["code"]:
            row["code"] = acc.party_code
        if acc.note and acc.note not in row["note"]:
            row["note"] = f"{row['note']} · {acc.note}" if row["note"] else acc.note
        for it in a.open_items:
            row["buckets"][bucket_of(it.date, cols)] += it.remaining
        row["prepayment"] += a.prepayment
        row["balance"] += a.balance
        row["file_balance"] += a.file_balance or ZERO
        row["by_book"][acc.book] = a.balance
        row["open_items"] += len(a.open_items)
        if a.oldest_open and (row["oldest_open"] is None or a.oldest_open < row["oldest_open"]):
            row["oldest_open"] = a.oldest_open
        if acc.last_date and (row["last_date"] is None or acc.last_date > row["last_date"]):
            row["last_date"] = acc.last_date
        has = acc.imported_at is not None
        row["has_ledger"] = row["has_ledger"] or has
        row["entries"] += len(a.entries)
        row["issues"] += [{"level": i.level, "message": i.message, "book": acc.book} for i in a.issues]
        row["reconciled"] = row["reconciled"] and a.reconciled
        if acc.manual_balance_rial is not None:
            row["manual_balance"] = (row["manual_balance"] or ZERO) + acc.manual_balance_rial
            row["manual_as_of"] = max(filter(None, [row["manual_as_of"], acc.manual_as_of]), default=None)
    out = sorted(rows.values(), key=lambda r: (-r["balance"], r["name"]))
    return {"columns": cols, "rows": out, "as_of": as_of, "book": book}


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="BFBFBF")
    return {
        "head_fill": PatternFill("solid", fgColor="1F3864"),
        "head_font": Font(bold=True, color="FFFFFF", size=11),
        "title_font": Font(bold=True, size=14),
        "total_fill": PatternFill("solid", fgColor="D9E1F2"),
        "bad_fill": PatternFill("solid", fgColor="FDE9E7"),
        "ok_fill": PatternFill("solid", fgColor="E2F0D9"),
        "bold": Font(bold=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "right": Alignment(horizontal="right", vertical="center"),
    }


RIAL = "#,##0;[Red]-#,##0"
#: A balance as آرپا prints it — the amount with «بد» or «بس» — still a number, so SUM works.
BALANCE = '#,##0 "بد";[Color10]#,##0 "بس";0'


def _table(ws, headers, body, money_cols=(), widths=None, start=1, total_label=None, sum_cols=(), balance_cols=()):
    """Header + rows, styled; optional totals row with real SUM formulas."""
    from openpyxl.utils import get_column_letter

    st = _styles()
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start, column=c, value=h)
        cell.fill, cell.font, cell.border, cell.alignment = st["head_fill"], st["head_font"], st["border"], st["center"]
    for r, values in enumerate(body, start + 1):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = st["border"]
            if c - 1 in money_cols or c - 1 in balance_cols:
                # Centred under its header, so the column reads straight down.
                cell.number_format = BALANCE if c - 1 in balance_cols else RIAL
                cell.alignment = st["center"]
    end = start + len(body)
    if total_label is not None and body:
        tr = end + 1
        ws.cell(row=tr, column=1, value=total_label)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=tr, column=c)
            cell.fill, cell.font, cell.border = st["total_fill"], st["bold"], st["border"]
            if c - 1 in sum_cols:
                L = get_column_letter(c)
                cell.value = f"=SUM({L}{start + 1}:{L}{end})"
                cell.number_format = BALANCE if c - 1 in balance_cols else RIAL
                cell.alignment = st["center"]
    for c, h in enumerate(headers, 1):
        w = (widths or {}).get(c - 1)
        if w is None:
            longest = max([len(str(h))] + [len(f"{v:,}") if isinstance(v, (int, Decimal)) else len(str(v or ""))
                                           for v in (row[c - 1] for row in body)]) if body else len(str(h))
            w = min(48, max(10, longest + 3))
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.sheet_view.rightToLeft = True
    return end


def _n(v) -> int | None:
    """A Rial figure for a cell: an int, or empty for zero (as the workbook leaves it)."""
    if v is None:
        return None
    v = int(v)
    return v or None


def _collections_sheet(ws, data: dict, title: str):
    """The «گزارش وصولی» sheet itself, column for column, with live formulas."""
    from openpyxl.utils import get_column_letter

    st = _styles()
    cols = data["columns"]
    headers = (["ردیف", "کد تفصیل", "نام طرف حساب", "تاریخ "] + [c["label"] for c in cols]
               + ["پیش دریافت", "مانده بدهکاران", "جمع", "توضیحات", "بازاریاب",
                  "مانده طبق آرپا", "کنترل با آرپا", "طبق گزارش دستی", "اختلاف با گزارش دستی",
                  "آخرین ردیف دفتر", "فاکتور باز"])
    first_bucket = 5
    last_bucket = first_bucket + len(cols) - 1
    c_pre = last_bucket + 1
    c_bal, c_sum = c_pre + 1, c_pre + 2
    c_arpa, c_check, c_manual, c_diff = c_pre + 5, c_pre + 6, c_pre + 7, c_pre + 8
    L = get_column_letter
    body = []
    for i, r in enumerate(data["rows"], 1):
        excel_row = i + 2
        body.append(
            [i, r["code"] or "", r["name"], jdate(r["oldest_open"])]
            + [_n(r["buckets"][c["key"]]) for c in cols]
            + [_n(r["prepayment"]),
               f"=SUM({L(first_bucket)}{excel_row}:{L(last_bucket)}{excel_row})-N({L(c_pre)}{excel_row})",
               f"={L(c_bal)}{excel_row}",
               r["note"], r["marketer"],
               int(r["file_balance"]) if r["has_ledger"] else None,
               ("✓" if r["reconciled"] else "✗ " + "؛ ".join(x["message"] for x in r["issues"] if x["level"] == "error"))
               if r["has_ledger"] else "دفتر وارد نشده",
               _n(r["manual_balance"]),
               f"=N({L(c_bal)}{excel_row})-N({L(c_manual)}{excel_row})" if r["manual_balance"] is not None else None,
               jdate(r["last_date"]), r["open_items"] or None]
        )
    balances = {c_bal - 1, c_sum - 1, c_arpa - 1, c_manual - 1}
    money = (set(range(first_bucket - 1, c_sum)) | {c_diff - 1}) - balances
    widths = {0: 6, 1: 11, 2: 34, 3: 12, c_pre + 2: 30, c_pre + 3: 18, c_check - 1: 28}
    for k in range(first_bucket - 1, c_sum):
        widths[k] = 16
    for k in (c_arpa - 1, c_manual - 1, c_diff - 1):
        widths[k] = 17
    end = _table(ws, headers, body, money_cols=money, widths=widths, start=2, total_label="مجموع",
                 sum_cols=money | balances, balance_cols=balances)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=c_sum)
    t = ws.cell(row=1, column=1, value=title)
    t.font, t.alignment = st["title_font"], st["center"]
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = ws.cell(row=3, column=4)
    for i, r in enumerate(data["rows"], 1):
        if r["has_ledger"]:
            ws.cell(row=i + 2, column=c_check).fill = st["ok_fill"] if r["reconciled"] else st["bad_fill"]
    if data["rows"]:
        ws.auto_filter.ref = f"A2:{L(len(headers))}{end}"


def _detail_rows(analyses: list[Analysis], inv_info: dict, as_of: date):
    inv_rows, rec_rows, issue_rows = [], [], []
    for a in analyses:
        acc = a.account
        who = [BOOK_SHORT[acc.book], acc.party_code, acc.customer.name_fa]
        names = {i.key: (i.number or ("مانده قبل از دوره" if i.kind == Kind.OPENING else f"ردیف {i.row_no}"))
                 for i in a.items}
        for it in a.items:
            info = inv_info.get((acc.customer_id, it.number)) if it.kind == Kind.INVOICE else None
            due = info["due_date"] if info else None
            open_ = it.remaining > 0
            overdue = (as_of - due).days if open_ and due and as_of > due else None
            settled = it.settled_on
            inv_rows.append(who + [
                "فاکتور" if it.kind == Kind.INVOICE else ("مانده قبل از دوره" if it.kind == Kind.OPENING else "سایر بدهکار"),
                it.number or "", jdate(it.date), it.doc_no, int(it.amount), _n(it.returned),
                _n(it.paid - it.returned), int(it.remaining) if open_ else 0,
                "باز" if open_ and not it.payments else ("جزئی" if open_ else "تسویه"),
                jdate(settled), (settled - it.date).days if settled else None,
                (as_of - it.date).days if open_ else None, jdate(due), overdue,
                (info or {}).get("owner", ""),
                "؛ ".join(f"{jdate(p.date)} سند {p.doc_no}: {int(p.amount):,}" +
                         ("" if p.how == "named" else " (قدیمی‌ترین)" if p.how == "fifo" else " (از مانده قبل)")
                         for p in it.payments),
                it.description,
            ])
        for c in a.credits:
            rec_rows.append(who + [
                jdate(c.date), c.doc_no, c.fixed_no,
                {Kind.RECEIPT: "دریافت", Kind.RETURN: "برگشت از فروش", Kind.OPENING: "مانده بستانکار قبل"}
                .get(c.kind, "سایر بستانکار"),
                dict(Method.choices).get(c.method, "") if c.method else "", c.cheque_count or None,
                int(c.amount), c.invoice_no,
                "؛ ".join(f"{names.get(k, k)}: {int(v):,}" for k, v, _ in c.allocations),
                _n(c.left), c.description,
            ])
        for i in a.issues:
            issue_rows.append(who + [{"error": "خطا", "warn": "هشدار", "info": "توضیح"}[i.level], i.row_no, i.message])
    return inv_rows, rec_rows, issue_rows


INV_HEADERS = ["دفتر", "کد تفصیل", "مشتری", "نوع", "شماره فاکتور", "تاریخ", "شماره سند", "مبلغ",
               "برگشت از فروش", "دریافت‌شده", "مانده", "وضعیت", "تاریخ تسویه", "روز تا تسویه",
               "روز باز", "سررسید (آرپا)", "روز گذشته از سررسید", "بازاریاب فاکتور", "ریز دریافت‌ها", "شرح"]
INV_MONEY = {7, 8, 9, 10}
REC_HEADERS = ["دفتر", "کد تفصیل", "مشتری", "تاریخ", "شماره سند", "شماره ثابت", "نوع", "روش",
               "تعداد چک", "مبلغ", "بابت فاکتور", "تسویه‌شده با", "پیش دریافت", "شرح"]
REC_MONEY = {9, 12}
ISSUE_HEADERS = ["دفتر", "کد تفصیل", "مشتری", "سطح", "ردیف", "شرح"]


def collections_workbook(book: str, as_of: date, months: int = 11) -> bytes:
    """
    «گزارش وصولی» as finance keeps it — one sheet per book in its exact
    layout, then what it never had: «تجمیع» of both books per customer,
    «ریز فاکتورها» (every فاکتور, what paid it and when), «ریز دریافت‌ها»
    (every دریافت and the فاکتورها it settled) and «کنترل» (every place the
    ledger and the computation disagree).
    """
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    books = [LedgerBook.MAIN, LedgerBook.SECOND] if book == "all" else [book]
    for b in books:
        data = report(b, as_of, months)
        ws = wb.create_sheet(BOOK_LABEL[b][:30])
        _collections_sheet(ws, data, f"گزارش وصولی منتهی به تاریخ {jdate(as_of)} — {BOOK_LABEL[b]}")
    if book == "all":
        data = report("all", as_of, months)
        ws = wb.create_sheet("تجمیع دو دفتر")
        body = [[i, r["codes"].get(LedgerBook.MAIN, ""), r["codes"].get(LedgerBook.SECOND, ""), r["name"],
                 _n(r["by_book"].get(LedgerBook.MAIN)), _n(r["by_book"].get(LedgerBook.SECOND)),
                 f"=N(E{i + 1})+N(F{i + 1})", jdate(r["oldest_open"]), r["marketer"]]
                for i, r in enumerate(data["rows"], 1)]
        _table(ws, ["ردیف", "کد تفصیل اصلی", "کد تفصیل دوم", "نام طرف حساب", "مانده دفتر اصلی",
                    "مانده دفتر دوم", "جمع", "قدیمی‌ترین فاکتور باز", "بازاریاب"], body,
               balance_cols={4, 5, 6}, widths={3: 34}, total_label="مجموع", sum_cols={4, 5, 6})
    analyses = [analyse(acc, as_of) for b in books for acc in accounts(b) if acc.imported_at]
    inv_info = sales_invoices({a.account.customer_id for a in analyses})
    inv_rows, rec_rows, issue_rows = _detail_rows(analyses, inv_info, as_of)
    _table(wb.create_sheet("ریز فاکتورها"), INV_HEADERS, inv_rows, money_cols=INV_MONEY,
           widths={2: 30, 18: 60, 19: 50}, total_label="مجموع", sum_cols=INV_MONEY)
    _table(wb.create_sheet("ریز دریافت‌ها"), REC_HEADERS, rec_rows, money_cols=REC_MONEY,
           widths={2: 30, 11: 50, 13: 60}, total_label="مجموع", sum_cols=REC_MONEY)
    _table(wb.create_sheet("کنترل"), ISSUE_HEADERS, issue_rows or [["", "", "", "—", None, "همه‌ی دفترها با آرپا می‌خوانند."]],
           widths={2: 30, 5: 90})
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def customer_workbook(customer, as_of: date) -> bytes:
    """One customer, both books: summary, فاکتورها, دریافت‌ها and the ledger itself."""
    from openpyxl import Workbook

    st = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "خلاصه"
    ws.sheet_view.rightToLeft = True
    accs = list(LedgerAccount.objects.filter(customer=customer).prefetch_related("entries"))
    analyses = [analyse(acc, as_of) for acc in accs]
    ws["A1"] = f"صورت مطالبات {customer.name_fa} — منتهی به {jdate(as_of)}"
    ws["A1"].font = st["title_font"]
    body = []
    for a in analyses:
        acc = a.account
        body.append([BOOK_LABEL[acc.book], acc.party_code, jdate(acc.first_date), jdate(acc.last_date),
                     int(acc.opening_rial), int(sum((e.debit_rial for e in a.entries), ZERO)),
                     int(sum((e.credit_rial for e in a.entries), ZERO)), int(a.balance),
                     int(a.file_balance) if a.file_balance is not None else None,
                     "✓" if a.reconciled else "✗", _n(a.prepayment), len(a.open_items),
                     jdate(a.oldest_open), acc.note])
    _table(ws, ["دفتر", "کد تفصیل", "از تاریخ", "تا تاریخ", "موجودی قبل از دوره", "جمع بدهکار",
                "جمع بستانکار", "مانده", "مانده طبق آرپا", "کنترل", "پیش دریافت", "فاکتور باز",
                "قدیمی‌ترین فاکتور باز", "توضیحات"], body, money_cols={5, 6, 10}, balance_cols={4, 7, 8},
           start=3, total_label="جمع" if len(body) > 1 else None, sum_cols={4, 5, 6, 7, 8, 10})
    ws.freeze_panes = None
    inv_rows, rec_rows, issue_rows = _detail_rows(analyses, sales_invoices({customer.pk}), as_of)
    _table(wb.create_sheet("فاکتورها"), INV_HEADERS, inv_rows, money_cols=INV_MONEY,
           widths={2: 26, 18: 70, 19: 50}, total_label="مجموع", sum_cols=INV_MONEY)
    _table(wb.create_sheet("دریافت‌ها"), REC_HEADERS, rec_rows, money_cols=REC_MONEY,
           widths={2: 26, 11: 50, 13: 70}, total_label="مجموع", sum_cols=REC_MONEY)
    for a in analyses:
        def side(v):
            return "بد" if v > 0 else ("بس" if v < 0 else "--")

        ob = a.account.opening_rial
        body = [[BOOK_SHORT[a.account.book], None, "", "", jdate(a.account.first_date), None, None,
                 abs(int(ob)), side(ob), abs(int(ob)), "✓", "", "", "موجودی قبل از دوره"]]
        for e, bal in a.running:
            fb = e.file_balance_rial
            body.append([BOOK_SHORT[a.account.book], e.row_no, e.doc_no, e.fixed_no, jdate(e.date),
                         _n(e.debit_rial), _n(e.credit_rial), abs(int(bal)), side(bal),
                         abs(int(fb)) if fb is not None else None,
                         "✓" if fb in (None, bal) else "✗",
                         e.get_kind_display(), e.invoice_no, e.description])
        ws2 = wb.create_sheet(f"دفتر {BOOK_SHORT[a.account.book]}")
        _table(ws2, ["دفتر", "ردیف", "شماره سند", "شماره ثابت", "تاریخ", "بدهکار", "بستانکار",
                     "مانده (محاسبه)", "تشخیص", "مانده طبق آرپا", "کنترل", "نوع", "فاکتور", "شرح"], body,
               money_cols={5, 6, 7, 9}, widths={13: 80})
    if issue_rows:
        _table(wb.create_sheet("کنترل"), ISSUE_HEADERS, issue_rows, widths={2: 26, 5: 90})
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# The hand-kept «گزارش وصولی» — notes, marketers and figures to compare with
# ---------------------------------------------------------------------------
@dataclass
class CollectionsImport:
    sheet: str
    as_of: date | None
    matched: int = 0
    created: int = 0
    unmatched: list = field(default_factory=list)   # (code, name, balance)


def import_collections(fileobj, book: str) -> CollectionsImport:
    """
    Take «توضیحات», «بازاریاب» and «مانده بدهکاران» from finance's own
    workbook (its newest sheet), so the notes carry over and every row can be
    compared with what the ledger says. Customers are matched by کد تفصیل
    (آرپا's code in CRM), then by name; nothing is created in CRM.
    """
    from openpyxl import load_workbook

    from apps.crm.matching import name_key
    from apps.crm.models import Customer, CustomerExternalRef, ExternalSource

    try:
        wb = load_workbook(fileobj, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise LedgerFileError("فایل اکسل خوانده نشد.") from exc

    def sheet_key(ws):
        digits = re.sub(r"\D", "", fold(ws.title))
        return int(digits) if len(digits) == 8 else 0

    ws = max(wb.worksheets, key=lambda w: (sheet_key(w), wb.worksheets.index(w)))
    rows = list(ws.iter_rows(values_only=True))
    head_i, cols = None, {}
    for i, r in enumerate(rows[:10]):
        names = [fold(c) for c in r]
        if any("نام طرف" in n for n in names):
            head_i = i
            for j, n in enumerate(names):
                if "کد تفصیل" in n:
                    cols["code"] = j
                elif "نام طرف" in n:
                    cols["name"] = j
                elif "مانده بدهکار" in n:
                    cols["balance"] = j
                elif "توضیح" in n:
                    cols["note"] = j
                elif "بازاریاب" in n:
                    cols["marketer"] = j
            break
    if head_i is None or "name" not in cols:
        raise LedgerFileError("ستون «نام طرف حساب» در برگه‌ی آخر پیدا نشد.")
    title = " ".join(fold(c) for c in rows[0] if c) if head_i > 0 else ""
    as_of = parse_jalali(title)
    if as_of is None and sheet_key(ws):
        k = str(sheet_key(ws))
        as_of = jalali.to_gregorian(int(k[:4]), int(k[4:6]), int(k[6:]))
    result = CollectionsImport(sheet=ws.title, as_of=as_of)

    by_code = dict(CustomerExternalRef.objects.filter(source=ExternalSource.ARPA)
                   .values_list("external_id", "customer_id"))
    by_name: dict[str, list[int]] = {}
    for cid, name in Customer.objects.values_list("id", "name_fa"):
        by_name.setdefault(name_key(name), []).append(cid)

    def cell(r, k):
        j = cols.get(k)
        return r[j] if j is not None and j < len(r) else None

    with transaction.atomic():
        for r in rows[head_i + 1:]:
            name = _text(cell(r, "name"))
            if not name or name.startswith("مجموع"):
                continue
            code = _text(cell(r, "code"))
            balance = cell(r, "balance")
            cid = by_code.get(code)
            if cid is None:
                key = name_key(name)
                hits = by_name.get(key, [])
                if not hits and len(key) >= 4:
                    # The workbook writes «اجاقلو» for «آقای اجاقلو (زنجان)».
                    hits = [i for k, ids in by_name.items() if key in k for i in ids]
                cid = hits[0] if len(hits) == 1 else None
            if cid is None:
                result.unmatched.append((code, name, int(_money(balance)) if balance not in (None, "") else None))
                continue
            acc, created = LedgerAccount.objects.get_or_create(book=book, customer_id=cid)
            note, marketer = _text(cell(r, "note")), _text(cell(r, "marketer"))
            if note and note != "0":
                acc.note = note[:400]
            if marketer:
                acc.marketer = marketer[:100]
            if code and not acc.party_code:
                acc.party_code = code[:32]
            if isinstance(balance, (int, float, Decimal)):
                acc.manual_balance_rial = _money(balance)
                acc.manual_as_of = as_of
            acc.save()
            result.matched += 1
            result.created += int(created)
    return result
