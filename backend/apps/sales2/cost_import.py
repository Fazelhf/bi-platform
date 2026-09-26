"""
Importing a month's فی حسابداری from accounting's Excel — preview, then confirm.

The file is read by header text, not column letter: accounting's workbooks
move columns around. Any sheet with «نام کالا» and a «فی» column is read (the
rightmost «فی» is accounting's), with «گرماژ» when it has one. That covers
both the monthly commission workbook and a plain three-column cost sheet.

Nothing is written until the preview has been seen: `preview` says, per
product and grammage, whether the month's cost is new, changed, the same or
missing from the file (which keeps the earlier cost), and lists every name the
catalogue does not know. `apply` writes exactly that.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal

from django.db import transaction

from apps.crm.models import Product
from apps.sales2 import pricing
from apps.sales2.models import AccountingCost

_NORM = str.maketrans({"ي": "ی", "ك": "ک", "‌": " "})


def norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "").translate(_NORM)).strip()


def product_key(text) -> str:
    """«57-40 -لوله 40» and «57 - 40 - لوله 40» are one product to a reader."""
    return re.sub(r"\s*-\s*", " - ", norm(text))


@dataclass
class Entry:
    name: str
    grammage: int | None
    cost: Decimal
    where: str


def read_workbook(fileobj) -> list[Entry]:
    from openpyxl import load_workbook

    # data_only: some rows carry the cost as a formula (=R9); its cached
    # result is the figure accounting meant.
    wb = load_workbook(fileobj, data_only=True)
    out: list[Entry] = []
    for ws in wb.worksheets:
        for hdr_row in ws.iter_rows(min_row=1, max_row=5):
            cols: dict[str, list[str]] = {}
            for c in hdr_row:
                if isinstance(c.value, (str, int)) and str(c.value).strip():
                    cols.setdefault(norm(c.value), []).append(c.column_letter)
            fi = cols.get("فی", []) + cols.get("فی حسابداری", [])
            if "نام کالا" not in cols or not fi:
                continue
            acc = max(fi, key=lambda c: (len(c), c))
            name_col, gram_col = cols["نام کالا"][0], cols.get("گرماژ", [None])[0]
            for r in range(hdr_row[0].row + 1, ws.max_row + 1):
                name = product_key(ws[f"{name_col}{r}"].value)
                cost = ws[f"{acc}{r}"].value
                if not name or not isinstance(cost, (int, float)) or cost <= 0:
                    continue
                gram = ws[f"{gram_col}{r}"].value if gram_col else None
                out.append(Entry(name, int(gram) if gram in (48, 55) else None,
                                 Decimal(str(round(cost))), f"{ws.title}:{r}"))
            break
    return out


def preview(entries: list[Entry], jy: int, jm: int) -> dict:
    """What confirming would do — per product/grammage — and what it cannot place."""
    products = {product_key(p.name_fa): p for p in Product.objects.all()}
    latest: dict[tuple[int, int], tuple[Decimal, str]] = {}
    unknown: dict[str, str] = {}
    for e in entries:
        prod = products.get(e.name)
        if not prod:
            unknown[e.name] = e.where
            continue
        is_roll = pricing.roll_size(prod) is not None
        if is_roll and e.grammage is None:
            unknown[f"{e.name} (بدون گرماژ)"] = e.where
            continue
        latest[(prod.id, e.grammage if is_roll else 0)] = (e.cost, e.where)  # later rows win

    key = jy * 100 + jm
    before: dict[tuple[int, int], AccountingCost] = {}
    for c in AccountingCost.objects.select_related("product").order_by("jalali_year", "jalali_month"):
        if c.month_key <= key:
            before[(c.product_id, c.grammage)] = c  # ascending: the last one kept is in force

    names = {p.id: p.name_fa for p in products.values()}
    rows = []
    for (pid, gram), (cost, where) in latest.items():
        prev = before.get((pid, gram))
        prev_cost = prev.cost_rial if prev else None
        status = "new" if prev is None else ("same" if prev_cost == cost else "changed")
        rows.append({"product": pid, "name": names[pid], "grammage": gram,
                     "previous_rial": None if prev_cost is None else str(prev_cost),
                     "previous_month": prev.month_key if prev else None,
                     "cost_rial": str(cost), "status": status, "where": where})
    for (pid, gram), prev in before.items():
        if (pid, gram) not in latest:
            rows.append({"product": pid, "name": prev.product.name_fa, "grammage": gram,
                         "previous_rial": str(prev.cost_rial), "previous_month": prev.month_key,
                         "cost_rial": None, "status": "kept", "where": ""})
    order = {"changed": 0, "new": 1, "kept": 2, "same": 3}
    rows.sort(key=lambda r: (order[r["status"]], r["name"]))
    counts = {s: sum(1 for r in rows if r["status"] == s) for s in order}
    return {"jalali_year": jy, "jalali_month": jm, "rows": rows, "counts": counts,
            "unknown": [{"name": n, "where": w} for n, w in sorted(unknown.items())]}


@transaction.atomic
def apply(entries: list[Entry], jy: int, jm: int, create_missing: bool = False,
          source: str = "") -> dict:
    """Write the month's costs. Earlier months are never touched."""
    created = 0
    if create_missing:
        known = {product_key(p.name_fa) for p in Product.objects.all()}
        for e in entries:
            if e.name not in known:
                n = 1
                while Product.objects.filter(code=f"s2-{n}").exists():
                    n += 1
                Product.objects.create(code=f"s2-{n}", name_fa=e.name, unit=Product.Unit.ROLL)
                known.add(e.name)
                created += 1
    result = preview(entries, jy, jm)
    written = 0
    for r in result["rows"]:
        if r["status"] in ("new", "changed", "same") and r["cost_rial"] is not None:
            AccountingCost.objects.update_or_create(
                product_id=r["product"], grammage=r["grammage"], jalali_year=jy, jalali_month=jm,
                defaults={"cost_rial": Decimal(r["cost_rial"]),
                          "source": (source or "اکسل حسابداری")[:80] + f" ({r['where']})"[:40]},
            )
            written += 1
    return {**result, "written": written, "products_created": created}
